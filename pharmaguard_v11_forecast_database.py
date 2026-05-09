# -*- coding: utf-8 -*-
# Single-file PharmaGuard build for IDLE
import sys
import math
import re
import traceback
import warnings

warnings.filterwarnings(
    "ignore",
    message=r"Failed to disconnect.*",
    category=RuntimeWarning,
)
from dataclasses import dataclass
from difflib import SequenceMatcher
import sqlite3
# Suppress noisy matplotlib layout warnings inside IDLE / Qt embedding
warnings.filterwarnings(
    "ignore",
    message=r"This figure was using a layout engine that is incompatible with subplots_adjust.*",
    category=UserWarning,
)
from datetime import date, datetime
from pathlib import Path
import pandas as pd
import numpy as np
from PySide6.QtCore import Qt, QSize, QTimer, QStringListModel
from PySide6.QtGui import QAction, QColor, QFont, QIcon, QPixmap
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QComboBox, QCompleter, QDockWidget, QFileDialog, QFormLayout,
    QInputDialog,
    QFrame, QGridLayout, QGroupBox, QHBoxLayout, QLabel, QLineEdit, QListWidget,
    QListWidgetItem, QMainWindow, QMessageBox, QPushButton, QPlainTextEdit,
    QScrollArea, QSplitter, QStackedWidget, QStatusBar, QTableWidget, QTableWidgetItem,
    QToolBar, QVBoxLayout, QWidget, QHeaderView, QTextEdit, QDoubleSpinBox,
    QAbstractItemView, QSplashScreen
)
try:
    from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
    from matplotlib.figure import Figure
    HAS_MATPLOTLIB = True
except Exception:
    HAS_MATPLOTLIB = False
    FigureCanvas = None
    Figure = None
try:
    import arabic_reshaper
    from bidi.algorithm import get_display as bidi_get_display
    HAS_ARABIC_SHAPING = True
except Exception:
    arabic_reshaper = None
    bidi_get_display = None
    HAS_ARABIC_SHAPING = False
try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except Exception:
    load_workbook = None
    Alignment = Font = PatternFill = None
    get_column_letter = None
    HAS_OPENPYXL = False
warnings.filterwarnings(
    "ignore",
    message="Tight layout not applied.*",
    category=UserWarning,
)
warnings.filterwarnings(
    "ignore",
    message="This figure was using a layout engine.*",
    category=UserWarning,
)
APP_VERSION = "V9 Stable Import Preview"
APP_TITLE_EN = f"PharmaGuard AI Pro {APP_VERSION}"
APP_TITLE_AR = f"فارماجارد برو {APP_VERSION}"
def _contains_arabic(text_value):
    s = str(text_value)
    return any("\u0600" <= ch <= "\u06FF" or "\u0750" <= ch <= "\u077F" or "\u08A0" <= ch <= "\u08FF" for ch in s)
def ui_text(text_value):
    s = "" if text_value is None else str(text_value)
    if HAS_ARABIC_SHAPING and _contains_arabic(s):
        try:
            return bidi_get_display(arabic_reshaper.reshape(s))
        except Exception:
            return s
    return s
def chart_font_family():
    # Common Windows fonts that usually render Arabic correctly.
    return "Tahoma"
def normalize_search_text(text_value):
    s = "" if text_value is None else str(text_value)
    s = s.strip().lower()
    if not s:
        return ""
    s = re.sub(r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED\u0640]", "", s)
    repl = {
        "أ": "ا", "إ": "ا", "آ": "ا",
        "ى": "ي", "ئ": "ي", "ؤ": "و",
        "ة": "ه",
    }
    for k, v in repl.items():
        s = s.replace(k, v)
    s = re.sub(r"[^\w\s\u0600-\u06FF]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s
def _latin_pharma_phonetic(value):
    s = normalize_search_text(value)
    if not s:
        return ""
    s = s.lower()
    pairs = [
        ("ph", "f"), ("qu", "k"), ("ck", "k"), ("x", "ks"), ("v", "f"),
        ("tion", "shn"), ("sion", "zhn"), ("ch", "k"), ("sh", "sh"),
        ("th", "t"), ("gh", "g"), ("ou", "u"), ("oo", "u"), ("ee", "i"),
    ]
    for a, b in pairs:
        s = s.replace(a, b)
    out = []
    i = 0
    while i < len(s):
        ch = s[i]
        nxt = s[i+1] if i + 1 < len(s) else ''
        if ch == 'c':
            out.append('s' if nxt in 'eiy' else 'k')
        elif ch == 'g':
            out.append('j' if nxt in 'eiy' else 'g')
        elif ch == 'q':
            out.append('k')
        elif ch == 'z':
            out.append('s')
        elif ch.isalnum() or ch.isspace():
            out.append(ch)
        i += 1
    s = ''.join(out)
    s = re.sub(r'[aeiou]', '', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s
def _arabic_to_latin_phonetic(value):
    s = normalize_search_text(value)
    if not s:
        return ""
    table = {
        'ا':'a','ب':'b','ت':'t','ث':'s','ج':'j','ح':'h','خ':'kh','د':'d','ذ':'z','ر':'r','ز':'z',
        'س':'s','ش':'sh','ص':'s','ض':'d','ط':'t','ظ':'z','ع':'a','غ':'g','ف':'f','ق':'k','ك':'k',
        'ل':'l','م':'m','ن':'n','ه':'h','و':'w','ي':'y','ئ':'y','ؤ':'w','ة':'h','ى':'a'
    }
    out = []
    for ch in s:
        if ch in table:
            out.append(table[ch])
        elif ch.isascii() and (ch.isalnum() or ch.isspace()):
            out.append(ch.lower())
        elif ch.isspace():
            out.append(' ')
    s = ''.join(out)
    s = re.sub(r'[aeiou]', '', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s
def build_search_blob_from_values(*values):
    parts = []
    for value in values:
        if value is None:
            continue
        raw = str(value)
        if not raw or raw.lower() == 'nan':
            continue
        norm = normalize_search_text(raw)
        if norm:
            parts.append(norm)
        lat = _latin_pharma_phonetic(raw)
        if lat:
            parts.append(lat)
        ar_lat = _arabic_to_latin_phonetic(raw)
        if ar_lat:
            parts.append(ar_lat)
    blob = ' '.join(parts)
    blob = re.sub(r'\s+', ' ', blob).strip()
    return blob
def build_search_blob_series(df, cols):
    cols = [c for c in cols if c in df.columns]
    if not cols:
        return pd.Series('', index=df.index)
    return df[cols].fillna('').astype(str).apply(lambda r: build_search_blob_from_values(*r.tolist()), axis=1)
def normalize_query_bundle(value):
    raw = str(value or '')
    norm = normalize_search_text(raw)
    return {
        'text': norm,
        'latin': _latin_pharma_phonetic(raw),
        'ar_latin': _arabic_to_latin_phonetic(raw),
    }
def _is_subsequence(needle, haystack):
    if not needle or not haystack:
        return False
    it = iter(haystack)
    return all(ch in it for ch in needle)
def query_match_score(bundle, blob):
    blob = str(blob or '')
    if not blob:
        return 0.0
    scores = []
    for key in ('text', 'latin', 'ar_latin'):
        q = bundle.get(key, '')
        if not q:
            continue
        if q in blob:
            scores.append(1.0)
            continue
        if any(part and part in blob for part in q.split()):
            scores.append(0.82)
            continue
        if _is_subsequence(q, blob):
            scores.append(0.72)
            continue
        ratio = SequenceMatcher(None, q, blob[: max(len(q) * 4, len(q))]).ratio()
        if ratio >= 0.72:
            scores.append(ratio)
    return max(scores) if scores else 0.0
def query_in_blob(bundle, blob):
    return query_match_score(bundle, blob) > 0.0
TEXT = {
    "en": {
        "title": APP_TITLE_EN,
        "welcome": "Government pharmacy supply chain decision-support app using scientific-name analysis",
        "subtitle": "Scroll-enabled, bilingual, chart-based, with smart Excel/CSV/Google Sheet import and stronger supply-chain analytics.",
        "toggle_nav": "Toggle navigation",
        "toggle_side": "Toggle tools panel",
        "collapse_all": "Collapse side panels",
        "expand_all": "Show side panels",
        "quick_screen": "Quick screen",
        "language": "Language",
        "theme": "Theme",
        "light": "Light",
        "dark": "Dark",
        "search": "Search medicine / branch / supplier",
        "alerts_only": "Alerts only",
        "import_excel": "📥 Import Excel / CSV",
        "import_google": "☁ Import Google Sheet",
        "purchase_horizon": "Consumption horizon",
        "months_1": "1 month",
        "months_2": "2 months",
        "months_3": "3 months",
        "months_6": "6 months",
        "months_12": "12 months",
        "stock_status": "Stock status",
        "overstock_risk": "Overstock risk",
        "predict_30": "Predicted stock (30d)",
        "predict_60": "Predicted stock (60d)",
        "predict_90": "Predicted stock (90d)",
        "google_sheet_prompt": "Paste a Google Sheet link (published CSV or standard sheet link).",
        "google_sheet_ok": "Google Sheet imported successfully.",
        "load_demo": "Load demo data",
        "refresh": "Refresh analysis",
        "export_workbook": "Export analysis workbook",
        "export_report": "Export final report",
        "export_compare": "Export branch comparison",
        "type_to_search": "Search by scientific/trade name or any close letters...",
        "final_report_saved": "Final report exported successfully.",
        "sheet_used": "Sheet used",
        "mapping_report": "Mapping report",
        "processing_log": "Processing log",
        "preview": "Preview",
        "raw_detected": "Messy / raw workbook was normalized successfully.",
        "import_ok": "Data imported successfully.",
        "import_fail": "Import failed.",
        "no_data": "No data available for this screen.",
        "save_ok": "File saved successfully.",
        "decision_only": "Decision support only. No purchasing or redistribution is executed automatically.",
        "kpi_total": "Total items",
        "kpi_critical": "Critical shortage alerts",
        "kpi_expiry": "High/Critical expiry alerts",
        "kpi_reorder": "Suggested reorder units",
        "kpi_red": "Red escalations",
        "nav_dashboard": "📊 Executive Dashboard",
        "nav_import": "📥 Import & Mapping",
        "nav_inventory": "🧠 Inventory Analyzer",
        "nav_purchase": "🛒 Purchase Requests",
        "nav_redistribution": "🔄 Redistribution",
        "nav_fefo": "FEFO & Expiry",
        "nav_abcven": "ABC-VEN",
        "nav_supplier": "🏭 Supplier Reliability",
        "nav_safety": "Safety Stock",
        "nav_emergency": "Emergency Escalation",
        "nav_quality": "Data Quality & Audit",
        "nav_governance": "Governance & Prompts",
        "filters": "Filters",
        "branch": "Branch",
        "risk": "Risk",
        "priority": "Clinical priority",
        "medicine_lookup": "Medicine lookup",
        "branch_compare": "Compare branches for selected medicine",
        "what_if": "What-if scenario",
        "lead_time": "Lead time (days)",
        "pending_po": "Pending PO quantity",
        "extra_stock": "Additional stock received",
        "run_scenario": "Run scenario",
        "generate_draft": "Generate draft",
        "save_txt": "Save as TXT",
        "data_quality": "Data quality checks",
        "prompt_library": "Prompt library",
        "governance_guardrails": "Governance guardrails",
        "failure_modes": "Failure modes",
        "dock_navigation": "Navigation",
        "dock_tools": "Tools",
        "quick_short": "Quick",
        "search_short": "Search",
        "all": "All",
        "inventory_selector": "Medicine list",
        "summary_generic_name": "Scientific name",
        "summary_trade_names": "Trade names",
        "summary_source_names": "Source item names",
        "branch_compare_title": "Branch comparison",
        "scenario_title": "What-if scenario",
        "import_group_title": "Import, mapping, and preview",
        "sheet_used_value": "Sheet used",
        "report_title": "Mapping report",
        "preview_title": "Preview",
        "purchase_draft_title": "Purchase request draft",
        "summary_item_code": "Item code",
        "summary_branch": "Branch",
        "summary_current_stock": "Current stock",
        "summary_dynamic_min": "Dynamic minimum stock",
        "summary_avg_daily": "Average daily consumption",
        "summary_avg_monthly": "Average monthly consumption",
        "summary_lead_time": "Lead time",
        "summary_pending_po": "Pending PO",
        "summary_shortage_risk": "Shortage risk",
        "summary_expiry_risk": "Expiry risk",
        "summary_escalation": "Escalation",
        "summary_supplier": "Supplier",
        "summary_on_time_rate": "On-time rate",
        "summary_reason": "Reason",
        "summary_recommended_action": "Recommended action",
        "scenario_result_title": "Scenario result for",
        "showing_rows": "Showing first {n} rows for performance.",
        "chart_reorder_title": "Suggested reorder quantity",
        "chart_risk_title": "Shortage risk distribution",
        "chart_abc_title": "ABC class distribution",
        "chart_ven_title": "VEN class distribution",
        "chart_supplier_title": "Supplier on-time reliability",
        "matplotlib_missing": "matplotlib not installed.",
    },
    "ar": {
        "title": APP_TITLE_AR,
        "welcome": "برنامج دعم قرار لسلسلة إمداد الصيدليات الحكومية مع التحليل بالاسم العلمي",
        "subtitle": "يدعم التمرير، عربي/إنجليزي، رسوم بيانية، واستيراد ذكي من Excel/CSV/Google Sheet مع تحليلات أقوى لسلسلة الإمداد.",
        "toggle_nav": "إظهار/إخفاء القائمة",
        "toggle_side": "إظهار/إخفاء لوحة الأدوات",
        "collapse_all": "إغلاق النوافذ الجانبية",
        "expand_all": "إظهار النوافذ الجانبية",
        "quick_screen": "اختيار الشاشة",
        "language": "اللغة",
        "theme": "النمط",
        "light": "فاتح",
        "dark": "داكن",
        "search": "بحث بالصنف / الفرع / المورد",
        "alerts_only": "التنبيهات فقط",
        "import_excel": "📥 استيراد Excel / CSV",
        "import_google": "☁ استيراد Google Sheet",
        "purchase_horizon": "فترة الاستهلاك",
        "months_1": "شهر",
        "months_2": "شهران",
        "months_3": "3 شهور",
        "months_6": "6 شهور",
        "months_12": "12 شهر",
        "stock_status": "حالة المخزون",
        "overstock_risk": "خطورة التكدس",
        "predict_30": "الرصيد المتوقع بعد 30 يوم",
        "predict_60": "الرصيد المتوقع بعد 60 يوم",
        "predict_90": "الرصيد المتوقع بعد 90 يوم",
        "google_sheet_prompt": "ضع رابط Google Sheet هنا.",
        "google_sheet_ok": "تم استيراد Google Sheet بنجاح.",
        "load_demo": "تحميل بيانات تجريبية",
        "refresh": "تحديث التحليل",
        "export_workbook": "تصدير ملف التحليل",
        "export_report": "تصدير التقرير النهائي",
        "export_compare": "تصدير مقارنة الفروع",
        "type_to_search": "ابحث بالاسم العلمي أو التجاري أو أي حروف قريبة...",
        "final_report_saved": "تم تصدير التقرير النهائي بنجاح.",
        "sheet_used": "الشيت المستخدم",
        "mapping_report": "تقرير الربط",
        "processing_log": "سجل المعالجة",
        "preview": "معاينة",
        "raw_detected": "تمت معالجة الملف الخام/غير المرتب بنجاح.",
        "import_ok": "تم استيراد البيانات بنجاح.",
        "import_fail": "فشل الاستيراد.",
        "no_data": "لا توجد بيانات متاحة لهذه الشاشة.",
        "save_ok": "تم حفظ الملف بنجاح.",
        "decision_only": "النظام للدعم واتخاذ القرار فقط. لا ينفذ شراء أو تحويل تلقائيًا.",
        "kpi_total": "إجمالي الأصناف",
        "kpi_critical": "تنبيهات النقص الحرجة",
        "kpi_expiry": "تنبيهات الصلاحية العالية/الحرجة",
        "kpi_reorder": "إجمالي المقترح للشراء",
        "kpi_red": "التصعيد الأحمر",
        "nav_dashboard": "لوحة القيادة",
        "nav_import": "📥 الاستيراد والربط",
        "nav_inventory": "🧠 تحليل المخزون",
        "nav_purchase": "🛒 طلبات الشراء",
        "nav_redistribution": "🔄 إعادة التوزيع",
        "nav_fefo": "الصلاحية و FEFO",
        "nav_abcven": "تحليل ABC-VEN",
        "nav_supplier": "🏭 موثوقية الموردين",
        "nav_safety": "مخزون الأمان",
        "nav_emergency": "التصعيد العاجل",
        "nav_quality": "جودة البيانات والتدقيق",
        "nav_governance": "الحوكمة والبرومبتات",
        "filters": "الفلاتر",
        "branch": "الفرع",
        "risk": "الخطورة",
        "priority": "الأولوية السريرية",
        "medicine_lookup": "بحث الصنف",
        "branch_compare": "مقارنة الفروع للصنف المختار",
        "what_if": "تحليل ماذا لو",
        "lead_time": "مدة التوريد (يوم)",
        "pending_po": "كمية أمر الشراء المعلق",
        "extra_stock": "زيادة في المخزون",
        "run_scenario": "تشغيل السيناريو",
        "generate_draft": "إنشاء مسودة",
        "save_txt": "حفظ TXT",
        "data_quality": "اختبارات جودة البيانات",
        "prompt_library": "مكتبة البرومبتات",
        "governance_guardrails": "ضوابط الحوكمة",
        "failure_modes": "أنماط الفشل",
        "dock_navigation": "القائمة",
        "dock_tools": "الأدوات",
        "quick_short": "تنقل سريع",
        "search_short": "بحث",
        "all": "الكل",
        "inventory_selector": "قائمة الأدوية",
        "summary_generic_name": "الاسم العلمي",
        "summary_trade_names": "الأسماء التجارية",
        "summary_source_names": "أسماء الأصناف الأصلية",
        "branch_compare_title": "مقارنة الفروع",
        "scenario_title": "سيناريو ماذا لو",
        "import_group_title": "الاستيراد والربط والمعاينة",
        "sheet_used_value": "الشيت المستخدم",
        "report_title": "تقرير الربط",
        "preview_title": "معاينة",
        "purchase_draft_title": "مسودة طلب الشراء",
        "summary_item_code": "كود الصنف",
        "summary_branch": "الفرع",
        "summary_current_stock": "الرصيد الحالي",
        "summary_dynamic_min": "الحد الأدنى الديناميكي",
        "summary_avg_daily": "متوسط الاستهلاك اليومي",
        "summary_avg_monthly": "متوسط الاستهلاك الشهري",
        "summary_lead_time": "مدة التوريد",
        "summary_pending_po": "الكمية المعلقة",
        "summary_shortage_risk": "خطورة النقص",
        "summary_expiry_risk": "خطورة الصلاحية",
        "summary_escalation": "مستوى التصعيد",
        "summary_supplier": "المورد",
        "summary_on_time_rate": "معدل الالتزام",
        "summary_reason": "سبب التنبيه",
        "summary_recommended_action": "الإجراء المقترح",
        "scenario_result_title": "نتيجة السيناريو للصنف",
        "showing_rows": "يتم عرض أول {n} صف لتحسين الأداء.",
        "chart_reorder_title": "الكميات المقترحة لإعادة الطلب",
        "chart_risk_title": "توزيع خطورة النقص",
        "chart_abc_title": "توزيع فئات ABC",
        "chart_ven_title": "توزيع فئات VEN",
        "chart_supplier_title": "موثوقية الموردين في الالتزام",
        "matplotlib_missing": "مكتبة matplotlib غير مثبتة.",
    },
}
NAV_ORDER = [
    "dashboard", "import", "inventory", "forecast", "stockbalance", "connector", "purchase", "redistribution",
    "fefo", "abcven", "supplier", "safety", "emergency", "quality", "governance"
]
NAV_TEXT = {k: f"nav_{k}" for k in NAV_ORDER}
RISK_TRANSLATIONS = {
    "Low": {"ar": "منخفض", "en": "Low"},
    "Moderate": {"ar": "متوسط", "en": "Moderate"},
    "High": {"ar": "مرتفع", "en": "High"},
    "Critical": {"ar": "حرج", "en": "Critical"},
    "Expired": {"ar": "منتهي", "en": "Expired"},
    "Unknown": {"ar": "غير معروف", "en": "Unknown"},
    "Green": {"ar": "أخضر", "en": "Green"},
    "Amber": {"ar": "أصفر", "en": "Amber"},
    "Red": {"ar": "أحمر", "en": "Red"},
    "Main Store": {"ar": "المخزن الرئيسي", "en": "Main Store"},
    "Main Branch": {"ar": "الفرع الرئيسي", "en": "Main Branch"},
    "Unknown Supplier": {"ar": "مورد غير معروف", "en": "Unknown Supplier"},
    "Balanced": {"ar": "متوازن", "en": "Balanced"},
    "Overstock": {"ar": "تكدس", "en": "Overstock"},
    "Shortage": {"ar": "نقص", "en": "Shortage"},
}
DOSAGE_FORM_TRANSLATIONS = {
    "Tablets": {"ar": "أقراص", "en": "Tablets"},
    "Ampoules": {"ar": "أمبولات", "en": "Ampoules"},
    "Solutions": {"ar": "محاليل", "en": "Solutions"},
    "Miscellaneous": {"ar": "متنوعات", "en": "Miscellaneous"},
}
COLUMN_LABELS = {
    "medicine_name": {"ar": "اسم الصنف التحليلي", "en": "Medicine"},
    "generic_name": {"ar": "الاسم العام", "en": "Generic"},
    "brand_name": {"ar": "اسم تجاري", "en": "Trade name"},
    "brand_names_display": {"ar": "الأسماء التجارية", "en": "Trade names"},
    "dosage_form_group": {"ar": "الشكل الصيدلاني", "en": "Dosage form"},
    "analysis_name": {"ar": "تحليل الاسم العلمي + الشكل", "en": "Analysis name"},
    "source_medicine_names": {"ar": "أسماء الأصناف الأصلية", "en": "Source names"},
    "item_code": {"ar": "كود الصنف", "en": "Item code"},
    "current_stock": {"ar": "الرصيد الحالي", "en": "Current stock"},
    "min_stock_level": {"ar": "الحد الأدنى", "en": "Min stock"},
    "avg_daily_consumption": {"ar": "متوسط يومي", "en": "Avg daily use"},
    "avg_monthly_consumption": {"ar": "متوسط شهري", "en": "Avg monthly use"},
    "lead_time_days": {"ar": "مدة التوريد", "en": "Lead time"},
    "pending_po_qty": {"ar": "كمية معلقة", "en": "Pending PO"},
    "expiry_date": {"ar": "تاريخ الصلاحية", "en": "expiry_date"},
    "clinical_priority": {"ar": "الأولوية السريرية", "en": "Clinical priority"},
    "storage_location": {"ar": "مكان التخزين", "en": "Storage location"},
    "branch_name": {"ar": "الفرع", "en": "Branch"},
    "unit_cost": {"ar": "تكلفة الوحدة", "en": "Unit cost"},
    "supplier_name": {"ar": "المورد", "en": "Supplier"},
    "supplier_on_time_rate": {"ar": "معدل التزام المورد", "en": "Supplier on-time rate"},
    "ven_class": {"ar": "فئة VEN", "en": "VEN class"},
    "abc_class": {"ar": "فئة ABC", "en": "ABC class"},
    "days_of_stock_left": {"ar": "أيام التغطية", "en": "Days left"},
    "days_to_expiry": {"ar": "أيام حتى الانتهاء", "en": "Days to expiry"},
    "dynamic_min_stock": {"ar": "الحد الأدنى الديناميكي", "en": "Dynamic min"},
    "expected_stock_at_lead_time": {"ar": "الرصيد المتوقع عند التوريد", "en": "Stock at lead time"},
    "recommended_reorder_qty": {"ar": "الكمية المقترحة", "en": "Suggested reorder qty"},
    "shortage_score": {"ar": "درجة النقص", "en": "Shortage score"},
    "shortage_risk": {"ar": "خطورة النقص", "en": "Shortage risk"},
    "expiry_risk": {"ar": "خطورة الصلاحية", "en": "Expiry risk"},
    "expected_unused_before_expiry": {"ar": "فائض متوقع قبل الانتهاء", "en": "expected_unused_before_expiry"},
    "reason_explanation": {"ar": "سبب التنبيه", "en": "Reason"},
    "escalation_level": {"ar": "مستوى التصعيد", "en": "Escalation"},
    "recommended_action": {"ar": "الإجراء المقترح", "en": "Recommended action"},
    "from_branch": {"ar": "من فرع", "en": "from_branch"},
    "to_branch": {"ar": "إلى فرع", "en": "to_branch"},
    "suggested_transfer_qty": {"ar": "كمية التحويل المقترحة", "en": "suggested_transfer_qty"},
    "source_surplus": {"ar": "فائض المصدر", "en": "source_surplus"},
    "target_gap": {"ar": "فجوة الهدف", "en": "target_gap"},
    "note": {"ar": "ملاحظة", "en": "note"},
    "items": {"ar": "عدد الأصناف", "en": "items"},
    "avg_on_time_rate": {"ar": "متوسط الالتزام", "en": "avg_on_time_rate"},
    "critical_items": {"ar": "أصناف حرجة", "en": "critical_items"},
    "critical_shortages": {"ar": "نواقص حرجة", "en": "critical_shortages"},
    "column": {"ar": "العمود", "en": "column"},
    "missing_count": {"ar": "عدد القيم المفقودة", "en": "missing_count"},
    "mapping_report": {"ar": "تقرير الربط", "en": "mapping_report"},
    "stock_status": {"ar": "حالة المخزون", "en": "Stock status"},
    "overstock_risk": {"ar": "خطورة التكدس", "en": "Overstock risk"},
    "predicted_stock_30d": {"ar": "الرصيد المتوقع 30 يوم", "en": "Predicted stock 30d"},
    "predicted_stock_60d": {"ar": "الرصيد المتوقع 60 يوم", "en": "Predicted stock 60d"},
    "predicted_stock_90d": {"ar": "الرصيد المتوقع 90 يوم", "en": "Predicted stock 90d"},
    "selected_horizon_months": {"ar": "فترة الاستهلاك", "en": "Horizon months"},
    "recommended_reorder_qty_horizon": {"ar": "كمية الشراء حسب الفترة", "en": "Reorder qty (horizon)"},
}
def translate_display_value(value, lang):
    if pd.isna(value):
        return ""
    text = str(value)
    if text in DOSAGE_FORM_TRANSLATIONS:
        return DOSAGE_FORM_TRANSLATIONS[text][lang]
    if lang == "ar" and text in RISK_TRANSLATIONS:
        return RISK_TRANSLATIONS[text]["ar"]
    return text
def localize_dataframe_for_display(df, lang):
    if df is None or df.empty:
        return df
    out = df.copy()
    rename_map = {}
    for col in out.columns:
        if col in COLUMN_LABELS:
            rename_map[col] = COLUMN_LABELS[col][lang]
    out = out.rename(columns=rename_map)
    for col in out.columns:
        out[col] = out[col].map(lambda x: translate_display_value(x, lang) if isinstance(x, str) or not pd.api.types.is_numeric_dtype(out[col]) else x)
    return out
CORE_COLUMNS = [
    "medicine_name",
    "item_code",
    "current_stock",
    "min_stock_level",
    "avg_daily_consumption",
    "lead_time_days",
    "pending_po_qty",
    "expiry_date",
    "clinical_priority",
    "storage_location",
]
OPTIONAL_COLUMNS = [
    "branch_name",
    "avg_monthly_consumption",
    "unit_cost",
    "supplier_name",
    "supplier_on_time_rate",
    "ven_class",
    "abc_class",
    "generic_name",
    "brand_name",
    "brand_names_display",
    "dosage_form_group",
    "analysis_name",
    "source_medicine_names",
    "stock_status",
    "overstock_risk",
    "predicted_stock_30d",
    "predicted_stock_60d",
    "predicted_stock_90d",
    "selected_horizon_months",
    "recommended_reorder_qty_horizon",
]
ALL_COLUMNS = CORE_COLUMNS + OPTIONAL_COLUMNS
MAX_TABLE_ROWS = 800
CHART_TOP_N = 8
PROMPTS = {
    "Shortage Risk Analysis": """You are PharmaGuard AI, a pharmacy supply chain assistant for a governmental hospital.
Analyze the provided medicine inventory, compute coverage, classify shortage risk, explain the main drivers, and suggest a safe human-reviewed action. If only monthly consumption is available, convert it to daily use appropriately. Never invent missing data.""",
    "Emergency Escalation Prompt": """Review one medicine item and classify escalation as Green, Amber, or Red. Explain why escalation is needed, whether branch redistribution should be checked, and whether an urgent purchase draft is required. Do not authorize any purchasing action.""",
    "Redistribution Prompt": """Review stock balances for the same medicine across branches. Identify branches below dynamic minimum stock and branches with transferable surplus. Recommend a safe redistribution suggestion while preserving minimum coverage in source branches.""",
}
COLUMN_ALIASES = {
    "medicine_name": [
        "medicine_name", "medicine", "item", "item_name", "drug", "drug_name", "name",
        "الصنف", "اسم الصنف", "الدواء", "اسم الدواء", "المستحضر", "الصنف / التركيز", "اسم الصنف/التركيز"
    ],
    "generic_name": [
        "generic_name", "generic", "scientific_name", "scientific name", "generic name", "active item name",
        "الاسم العلمي", "اسم علمي", "اسم المادة", "اسم المستحضر العلمي"
    ],
    "brand_name": [
        "brand_name", "brand", "trade_name", "trade name", "brand name",
        "الاسم التجاري", "اسم تجاري", "الاسم التجاري/الشركة"
    ],
    "brand_names_display": [
        "brand_names_display", "trade_names", "trade names", "brands", "brand list",
        "الأسماء التجارية", "الاسماء التجارية", "العلامات التجارية"
    ],
    "item_code": [
        "item_code", "code", "item code", "sku", "barcode", "كود", "الكود", "كود الصنف", "رقم الصنف"
    ],
    "current_stock": [
        "current_stock", "stock", "balance", "on_hand", "available", "qty", "quantity",
        "الرصيد", "الرصيد الحالي", "المتاح", "الكميه", "الكمية", "رصيد", "مخزون"
    ],
    "min_stock_level": [
        "min_stock_level", "min_stock", "minimum stock", "reorder_point", "safety_min",
        "الحد الادنى", "الحد الأدنى", "نقطة الطلب", "حد ادنى", "حد أدنى"
    ],
    "avg_daily_consumption": [
        "avg_daily_consumption", "daily_consumption", "adc", "avg_daily", "consumption_per_day",
        "متوسط الاستهلاك اليومي", "استهلاك يومي", "متوسط يومي"
    ],
    "avg_monthly_consumption": [
        "avg_monthly_consumption", "monthly_consumption", "amc", "avg_monthly", "monthly avg",
        "متوسط الاستهلاك الشهري", "استهلاك شهري", "متوسط شهري", "منصرف شهري", "متوسط المنصرف الشهري",
        "منصرفات", "منصرف"
    ],
    "dosage_form_group": [
        "dosage_form_group", "dosage_form", "dosage form", "form", "pharmaceutical form",
        "الشكل الصيدلاني", "شكل صيدلاني", "الشكل", "الهيئة الصيدلانية"
    ],
    "lead_time_days": [
        "lead_time_days", "lead_time", "supplier_lead_time", "days lead", "مدة التوريد", "زمن التوريد"
    ],
    "pending_po_qty": [
        "pending_po_qty", "pending_po", "open_po", "po_qty", "on_order",
        "امر شراء معلق", "طلب شراء معلق", "كمية قيد التوريد", "كمية طلب شراء"
    ],
    "expiry_date": [
        "expiry_date", "exp", "expiry", "expiration", "expiration_date",
        "تاريخ الصلاحية", "الصلاحية", "انتهاء الصلاحية", "exp date"
    ],
    "clinical_priority": [
        "clinical_priority", "priority", "criticality", "clinical importance",
        "الأولوية", "أهمية سريرية", "الأولوية السريرية"
    ],
    "storage_location": [
        "storage_location", "store", "location", "warehouse", "storage", "المخزن", "الموقع", "مكان التخزين"
    ],
    "branch_name": [
        "branch_name", "branch", "pharmacy", "satellite", "department", "location_name",
        "الفرع", "الصيدلية", "الوحدة", "القسم", "الموقع"
    ],
    "supplier_name": [
        "supplier_name", "supplier", "vendor", "المورد", "اسم المورد"
    ],
    "supplier_on_time_rate": [
        "supplier_on_time_rate", "on_time_rate", "supplier_score", "نسبة الالتزام", "التزام المورد"
    ],
    "unit_cost": [
        "unit_cost", "cost", "price", "unit price", "سعر الوحدة", "التكلفة", "السعر"
    ],
    "ven_class": [
        "ven_class", "ven", "VEN", "تصنيف ven", "ven classification"
    ],
    "abc_class": [
        "abc_class", "abc", "ABC", "تصنيف abc", "abc classification"
    ],
}
DEMO_DATA = [
    {
        "medicine_name": "Ceftriaxone 1g",
        "item_code": "MED-001",
        "current_stock": 180,
        "min_stock_level": 120,
        "avg_daily_consumption": 22,
        "lead_time_days": 10,
        "pending_po_qty": 0,
        "expiry_date": "2026-06-20",
        "clinical_priority": "High",
        "storage_location": "Main Pharmacy",
        "branch_name": "Main Pharmacy",
        "avg_monthly_consumption": 660,
        "unit_cost": 18,
        "supplier_name": "MediSupply",
        "supplier_on_time_rate": 0.78,
        "ven_class": "E",
    },
    {
        "medicine_name": "Insulin Regular",
        "item_code": "MED-002",
        "current_stock": 90,
        "min_stock_level": 100,
        "avg_daily_consumption": 8,
        "lead_time_days": 14,
        "pending_po_qty": 50,
        "expiry_date": "2026-05-15",
        "clinical_priority": "Critical",
        "storage_location": "Cold Chain Store",
        "branch_name": "Cold Chain Store",
        "avg_monthly_consumption": 240,
        "unit_cost": 95,
        "supplier_name": "BioCare",
        "supplier_on_time_rate": 0.62,
        "ven_class": "V",
    },
    {
        "medicine_name": "Paracetamol 500mg",
        "item_code": "MED-003",
        "current_stock": 1500,
        "min_stock_level": 500,
        "avg_daily_consumption": 40,
        "lead_time_days": 7,
        "pending_po_qty": 0,
        "expiry_date": "2026-12-30",
        "clinical_priority": "Medium",
        "storage_location": "Main Pharmacy",
        "branch_name": "ER Satellite",
        "avg_monthly_consumption": 1200,
        "unit_cost": 2.5,
        "supplier_name": "MediSupply",
        "supplier_on_time_rate": 0.78,
        "ven_class": "E",
    },
    {
        "medicine_name": "Ondansetron 8mg",
        "item_code": "MED-004",
        "current_stock": 70,
        "min_stock_level": 80,
        "avg_daily_consumption": 7,
        "lead_time_days": 12,
        "pending_po_qty": 0,
        "expiry_date": "2026-04-28",
        "clinical_priority": "High",
        "storage_location": "Oncology Satellite",
        "branch_name": "Oncology Satellite",
        "avg_monthly_consumption": 210,
        "unit_cost": 34,
        "supplier_name": "HealthPoint",
        "supplier_on_time_rate": 0.55,
        "ven_class": "V",
    },
    {
        "medicine_name": "Heparin 5000 IU",
        "item_code": "MED-005",
        "current_stock": 45,
        "min_stock_level": 60,
        "avg_daily_consumption": 6,
        "lead_time_days": 15,
        "pending_po_qty": 0,
        "expiry_date": "2026-07-10",
        "clinical_priority": "Critical",
        "storage_location": "ICU Pharmacy",
        "branch_name": "ICU Pharmacy",
        "avg_monthly_consumption": 180,
        "unit_cost": 48,
        "supplier_name": "CardioPharm",
        "supplier_on_time_rate": 0.71,
        "ven_class": "V",
    },
    {
        "medicine_name": "Amoxicillin Suspension",
        "item_code": "MED-006",
        "current_stock": 35,
        "min_stock_level": 50,
        "avg_daily_consumption": 3,
        "lead_time_days": 20,
        "pending_po_qty": 30,
        "expiry_date": "2026-04-18",
        "clinical_priority": "Medium",
        "storage_location": "Pediatrics",
        "branch_name": "Pediatrics",
        "avg_monthly_consumption": 90,
        "unit_cost": 14,
        "supplier_name": "PediaTrade",
        "supplier_on_time_rate": 0.83,
        "ven_class": "E",
    },
    {
        "medicine_name": "Salbutamol Nebules",
        "item_code": "MED-007",
        "current_stock": 260,
        "min_stock_level": 100,
        "avg_daily_consumption": 12,
        "lead_time_days": 8,
        "pending_po_qty": 40,
        "expiry_date": "2026-11-12",
        "clinical_priority": "High",
        "storage_location": "Chest Clinic",
        "branch_name": "Chest Clinic",
        "avg_monthly_consumption": 360,
        "unit_cost": 11,
        "supplier_name": "Respira",
        "supplier_on_time_rate": 0.9,
        "ven_class": "E",
    },
]
DARK_QSS = """
QMainWindow, QWidget { background: #111827; color: #e5e7eb; }
QFrame#Card { background: qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 #1f2937, stop:1 #162235); border: 1px solid #334155; border-radius: 16px; }
QGroupBox { border: 1px solid #374151; border-radius: 12px; margin-top: 10px; padding-top: 12px; font-weight: 600; }
QGroupBox::title { left: 12px; padding: 0 4px 0 4px; }
QPushButton { background: #2563eb; color: white; border: none; border-radius: 12px; padding: 8px 14px; min-height: 34px; font-weight: 700; }
QPushButton#Ghost { background: transparent; color: #bfdbfe; border: 1px solid #60a5fa; }
QPushButton#Ghost { background: transparent; color: #2563eb; border: 1px solid #93c5fd; }
QPushButton:hover { background: #1d4ed8; } QPushButton:pressed { background: #1e40af; }
QLineEdit, QComboBox, QPlainTextEdit, QTextEdit, QTableWidget, QListWidget, QDoubleSpinBox { background: #0f172a; color: #e5e7eb; border: 1px solid #334155; border-radius: 8px; }
QHeaderView::section { background: #334155; color: #f8fafc; padding: 6px; border: none; }
QToolBar { background: #0f172a; border-bottom: 1px solid #334155; spacing: 6px; padding: 4px; }
QDockWidget { font-size: 13px; }
QDockWidget::title { background: #0b1220; color: #f8fafc; padding: 10px 12px; font-weight: 800; border-bottom: 1px solid #334155; }
QListWidget { padding: 6px; outline: 0; }
QListWidget::item { padding: 8px 10px; margin: 3px 3px; border-radius: 10px; font-weight: 600; }
QListWidget::item:selected { background: #2563eb; color: white; font-weight: 700; }
QListWidget::item:hover:!selected { background: #1e293b; }
QTabWidget::pane { border: 1px solid #374151; border-radius: 12px; top: -1px; }
QTabBar::tab { background: #172033; color: #cbd5e1; padding: 10px 18px; margin-right: 4px; border-top-left-radius: 10px; border-top-right-radius: 10px; }
QTabBar::tab:selected { background: #1f2937; color: #ffffff; font-weight: 700; }
QScrollArea { border: none; }
"""
LIGHT_QSS = """
QMainWindow, QWidget { background: #f8fafc; color: #111827; }
QFrame#Card { background: qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 #ffffff, stop:1 #f3f7ff); border: 1px solid #d5dceb; border-radius: 16px; }
QGroupBox { border: 1px solid #d1d5db; border-radius: 12px; margin-top: 10px; padding-top: 12px; font-weight: 600; }
QGroupBox::title { left: 12px; padding: 0 4px 0 4px; }
QPushButton { background: #2563eb; color: white; border: none; border-radius: 12px; padding: 8px 14px; min-height: 34px; font-weight: 700; }
QPushButton:hover { background: #1d4ed8; } QPushButton:pressed { background: #1e40af; }
QLineEdit, QComboBox, QPlainTextEdit, QTextEdit, QTableWidget, QListWidget, QDoubleSpinBox { background: white; color: #111827; border: 1px solid #cbd5e1; border-radius: 8px; }
QHeaderView::section { background: #e2e8f0; color: #0f172a; padding: 6px; border: none; }
QToolBar { background: white; border-bottom: 1px solid #cbd5e1; spacing: 6px; padding: 4px; }
QDockWidget { font-size: 13px; }
QDockWidget::title { background: #eef4fb; color: #0f172a; padding: 10px 12px; font-weight: 800; border-bottom: 1px solid #dbe4ef; }
QListWidget { padding: 6px; outline: 0; }
QListWidget::item { padding: 8px 10px; margin: 3px 3px; border-radius: 10px; font-weight: 600; }
QListWidget::item:selected { background: #2563eb; color: white; font-weight: 700; }
QListWidget::item:hover:!selected { background: #eaf1ff; }
QTabWidget::pane { border: 1px solid #d1d5db; border-radius: 12px; top: -1px; }
QTabBar::tab { background: #edf2ff; color: #334155; padding: 10px 18px; margin-right: 4px; border-top-left-radius: 10px; border-top-right-radius: 10px; }
QTabBar::tab:selected { background: white; color: #0f172a; font-weight: 700; }
QScrollArea { border: none; }
"""
def t(lang, key):
    return TEXT.get(lang, TEXT["en"]).get(key, key)
TEXT["en"].update({
    "nav_forecast": "📈 Forecast Center",
    "nav_stockbalance": "⚖ Shortage vs Overstock",
    "nav_connector": "🔌 Google Sheet / DB Connector",
    "forecast_center": "Forecast Center",
    "forecast_horizon": "Forecast horizon",
    "page_size": "Page size",
    "prev_page": "◀ Previous",
    "next_page": "Next ▶",
    "connector_center": "Connector Center",
    "import_db": "🗄 Import Database / Parquet",
    "google_sync_note": "Use Google Sheets for light/medium operational sync. For large volumes use SQLite/DuckDB/Parquet and aggregate before loading the UI.",
    "stock_balance_center": "Shortage vs Overstock",
    "status_view": "Status view",
    "balance_shortage": "Shortage",
    "balance_overstock": "Overstock",
    "balance_balanced": "Balanced",
    "connector_loaded": "Connector import loaded successfully.",
    "db_prompt": "Choose source type",
    "db_sqlite": "SQLite database",
    "db_parquet": "Parquet / Feather",
    "db_csv": "CSV / TSV",
    "db_table_prompt": "Choose table or write SQL query",
    "connector_rows": "Loaded rows",
    "connector_mode": "Scalability mode",
    "forecast_gap": "Predicted gap",
    "forecast_balance": "Predicted balance",
    "top_gap_chart": "Top predicted gaps",
    "stock_mix_chart": "Shortage vs Overstock mix",
})
TEXT["ar"].update({
    "nav_forecast": "📈 مركز التوقعات",
    "nav_stockbalance": "⚖ النقص مقابل التكدس",
    "nav_connector": "🔌 ربط Google Sheet / قواعد البيانات",
    "forecast_center": "مركز التوقعات",
    "forecast_horizon": "أفق التوقع",
    "page_size": "حجم الصفحة",
    "prev_page": "◀ السابق",
    "next_page": "التالي ▶",
    "connector_center": "مركز الربط",
    "import_db": "🗄 استيراد قاعدة بيانات / Parquet",
    "google_sync_note": "Google Sheets مناسب للمزامنة اليومية الصغيرة والمتوسطة. للأحجام الكبيرة استخدم SQLite أو DuckDB أو Parquet مع تجميع البيانات قبل العرض.",
    "stock_balance_center": "النقص مقابل التكدس",
    "status_view": "عرض الحالة",
    "balance_shortage": "نقص",
    "balance_overstock": "تكدس",
    "balance_balanced": "متوازن",
    "connector_loaded": "تم تحميل بيانات الموصل بنجاح.",
    "db_prompt": "اختر نوع المصدر",
    "db_sqlite": "قاعدة SQLite",
    "db_parquet": "ملف Parquet / Feather",
    "db_csv": "ملف CSV / TSV",
    "db_table_prompt": "اختر الجدول أو اكتب استعلام SQL",
    "connector_rows": "عدد الصفوف المحملة",
    "connector_mode": "نمط التوسع",
    "forecast_gap": "الفجوة المتوقعة",
    "forecast_balance": "الرصيد المتوقع",
    "top_gap_chart": "أعلى فجوات متوقعة",
    "stock_mix_chart": "مزيج النقص والتكدس",
})
def clean_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", text)
def slug(text):
    text = clean_text(text).lower()
    text = text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    text = re.sub(r"[^a-z0-9\u0600-\u06ff]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()
def is_code_like_text(value):
    s = clean_text(value)
    if not s:
        return False
    return bool(re.fullmatch(r"[A-Za-z]{2,8}[-_]?\d{1,8}[A-Za-z0-9_-]*", s))
def scientific_candidate_score(value, medicine_name="", item_code=""):
    s = clean_text(value)
    if not s:
        return -999
    norm = normalize_search_text(s)
    if norm in {normalize_search_text(medicine_name), normalize_search_text(item_code)}:
        return -50
    score = 0
    if re.search(r"\b\d+(?:\.\d+)?\s*(mg|mcg|g|gm|ml|iu|unit|units|%)\b", s, flags=re.I):
        score += 40
    if "/" in s:
        score += 18
    if len(s) >= 6:
        score += min(len(s), 30)
    if re.search(r"[A-Za-z]", s) or _contains_arabic(s):
        score += 10
    if is_code_like_text(s):
        score -= 80
    bad_tokens = ["supplier", "vendor", "branch", "pharmacy", "store", "warehouse", "مورد", "فرع", "صيدلية", "مخزن", "location"]
    if any(tok in norm for tok in bad_tokens):
        score -= 35
    return score
def choose_best_scientific_name(preferred, source_names, analysis_name, dosage=""):
    candidates = []
    for raw in [preferred, source_names, analysis_name]:
        for part in re.split(r"\s*[|;,]+\s*", clean_text(raw)):
            part = clean_text(part)
            if not part:
                continue
            candidates.append(part)
    if dosage:
        dosage_norm = normalize_search_text(dosage)
    else:
        dosage_norm = ""
    ranked = sorted({c for c in candidates}, key=lambda c: (scientific_candidate_score(c, analysis_name, ""), len(c)), reverse=True)
    if not ranked:
        return clean_text(preferred) or clean_text(source_names) or clean_text(analysis_name)
    best = ranked[0]
    if dosage_norm and dosage_norm not in normalize_search_text(best) and not is_code_like_text(best):
        return best
    return best
def infer_priority(val):
    s = clean_text(val).lower()
    if s in {"critical", "حرج", "حرجة", "v"}:
        return "Critical"
    if s in {"high", "عالي", "عالية"}:
        return "High"
    if s in {"low", "منخفض", "منخفضة"}:
        return "Low"
    return "Medium"
def infer_branch_from_context(text):
    s = slug(text)
    for token in ["صيدلية", "pharmacy", "branch", "فرع", "store", "satellite", "clinic"]:
        if token in s:
            return clean_text(text)
    return None
def _extract_bracket_brand(raw_name: str) -> str:
    if not raw_name:
        return ""
    m = re.search(r"\[([^\]]+)\]", raw_name)
    if m:
        return clean_text(m.group(1))
    m = re.search(r"\(([^\)]*(?:brand|trade|generic|plavix|clopidogrel|amaryl|augmentin)[^\)]*)\)", raw_name, flags=re.I)
    if m:
        return clean_text(m.group(1))
    return ""
def infer_dosage_form_group(raw_name: str, explicit_form: str = "") -> str:
    text = f"{clean_text(raw_name)} {clean_text(explicit_form)}".lower()
    if not text.strip():
        return "Miscellaneous"
    tablet_tokens = ["tablet", "tablets", "tab", "tabs", "film coated", "film-coated", "capsule", "capsules", "cap", "caps", "قرص", "اقراص", "أقراص", "كبسول", "كبسولات"]
    ampoule_tokens = ["amp", "ampoule", "ampoules", "vial", "vials", "inj", "injection", "powder for injection", "امبول", "أمبول", "أمبولات", "حقن", "فيال"]
    solution_tokens = ["solution", "solutions", "syrup", "suspension", "drops", "drop", "oral solution", "infusion", "nebulizer", "nebules", "محلول", "محاليل", "شراب", "معلق", "نقط", "قطرة", "محلول وريدي"]
    misc_tokens = ["cream", "ointment", "gel", "spray", "suppository", "suppositories", "patch", "inhaler", "powder", "granules", "lozenge", "mouthwash", "misc", "متنوع", "متنوعات", "لبوس", "جل", "مرهم", "كريم", "بخاخ"]
    if any(tok in text for tok in tablet_tokens):
        return "Tablets"
    if any(tok in text for tok in ampoule_tokens):
        return "Ampoules"
    if any(tok in text for tok in solution_tokens):
        return "Solutions"
    return "Miscellaneous"
def parse_scientific_and_trade_names(raw_name: str, explicit_form: str = ""):
    raw = clean_text(raw_name)
    if not raw:
        return "Unknown Item", "", infer_dosage_form_group(raw_name, explicit_form), "Unknown Item - Miscellaneous"
    brand = _extract_bracket_brand(raw)
    working = re.sub(r"\[[^\]]+\]", " ", raw)
    working = re.sub(r"\([^\)]*\)", " ", working)
    working = re.sub(r"^[A-Za-z]{0,4}\d+[A-Za-z0-9\-]*\s*[\/\-]\s*", "", working).strip()
    if "/" in working:
        parts = [clean_text(p) for p in working.split("/") if clean_text(p)]
        if len(parts) >= 2:
            working = sorted(parts, key=lambda x: (sum(ch.isalpha() for ch in x), -sum(ch.isdigit() for ch in x)), reverse=True)[0]
    dosage_form_group = infer_dosage_form_group(raw, explicit_form)
    split_parts = [clean_text(p) for p in re.split(r"\s+-\s+", working) if clean_text(p)]
    generic = split_parts[0] if split_parts else working
    generic = re.sub(r"\d+(?:\.\d+)?\s*(?:mg|mcg|g|gram|grams|ml|iu|units?|%|gm|tablet|tab|capsule|cap|vial|amp|syrup|suspension|film coated tablet|film-coated tablet)", " ", generic, flags=re.I)
    generic = re.sub(r"(?:film coated tablet|film-coated tablet|tablet|tab|capsule|cap|vial|ampoule|amp|syrup|suspension|injection|inj|solution|oral|powder|cream|ointment|drops?)", " ", generic, flags=re.I)
    generic = re.sub(r"\s+", " ", generic).strip(' -_/')
    # If generic is only a generated code such as NAS-001, fallback to a more descriptive text.
    is_code_like = bool(re.fullmatch(r"[A-Za-z]{2,8}[\-_]?\d{1,6}[A-Za-z0-9\-_]*", generic or ""))
    if is_code_like:
        fallback = working.strip(' -_/') or raw
        if split_parts and len(split_parts) >= 2:
            fallback = (split_parts[0] + " " + split_parts[1]).strip()
        generic = fallback
    if not generic:
        generic = working.strip(' -_/') or raw
    generic_for_analysis = generic
    if normalize_search_text(dosage_form_group) in normalize_search_text(generic_for_analysis):
        analysis_name = generic_for_analysis
    else:
        analysis_name = f"{generic_for_analysis} - {dosage_form_group}"
    if not _contains_arabic(generic):
        generic = ' '.join(w.capitalize() for w in generic.split())
    return generic, brand, dosage_form_group, analysis_name
def enrich_scientific_identity(df):
    out = df.copy()
    explicit_form = out["dosage_form_group"] if "dosage_form_group" in out.columns else pd.Series([""] * len(out), index=out.index)
    parsed = [parse_scientific_and_trade_names(name, explicit_form.loc[idx] if idx in explicit_form.index else "") for idx, name in out["medicine_name"].fillna("").items()]
    parsed_df = pd.DataFrame(parsed, index=out.index, columns=["generic_name", "brand_name", "dosage_form_group", "analysis_name"])
    for col in ["generic_name", "brand_name", "dosage_form_group"]:
        if col in out.columns and out[col].notna().any():
            base = out[col].astype(str).replace("nan", "")
            parsed_df[col] = base.where(base.astype(str).str.strip() != "", parsed_df[col])
    generic_series = parsed_df["generic_name"].copy()
    generic_series = generic_series.where(generic_series.astype(str).str.strip() != "", out["medicine_name"])
    dosage_series = parsed_df["dosage_form_group"].copy()
    dosage_series = dosage_series.where(dosage_series.astype(str).str.strip() != "", "Miscellaneous")
    out["generic_name"] = generic_series.fillna(out["medicine_name"])
    out["brand_name"] = parsed_df["brand_name"].fillna("")
    out["dosage_form_group"] = dosage_series.fillna("Miscellaneous")
    out["analysis_name"] = out["generic_name"].astype(str).str.strip() + " - " + out["dosage_form_group"].astype(str).str.strip()
    out["source_medicine_names"] = out["medicine_name"].astype(str)
    out["medicine_name"] = out["analysis_name"]
    return out
def _priority_rank(val):
    return {"Low": 1, "Medium": 2, "High": 3, "Critical": 4}.get(str(val), 2)
def _pick_max_priority(series):
    if len(series) == 0:
        return "Medium"
    vals = [str(v) for v in series if pd.notna(v)]
    if not vals:
        return "Medium"
    return max(vals, key=_priority_rank)
def _pick_first_nonempty(series, default=""):
    for v in series:
        if pd.notna(v) and str(v).strip() != "":
            return clean_text(v)
    return default
def _join_unique(series, sep=", "):
    vals = []
    seen = set()
    for v in series:
        if pd.isna(v):
            continue
        s = clean_text(v)
        if not s or s in seen:
            continue
        seen.add(s)
        vals.append(s)
    return sep.join(vals)
def aggregate_by_generic_branch(df):
    if df.empty:
        return df
    working = df.copy()
    if "generic_name" not in working.columns:
        working = enrich_scientific_identity(working)
    grouped = []
    for (analysis_name, branch_name), g in working.groupby(["medicine_name", "branch_name"], dropna=False):
        brands = _join_unique(g.get("brand_name", pd.Series(dtype=str)))
        grouped.append({
            "medicine_name": analysis_name,
            "analysis_name": analysis_name,
            "generic_name": _pick_first_nonempty(g.get("generic_name", pd.Series(dtype=str)), default=analysis_name),
            "dosage_form_group": _pick_first_nonempty(g.get("dosage_form_group", pd.Series(dtype=str)), default="Miscellaneous"),
            "brand_name": _pick_first_nonempty(g.get("brand_name", pd.Series(dtype=str))),
            "brand_names_display": brands,
            "source_medicine_names": _join_unique(g.get("source_medicine_names", g["medicine_name"]), sep=" | "),
            "item_code": _join_unique(g["item_code"], sep=" | "),
            "current_stock": pd.to_numeric(g["current_stock"], errors="coerce").fillna(0).sum(),
            "min_stock_level": pd.to_numeric(g["min_stock_level"], errors="coerce").fillna(0).sum(),
            "avg_daily_consumption": pd.to_numeric(g["avg_daily_consumption"], errors="coerce").fillna(0).sum(),
            "avg_monthly_consumption": pd.to_numeric(g["avg_monthly_consumption"], errors="coerce").fillna(0).sum(),
            "lead_time_days": pd.to_numeric(g["lead_time_days"], errors="coerce").fillna(14).max(),
            "pending_po_qty": pd.to_numeric(g["pending_po_qty"], errors="coerce").fillna(0).sum(),
            "expiry_date": pd.to_datetime(g["expiry_date"], errors="coerce").min(),
            "clinical_priority": _pick_max_priority(g["clinical_priority"]),
            "storage_location": _join_unique(g["storage_location"], sep=" | "),
            "branch_name": branch_name,
            "unit_cost": pd.to_numeric(g["unit_cost"], errors="coerce").fillna(0).mean(),
            "supplier_name": _join_unique(g["supplier_name"], sep=" | "),
            "supplier_on_time_rate": pd.to_numeric(g["supplier_on_time_rate"], errors="coerce").fillna(0.75).mean(),
            "ven_class": _pick_first_nonempty(g["ven_class"], default="E") or "E",
            "abc_class": _pick_first_nonempty(g.get("abc_class", pd.Series(dtype=str)), default=""),
        })
    out = pd.DataFrame(grouped)
    for col in ALL_COLUMNS:
        if col not in out.columns:
            out[col] = None
    return out
def coalesce_duplicate_columns(df):
    out = df.copy()
    out.columns = [clean_text(c) for c in out.columns]
    seen_positions = {}
    for i, c in enumerate(list(out.columns)):
        seen_positions.setdefault(c, []).append(i)
    merged = pd.DataFrame(index=out.index)
    for col, positions in seen_positions.items():
        if len(positions) == 1:
            merged[col] = out.iloc[:, positions[0]]
        else:
            base = out.iloc[:, positions[0]].copy()
            for pos in positions[1:]:
                other = out.iloc[:, pos]
                base = base.where(
                    base.notna() & (base.astype(str).str.strip() != ""),
                    other
                )
            merged[col] = base
    return merged
def normalize_dataframe(df):
    out = coalesce_duplicate_columns(df)
    numeric_cols = [
        "current_stock", "min_stock_level", "avg_daily_consumption",
        "avg_monthly_consumption", "lead_time_days", "pending_po_qty",
        "unit_cost", "supplier_on_time_rate"
    ]
    for col in ALL_COLUMNS:
        if col not in out.columns:
            out[col] = None
    for col in numeric_cols:
        out[col] = pd.to_numeric(out[col], errors="coerce")
    if out["avg_daily_consumption"].isna().all() and out["avg_monthly_consumption"].notna().any():
        out["avg_daily_consumption"] = out["avg_monthly_consumption"] / 30.0
    out["avg_daily_consumption"] = out["avg_daily_consumption"].fillna(0)
    out["avg_monthly_consumption"] = out["avg_monthly_consumption"].fillna(out["avg_daily_consumption"] * 30)
    out["lead_time_days"] = out["lead_time_days"].fillna(14)
    out["pending_po_qty"] = out["pending_po_qty"].fillna(0)
    out["current_stock"] = out["current_stock"].fillna(out["avg_monthly_consumption"]).fillna(0)
    out["min_stock_level"] = out["min_stock_level"].fillna((out["avg_monthly_consumption"] * 0.5).round()).fillna(0)
    out["unit_cost"] = out["unit_cost"].fillna(0)
    out["supplier_on_time_rate"] = out["supplier_on_time_rate"].fillna(0.75)
    out["expiry_date"] = pd.to_datetime(out["expiry_date"], errors="coerce")
    out["medicine_name"] = out["medicine_name"].fillna("Unknown Item").map(clean_text)
    # Heuristic recovery for imported scientific names when the workbook uses code columns like NAS-002.
    text_like_cols = [c for c in out.columns if c not in numeric_cols + ["expiry_date"]]
    recovered_generic = []
    current_generic = out.get("generic_name", pd.Series([None] * len(out), index=out.index))
    for idx, row in out.iterrows():
        current_val = clean_text(current_generic.loc[idx] if idx in current_generic.index else "")
        if current_val and not is_code_like_text(current_val):
            recovered_generic.append(current_val)
            continue
        best_val = ""
        best_score = -999
        for col in text_like_cols:
            val = clean_text(row.get(col, ""))
            score = scientific_candidate_score(val, row.get("medicine_name", ""), row.get("item_code", ""))
            if score > best_score:
                best_score = score
                best_val = val
        recovered_generic.append(best_val or current_val or clean_text(row.get("medicine_name", "")))
    out["generic_name"] = pd.Series(recovered_generic, index=out.index)
    out = enrich_scientific_identity(out)
    missing_code_mask = out["item_code"].isna() | (out["item_code"].astype(str).str.strip() == "")
    out.loc[missing_code_mask, "item_code"] = [
        f"AUTO-{i+1:04d}" for i in range(missing_code_mask.sum())
    ]
    out["item_code"] = out["item_code"].astype(str).map(clean_text)
    out["clinical_priority"] = out["clinical_priority"].map(infer_priority)
    out["storage_location"] = out["storage_location"].fillna(out["branch_name"]).fillna("Main Store").map(clean_text)
    out["branch_name"] = out["branch_name"].fillna(out["storage_location"]).fillna("Main Branch").map(clean_text)
    out["supplier_name"] = out["supplier_name"].fillna("Unknown Supplier").map(clean_text)
    out["ven_class"] = out["ven_class"].fillna("E").astype(str).str.upper().replace({"": "E"})
    out["abc_class"] = out["abc_class"].fillna("").astype(str).str.upper()
    out = out[out["generic_name"].astype(str).str.strip() != ""].copy()
    out = out.drop_duplicates(subset=["source_medicine_names", "item_code", "branch_name"], keep="first")
    return out[ALL_COLUMNS]
def compute_abc(df):
    out = df.copy()
    if out.empty:
        return out
    usage_value = (out["avg_monthly_consumption"].fillna(0) * out["unit_cost"].fillna(0)).astype(float)
    total = usage_value.sum()
    if total <= 0:
        out["abc_class"] = out["abc_class"].replace("", "C")
        return out
    out["_annual_value"] = usage_value
    out = out.sort_values("_annual_value", ascending=False).copy()
    out["_cum_pct"] = out["_annual_value"].cumsum() / total
    def classify(p):
        if p <= 0.7:
            return "A"
        if p <= 0.9:
            return "B"
        return "C"
    out["abc_class"] = out["_cum_pct"].apply(classify)
    return out.drop(columns=["_annual_value", "_cum_pct"])
def compute_priority_multiplier(priority):
    return {"Critical": 1.8, "High": 1.5, "Medium": 1.2, "Low": 1.0}.get(priority, 1.2)
def build_reason_explanation_from_row(row):
    reasons = []
    if row.get("current_stock", 0) < row.get("dynamic_min_stock", 0):
        reasons.append("current stock is below dynamic minimum")
    if row.get("days_of_stock_left", math.inf) != math.inf and row.get("days_of_stock_left", math.inf) < row.get("lead_time_days", 0):
        reasons.append("stock coverage is below supplier lead time")
    if row.get("pending_po_qty", 0) <= 0:
        reasons.append("there is no pending purchase order")
    if row.get("expiry_risk") in {"High", "Critical", "Expired"}:
        reasons.append("expiry timeline needs review")
    if row.get("clinical_priority") in {"High", "Critical"}:
        reasons.append(f"clinical priority is {str(row.get('clinical_priority')).lower()}")
    return "; ".join(reasons) if reasons else "no major rule-based risk signals detected"
def compute_metrics(df):
    out = normalize_dataframe(df)
    if out.empty:
        return out
    out = aggregate_by_generic_branch(out)
    out = compute_abc(out)
    today = pd.Timestamp(date.today())
    priority_mult = out["clinical_priority"].map({"Critical": 1.8, "High": 1.5, "Medium": 1.2, "Low": 1.0}).fillna(1.2)
    daily = out["avg_daily_consumption"].astype(float).fillna(0)
    monthly = out["avg_monthly_consumption"].astype(float)
    monthly = monthly.fillna(daily * 30)
    monthly = monthly.mask(monthly == 0, daily * 30)
    current = out["current_stock"].astype(float).fillna(0)
    lead = out["lead_time_days"].astype(float).fillna(0)
    pending = out["pending_po_qty"].astype(float).fillna(0)
    min_stock = out["min_stock_level"].astype(float).fillna(0)
    available = current + pending
    with np.errstate(divide='ignore', invalid='ignore'):
        coverage = np.where(daily > 0, np.round(current / daily, 1), np.inf)
        coverage_months = np.where(monthly > 0, np.round(available / monthly, 2), np.inf)
    out["days_of_stock_left"] = coverage
    out["coverage_months"] = coverage_months
    out["days_to_expiry"] = (out["expiry_date"] - today).dt.days
    dyn = np.maximum(np.ceil(daily * lead * priority_mult), min_stock).astype(int)
    out["dynamic_min_stock"] = dyn
    out["expected_stock_at_lead_time"] = np.round(available - (daily * lead), 1)
    target_days = np.maximum(lead * priority_mult, 14)
    reorder = np.ceil((daily * target_days) + dyn - available)
    out["recommended_reorder_qty"] = np.maximum(reorder, 0).astype(int)
    out["predicted_stock_30d"] = np.round(available - (daily * 30), 1)
    out["predicted_stock_60d"] = np.round(available - (daily * 60), 1)
    out["predicted_stock_90d"] = np.round(available - (daily * 90), 1)
    score = np.zeros(len(out), dtype=int)
    score += (out["expected_stock_at_lead_time"] < 0).astype(int) * 3
    score += (current < dyn).astype(int) * 2
    score += (np.isfinite(coverage) & (coverage < lead)).astype(int) * 3
    score += (pending <= 0).astype(int) * 1
    score += (out["clinical_priority"] == "Critical").astype(int) * 2
    score += (out["clinical_priority"] == "High").astype(int) * 1
    out["shortage_score"] = score
    out["shortage_risk"] = np.select([score >= 7, score >= 5, score >= 3], ["Critical", "High", "Moderate"], default="Low")
    overstock_flags = (coverage_months >= 3) | (available >= (dyn * 2.2))
    severe_overstock = (coverage_months >= 6) | (available >= (dyn * 3.5))
    out["overstock_risk"] = np.select([severe_overstock, overstock_flags], ["High", "Moderate"], default="Low")
    out["stock_status"] = np.select(
        [out["shortage_risk"].isin(["Critical", "High"]), out["overstock_risk"].isin(["High", "Moderate"])],
        ["Shortage", "Overstock"],
        default="Balanced"
    )
    dte = out["days_to_expiry"]
    out["expiry_risk"] = np.select([dte.isna(), dte < 0, dte <= 30, dte <= 90, dte <= 180], ["Unknown", "Expired", "Critical", "High", "Moderate"], default="Low")
    usable = daily * dte.fillna(0)
    expected_unused = np.where(dte.isna() | (dte <= 0), current, np.maximum(np.ceil(current - usable), 0))
    out["expected_unused_before_expiry"] = expected_unused.astype(int)
    below_dyn = current < dyn
    below_lead = np.isfinite(coverage) & (coverage < lead)
    no_po = pending <= 0
    expiry_review = out["expiry_risk"].isin(["High", "Critical", "Expired"])
    high_clinical = out["clinical_priority"].isin(["High", "Critical"])
    reason = np.where(below_dyn, "current stock is below dynamic minimum; ", "")
    reason = np.where(below_lead, reason + "stock coverage is below supplier lead time; ", reason)
    reason = np.where(no_po, reason + "there is no pending purchase order; ", reason)
    reason = np.where(expiry_review, reason + "expiry timeline needs review; ", reason)
    reason = np.where(out["overstock_risk"].isin(["Moderate", "High"]), reason + "stock coverage is above expected demand; ", reason)
    reason = np.where(high_clinical, reason + "clinical priority is " + out["clinical_priority"].astype(str).str.lower() + "; ", reason)
    out["reason_explanation"] = pd.Series(reason, index=out.index).str.rstrip("; ").replace("", "no major rule-based risk signals detected")
    out["escalation_level"] = np.select([
        (out["shortage_risk"] == "Critical") & out["clinical_priority"].isin(["Critical", "High"]),
        out["shortage_risk"].isin(["High", "Critical"]) | out["expiry_risk"].isin(["High", "Critical", "Expired"]),
        out["overstock_risk"] == "High"
    ], ["Red", "Amber", "Amber"], default="Green")
    out["recommended_action"] = np.select([
        out["escalation_level"] == "Red",
        out["shortage_risk"] == "High",
        out["overstock_risk"].isin(["Moderate", "High"]),
        out["expiry_risk"].isin(["High", "Critical"]) & (out["expected_unused_before_expiry"] > 0)
    ], [
        "Escalate to Pharmacy Director, verify physical stock, check redistribution, and prepare urgent draft.",
        "Review within 24 hours, confirm supplier status, and prepare purchase request draft.",
        "Review redistribution, supplier pacing, and purchasing freeze risk for overstock.",
        "Review FEFO issue plan or redistribution before expiry."
    ], default="Continue routine monitoring.")
    # default purchasing horizon = 1 month
    out["selected_horizon_months"] = 1
    out["recommended_reorder_qty_horizon"] = np.maximum(np.ceil((monthly * out["selected_horizon_months"]) + dyn - available), 0).astype(int)
    return out
def draft_purchase_request(row):
    horizon = int(float(row.get("selected_horizon_months", 1) or 1))
    reorder_h = int(float(row.get("recommended_reorder_qty_horizon", row.get("recommended_reorder_qty", 0)) or 0))
    return f"""Purchase Request Draft
Scientific name: {row['generic_name']}
Trade names: {row.get('brand_names_display', '')}
Item code(s): {row['item_code']}
Branch: {row['branch_name']}
Current stock: {row['current_stock']}
Dynamic minimum stock: {row['dynamic_min_stock']}
Average daily consumption: {row['avg_daily_consumption']}
Average monthly consumption: {row['avg_monthly_consumption']}
Lead time: {row['lead_time_days']} days
Pending PO quantity: {row['pending_po_qty']}
Shortage risk: {row['shortage_risk']}
Overstock risk: {row.get('overstock_risk', 'Low')}
Predicted stock after 30d: {row.get('predicted_stock_30d', '')}
Predicted stock after 60d: {row.get('predicted_stock_60d', '')}
Predicted stock after 90d: {row.get('predicted_stock_90d', '')}
Consumption horizon: {horizon} month(s)
Recommended reorder quantity: {reorder_h}
Clinical priority: {row['clinical_priority']}
Justification:
The item shows {str(row['shortage_risk']).lower()} shortage risk because {row['reason_explanation']}. The proposed quantity is aligned to a {horizon}-month consumption horizon while preserving a safety buffer suited to the item's operational and clinical importance.
Governance note:
This draft is advisory only. Final review and approval must be completed by authorized pharmacy and procurement staff according to policy and public-sector procurement controls.
"""
def detect_header_and_sheet(path):
    log = []
    source = clean_text(path)
    lower_path = source.lower()
    if lower_path.startswith(("http://", "https://")):
        csv_url = google_sheet_to_csv_url(source)
        df = pd.read_csv(csv_url)
        log.append("URL source detected; CSV parser used.")
        if csv_url != source:
            log.append("Google Sheet URL normalized to CSV export endpoint.")
            return df, "Google Sheet", log
        return df, "URL CSV", log
    if lower_path.endswith('.csv'):
        df = pd.read_csv(source)
        log.append("CSV file detected; standard parser used.")
        return df, "CSV", log
    # Fast path for xlsx/xlsm using openpyxl read-only scanning
    if HAS_OPENPYXL and lower_path.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(source, read_only=True, data_only=True)
        best_sheet = None
        best_header = 0
        best_score = -1
        for ws in wb.worksheets:
            for header_row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), start=1):
                header_vals = [clean_text(v) for v in row]
                score = 0
                for h in header_vals:
                    hs = slug(h)
                    if not hs:
                        continue
                    for aliases in COLUMN_ALIASES.values():
                        if any(slug(a) == hs or slug(a) in hs or hs in slug(a) for a in aliases):
                            score += 1
                            break
                joined = " | ".join(header_vals)
                sj = slug(joined)
                if any(x in sj for x in ["الصنف", "medicine", "drug", "item", "منصرف"]):
                    score += 2
                if score > best_score:
                    best_score = score
                    best_sheet = ws.title
                    best_header = header_row_idx - 1
        if best_sheet is None:
            raise ValueError("Could not detect a usable sheet/header.")
        log.append(f"Best sheet: {best_sheet}")
        log.append(f"Detected header row: {best_header + 1}")
        log.append(f"Header match score: {best_score}")
        df = pd.read_excel(source, sheet_name=best_sheet, header=best_header)
        return df, best_sheet, log
    # Fallback generic parser
    xls = pd.ExcelFile(source)
    best = None
    best_score = -1
    best_sheet = None
    best_header = 0
    for sheet in xls.sheet_names:
        try:
            raw = pd.read_excel(source, sheet_name=sheet, header=None, nrows=12)
        except Exception as exc:
            log.append(f"Could not read sheet {sheet}: {exc}")
            continue
        for header_row in range(len(raw)):
            header_vals = [clean_text(v) for v in raw.iloc[header_row].tolist()]
            score = 0
            for h in header_vals:
                hs = slug(h)
                if not hs:
                    continue
                for aliases in COLUMN_ALIASES.values():
                    if any(slug(a) == hs or slug(a) in hs or hs in slug(a) for a in aliases):
                        score += 1
                        break
            if score > best_score:
                best_score = score
                best_sheet = sheet
                best_header = header_row
                best = raw
    if best_sheet is None:
        raise ValueError("Could not detect a usable sheet/header.")
    log.append(f"Best sheet: {best_sheet}")
    log.append(f"Detected header row: {best_header + 1}")
    log.append(f"Header match score: {best_score}")
    df = pd.read_excel(source, sheet_name=best_sheet, header=best_header)
    return df, best_sheet, log


def map_columns_generic(df):
    original_columns = list(df.columns)
    mapped = {}
    processing_log = []
    cols_clean = {col: slug(col) for col in df.columns}

    for target, aliases in COLUMN_ALIASES.items():
        for col, col_slug in cols_clean.items():
            if any(slug(a) == col_slug or slug(a) in col_slug or col_slug in slug(a) for a in aliases if col_slug):
                mapped[col] = target
                break

    out = df.rename(columns=mapped).copy()
    out = coalesce_duplicate_columns(out)
    processing_log.append(f"Mapped columns: {mapped}")

    out = out.dropna(axis=1, how="all")
    out = out.dropna(axis=0, how="all")

    if "medicine_name" not in out.columns:
        candidates = [c for c in out.columns if not pd.api.types.is_numeric_dtype(out[c])]
        if candidates:
            out = out.rename(columns={candidates[0]: "medicine_name"})
            processing_log.append(f"Fallback medicine_name used from column: {candidates[0]}")

    numeric_candidates = [
        c for c in out.columns
        if c not in mapped.values() and pd.to_numeric(out[c], errors="coerce").notna().sum() >= max(3, len(out)//5)
    ]
    branch_like = [c for c in numeric_candidates if infer_branch_from_context(c) or not str(c).lower().startswith("unnamed")]
    if (
        "branch_name" not in out.columns
        and "current_stock" not in out.columns
        and "avg_monthly_consumption" not in out.columns
        and len(branch_like) >= 2
        and "medicine_name" in out.columns
    ):
        melted = out.melt(id_vars=["medicine_name"], value_vars=branch_like, var_name="branch_name", value_name="avg_monthly_consumption")
        melted = melted[melted["avg_monthly_consumption"].notna()].copy()
        melted["branch_name"] = melted["branch_name"].map(clean_text)
        processing_log.append("Detected wide branch sheet and unpivoted to long format using numeric branch columns.")
        out = melted

    if "clinical_priority" not in out.columns:
        out["clinical_priority"] = "Medium"
    if "lead_time_days" not in out.columns:
        out["lead_time_days"] = 14
    if "pending_po_qty" not in out.columns:
        out["pending_po_qty"] = 0
    if "storage_location" not in out.columns:
        out["storage_location"] = out["branch_name"] if "branch_name" in out.columns else "Main Store"

    out = coalesce_duplicate_columns(out)

    if "current_stock" not in out.columns and "avg_monthly_consumption" in out.columns:
        out["current_stock"] = pd.to_numeric(out["avg_monthly_consumption"], errors="coerce")
        processing_log.append("current_stock inferred from avg_monthly_consumption for prototype use.")
    if "min_stock_level" not in out.columns and "avg_monthly_consumption" in out.columns:
        out["min_stock_level"] = pd.to_numeric(out["avg_monthly_consumption"], errors="coerce") * 0.5
        processing_log.append("min_stock_level inferred from avg_monthly_consumption * 0.5.")

    if "supplier_name" not in out.columns:
        out["supplier_name"] = "Unknown Supplier"
    if "supplier_on_time_rate" not in out.columns:
        out["supplier_on_time_rate"] = 0.75
    if "unit_cost" not in out.columns:
        out["unit_cost"] = 0
    if "ven_class" not in out.columns:
        out["ven_class"] = "E"

    to_drop = [c for c in out.columns if str(c).lower().startswith("unnamed")]
    if to_drop:
        out = out.drop(columns=to_drop, errors="ignore")

    return out, processing_log, original_columns



# ----------------------------------------------------------------------
# SPECIAL IMPORTER: Egyptian Ministry / Nasser consumption workbook
# Handles files with multi-row Arabic headers and branch pairs:
# Branch name row -> "منصرف" / "رصيد" sub-columns.
# It converts the wide workbook into PharmaGuard's normal long format
# without removing any existing import features.
# ----------------------------------------------------------------------
def _pg_is_blank(value):
    try:
        return pd.isna(value)
    except Exception:
        return value is None


def _pg_to_number(value):
    if value is None or _pg_is_blank(value):
        return np.nan
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    s = clean_text(value)
    if not s or s in {"-", "--", "—", "_"}:
        return np.nan
    s = s.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    s = s.replace(",", "").replace("٬", "").replace("٫", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return np.nan
    try:
        return float(m.group(0))
    except Exception:
        return np.nan



# ----------------------------------------------------------------------
# SAFE FILE PATH RESOLVER
# Fixes common Windows beginner issues:
# - user selects/moves file and old path no longer exists
# - Arabic file names / long paths sometimes fail in Excel engines
# - file is in the same PharmaGuard folder but Windows returns another path
# ----------------------------------------------------------------------
def _pg_resolve_existing_input_path(path):
    raw = str(path or "").strip().strip('"').strip("'")
    if not raw:
        raise FileNotFoundError("No file path was selected.")

    if raw.lower().startswith(("http://", "https://")):
        return raw

    p = Path(raw)
    if p.exists():
        return str(p)

    name = p.name
    candidates = []
    try:
        candidates.append(Path.cwd() / name)
    except Exception:
        pass
    try:
        candidates.append(Path(__file__).resolve().parent / name)
    except Exception:
        pass
    try:
        home = Path.home()
        candidates.extend([
            home / "Desktop" / "PharmaGuard" / name,
            home / "Desktop" / name,
            home / "Downloads" / name,
            home / "Documents" / name,
        ])
    except Exception:
        pass

    for c in candidates:
        try:
            if c.exists():
                return str(c)
        except Exception:
            continue

    search_dirs = []
    try:
        search_dirs.extend([
            Path.cwd(),
            Path(__file__).resolve().parent,
            Path.home() / "Desktop" / "PharmaGuard",
            Path.home() / "Desktop",
            Path.home() / "Downloads",
            Path.home() / "Documents",
        ])
    except Exception:
        pass
    wanted = name.casefold()
    for folder in search_dirs:
        try:
            if not folder.exists() or not folder.is_dir():
                continue
            for child in folder.iterdir():
                if child.name.casefold() == wanted:
                    return str(child)
        except Exception:
            continue

    raise FileNotFoundError(
        "PharmaGuard could not find the selected Excel/CSV file.\n\n"
        f"Selected path was:\n{raw}\n\n"
        "Easy fix:\n"
        "1) Put the Excel file inside the same PharmaGuard folder as the .py file.\n"
        "2) Rename it to a simple English name such as nasser.xlsx.\n"
        "3) Open Import again and select that file from the PharmaGuard folder."
    )


def _pg_make_safe_import_copy(path):
    """Return a local temporary ASCII path for pandas/openpyxl reads."""
    import shutil
    import tempfile

    resolved = _pg_resolve_existing_input_path(path)
    if resolved.lower().startswith(("http://", "https://")):
        return resolved, None

    p = Path(resolved)
    suffix = p.suffix or ".xlsx"
    tmp_dir = Path(tempfile.mkdtemp(prefix="pharmaguard_import_"))
    safe_path = tmp_dir / ("source" + suffix.lower())
    shutil.copy2(str(p), str(safe_path))
    return str(safe_path), tmp_dir

def _pg_sheet_form_group(sheet_name):
    s = slug(sheet_name)
    if "اقراص" in s or "قرص" in s:
        return "Tablets"
    if "امبول" in s or "امبولات" in s or "حقن" in s or "فيال" in s:
        return "Ampoules"
    if "محاليل" in s or "محلول" in s:
        return "Solutions"
    return "Miscellaneous"


def _pg_find_header_rows(raw):
    max_scan = min(14, len(raw))
    best_meta = 0
    best_sub = 0
    best_score = -1
    for i in range(max_scan):
        vals = [clean_text(x) for x in raw.iloc[i].tolist()]
        joined = slug(" ".join(vals))
        score = 0
        if "الاسم كما" in joined:
            score += 4
        if "الاسم العلمي" in joined or "الاسم العلمى" in joined:
            score += 3
        if "الاستهلاك" in joined:
            score += 2
        if "رصيد شهر" in joined or "رصيد نوفمبر" in joined:
            score += 2
        if "صيدليه" in joined or "صيدلية" in joined:
            score += 2
        sub_candidate = i + 1
        for j in range(i, min(i + 6, max_scan)):
            row_joined = slug(" ".join([clean_text(x) for x in raw.iloc[j].tolist()]))
            if "منصرف" in row_joined and "رصيد" in row_joined:
                score += 4
                sub_candidate = j
                break
        if score > best_score:
            best_score = score
            best_meta = i
            best_sub = sub_candidate
    return best_meta, best_sub, best_score


def _pg_find_col(raw, header_start, sub_row, keywords):
    for c in range(raw.shape[1]):
        parts = []
        for r in range(header_start, min(sub_row + 1, len(raw))):
            parts.append(clean_text(raw.iat[r, c]))
        cell_text = slug(" ".join(parts))
        if any(k in cell_text for k in keywords):
            return c
    return None


def _pg_branch_name_for_col(raw, meta_row, sub_row, col):
    for c in [col, col - 1, col - 2]:
        if c < 0 or c >= raw.shape[1]:
            continue
        vals = []
        for r in range(meta_row, min(sub_row, len(raw))):
            val = clean_text(raw.iat[r, c])
            if val:
                vals.append(val)
        for val in reversed(vals):
            sl = slug(val)
            if any(token in sl for token in ["صيدليه", "صيدلية", "خارجي", "اورام", "داخلي", "داخلى", "عمليات", "التصلب", "رعايه", "رعاية", "القلب", "السادس"]):
                return val
        if vals:
            return vals[-1]
    return f"Branch column {col + 1}"


def _pg_build_branch_pairs(raw, meta_row, sub_row, after_col):
    pairs = []
    last_branch = ""
    c = max(int(after_col) + 1, 0)
    while c < raw.shape[1]:
        sub = slug(raw.iat[sub_row, c]) if sub_row < len(raw) else ""
        if "منصرف" in sub:
            branch = _pg_branch_name_for_col(raw, meta_row, sub_row, c) or last_branch or f"Branch {c + 1}"
            last_branch = branch
            stock_col = None
            if c + 1 < raw.shape[1]:
                nxt = slug(raw.iat[sub_row, c + 1]) if sub_row < len(raw) else ""
                if "رصيد" in nxt:
                    stock_col = c + 1
            pairs.append((clean_text(branch), c, stock_col))
            c += 2 if stock_col is not None else 1
            continue
        if "رصيد" in sub and last_branch:
            pairs.append((clean_text(last_branch), None, c))
        c += 1
    cleaned = []
    seen = set()
    for branch, dispense_col, stock_col in pairs:
        if not branch:
            continue
        key = (branch, dispense_col, stock_col)
        if key in seen:
            continue
        seen.add(key)
        cleaned.append((branch, dispense_col, stock_col))
    return cleaned


def _pg_is_ministry_consumption_workbook(path):
    src = clean_text(path)
    if src.lower().startswith(("http://", "https://")) or src.lower().endswith(".csv"):
        return False
    if not src.lower().endswith((".xlsx", ".xlsm", ".xls")):
        return False
    try:
        xls = pd.ExcelFile(src)
        for sheet in xls.sheet_names[:6]:
            raw = pd.read_excel(src, sheet_name=sheet, header=None, nrows=14, dtype=object)
            meta_row, sub_row, score = _pg_find_header_rows(raw)
            joined = slug(" ".join(raw.fillna("").astype(str).head(14).values.ravel().tolist()))
            if score >= 8 and "منصرف" in joined and "رصيد" in joined and ("الاسم العلمي" in joined or "الاسم العلمى" in joined):
                return True
    except Exception:
        return False
    return False


def import_ministry_consumption_workbook(path):
    xls = pd.ExcelFile(path)
    rows = []
    report = ["Special importer used: Ministry/Nasser multi-row Arabic consumption workbook."]
    skipped = []
    for sheet in xls.sheet_names:
        sheet_slug = slug(sheet)
        if sheet.lower().startswith("copy") or "copy of" in sheet.lower() or sheet_slug.startswith("copy"):
            skipped.append(sheet)
            continue
        try:
            raw = pd.read_excel(path, sheet_name=sheet, header=None, dtype=object)
        except Exception as exc:
            report.append(f"Could not read sheet {sheet}: {exc}")
            continue
        if raw.dropna(how="all").empty:
            continue
        meta_row, sub_row, score = _pg_find_header_rows(raw)
        if score < 6:
            report.append(f"Skipped sheet {sheet}: header pattern was not recognized.")
            continue

        source_col = _pg_find_col(raw, meta_row, sub_row, ["الاسم كما", "مدون علي المنظومه", "مدون على المنظومه"])
        generic_col = _pg_find_col(raw, meta_row, sub_row, ["الاسم العلمي", "الاسم العلمى"])
        brand_col = _pg_find_col(raw, meta_row, sub_row, ["الاسم التجاري", "الاسم التجارى"])
        unit_col = _pg_find_col(raw, meta_row, sub_row, ["الوحده"])
        monthly_col = _pg_find_col(raw, meta_row, sub_row, ["معدل الاستهلاك", "متوسط الاستهلاك"])
        total_stock_col = _pg_find_col(raw, meta_row, sub_row, ["رصيد شهر", "رصيد نوفمبر"])
        serial_col = _pg_find_col(raw, meta_row, sub_row, ["م"]) or 0

        after_col = total_stock_col if total_stock_col is not None else (monthly_col if monthly_col is not None else 6)
        branch_pairs = _pg_build_branch_pairs(raw, meta_row, sub_row, after_col)
        form_group = _pg_sheet_form_group(sheet)

        report.append(
            f"Sheet '{sheet}': header row {meta_row + 1}, subheader row {sub_row + 1}, "
            f"branches detected {len(branch_pairs)}."
        )

        last_source = ""
        last_generic = ""
        last_monthly = np.nan
        start_row = sub_row + 1

        for ri in range(start_row, len(raw)):
            row = raw.iloc[ri]
            source_name = clean_text(row.iloc[source_col]) if source_col is not None and source_col < len(row) else ""
            generic_name = clean_text(row.iloc[generic_col]) if generic_col is not None and generic_col < len(row) else ""
            brand_name = clean_text(row.iloc[brand_col]) if brand_col is not None and brand_col < len(row) else ""
            unit_name = clean_text(row.iloc[unit_col]) if unit_col is not None and unit_col < len(row) else ""
            monthly_value = _pg_to_number(row.iloc[monthly_col]) if monthly_col is not None and monthly_col < len(row) else np.nan
            total_stock = _pg_to_number(row.iloc[total_stock_col]) if total_stock_col is not None and total_stock_col < len(row) else np.nan

            if source_name:
                last_source = source_name
            else:
                source_name = last_source
            if generic_name:
                last_generic = generic_name
            else:
                generic_name = last_generic
            if not np.isnan(monthly_value):
                last_monthly = monthly_value
            else:
                monthly_value = last_monthly

            if not source_name and not generic_name and not brand_name and np.isnan(total_stock):
                continue

            base_name = source_name or generic_name or brand_name
            if not base_name:
                continue

            if unit_name and unit_name not in str(base_name):
                source_display = f"{base_name} | Unit: {unit_name}"
            else:
                source_display = base_name

            if not np.isnan(total_stock) or not np.isnan(monthly_value):
                monthly_for_main = 0 if np.isnan(monthly_value) else float(monthly_value)
                stock_for_main = 0 if np.isnan(total_stock) else float(total_stock)
                rows.append({
                    "medicine_name": source_display,
                    "generic_name": generic_name or base_name,
                    "brand_name": brand_name,
                    "dosage_form_group": form_group,
                    "source_medicine_names": source_display,
                    "item_code": f"{sheet}-{ri + 1}-MAIN",
                    "current_stock": stock_for_main,
                    "min_stock_level": max(monthly_for_main * 0.5, 0),
                    "avg_daily_consumption": monthly_for_main / 30.0 if monthly_for_main else 0,
                    "avg_monthly_consumption": monthly_for_main,
                    "lead_time_days": 14,
                    "pending_po_qty": 0,
                    "expiry_date": None,
                    "clinical_priority": "Medium",
                    "storage_location": "رصيد شهر نوفمبر 2023",
                    "branch_name": "رصيد شهر نوفمبر 2023",
                    "unit_cost": 0,
                    "supplier_name": "Unknown Supplier",
                    "supplier_on_time_rate": 0.75,
                    "ven_class": "E",
                })

            for branch_name, dispense_col, stock_col in branch_pairs:
                dispensed = _pg_to_number(row.iloc[dispense_col]) if dispense_col is not None and dispense_col < len(row) else np.nan
                branch_stock = _pg_to_number(row.iloc[stock_col]) if stock_col is not None and stock_col < len(row) else np.nan
                if np.isnan(dispensed) and np.isnan(branch_stock):
                    continue
                monthly_for_branch = float(dispensed) if not np.isnan(dispensed) else (0 if np.isnan(monthly_value) else float(monthly_value))
                stock_for_branch = float(branch_stock) if not np.isnan(branch_stock) else 0
                rows.append({
                    "medicine_name": source_display,
                    "generic_name": generic_name or base_name,
                    "brand_name": brand_name,
                    "dosage_form_group": form_group,
                    "source_medicine_names": source_display,
                    "item_code": f"{sheet}-{ri + 1}-{dispense_col if dispense_col is not None else 'S'}",
                    "current_stock": stock_for_branch,
                    "min_stock_level": max(monthly_for_branch * 0.5, 0),
                    "avg_daily_consumption": monthly_for_branch / 30.0 if monthly_for_branch else 0,
                    "avg_monthly_consumption": monthly_for_branch,
                    "lead_time_days": 14,
                    "pending_po_qty": 0,
                    "expiry_date": None,
                    "clinical_priority": "Medium",
                    "storage_location": branch_name,
                    "branch_name": branch_name,
                    "unit_cost": 0,
                    "supplier_name": "Unknown Supplier",
                    "supplier_on_time_rate": 0.75,
                    "ven_class": "E",
                })

    if skipped:
        report.append(f"Skipped duplicated/copy sheets: {', '.join(skipped)}")
    if not rows:
        raise ValueError("No usable pharmacy consumption rows were extracted from this workbook.")

    parsed = pd.DataFrame(rows)
    report.append(f"Rows extracted before PharmaGuard normalization: {len(parsed):,}")
    report.append(f"Branches/locations detected: {parsed['branch_name'].nunique():,}")
    report.append("Note: this source workbook has no expiry/batch columns, so expiry risk will appear as Unknown until a batch-expiry file is imported.")
    normalized = normalize_dataframe(parsed)
    report.append(f"Rows after PharmaGuard normalization: {len(normalized):,}")
    return normalized, "All recognized sheets", "\n".join(report), parsed.head(25)

def import_any_file(path):
    # V9 robust import flow:
    # 1) Resolve/copy the selected file to a safe temporary path.
    # 2) Try the special Egyptian hospital workbook parser when the shape matches.
    # 3) Try the generic Excel/CSV parser.
    # 4) If generic parsing fails on Excel, try the special parser as a fallback.
    # No existing feature is removed; this only makes import safer and easier.
    safe_path, tmp_dir = _pg_make_safe_import_copy(path)
    try:
        report_notes = []
        is_excel = str(safe_path).lower().endswith((".xlsx", ".xlsm", ".xls"))

        try:
            if is_excel and _pg_is_ministry_consumption_workbook(safe_path):
                df, sheet_used, report, preview = import_ministry_consumption_workbook(safe_path)
                report = "Importer selected: SPECIAL GOVERNMENT EXCEL / Arabic branch مصرف-رصيد layout.\n" + str(report)
                return df, sheet_used, report, preview
        except Exception as special_detect_exc:
            report_notes.append(f"Special importer detection warning: {special_detect_exc}")

        generic_exc = None
        try:
            report = []
            if report_notes:
                report.extend(report_notes)
            df0, sheet_used, log1 = detect_header_and_sheet(safe_path)
            report.extend(log1)
            mapped_df, log2, original_columns = map_columns_generic(df0)
            report.extend(log2)
            report.append(f"Original columns: {original_columns}")
            report.append(f"Final columns before normalization: {list(mapped_df.columns)}")
            normalized = normalize_dataframe(mapped_df)
            report.append(f"Rows after normalization: {len(normalized)}")
            return normalized, sheet_used, "\n".join(report), df0.head(50)
        except Exception as exc:
            generic_exc = exc

        if is_excel:
            try:
                df, sheet_used, report, preview = import_ministry_consumption_workbook(safe_path)
                report = (
                    "Importer selected: SPECIAL GOVERNMENT EXCEL fallback after generic import failed.\n"
                    f"Generic import error was: {generic_exc}\n\n" + str(report)
                )
                return df, sheet_used, report, preview
            except Exception as special_exc:
                raise ValueError(
                    "Could not import this workbook using either generic import or government workbook import.\n\n"
                    f"Generic error: {generic_exc}\n"
                    f"Government-layout error: {special_exc}\n\n"
                    "Beginner fix:\n"
                    "1) Close the Excel file if it is open.\n"
                    "2) Put the Excel file in the same PharmaGuard folder.\n"
                    "3) Rename it to a simple English name like nasser.xlsx.\n"
                    "4) Open Import again and select it from the PharmaGuard folder."
                )

        raise generic_exc
    finally:
        if tmp_dir is not None:
            try:
                import shutil
                shutil.rmtree(str(tmp_dir), ignore_errors=True)
            except Exception:
                pass



# ----------------------------------------------------------------------
# IMPORT DIAGNOSIS / PREVIEW HELPERS (V9)
# These helpers do not change the analytical features. They only make the
# import screen clearer for beginners and safer for hospital Excel files.
# ----------------------------------------------------------------------
def build_import_diagnosis(df, preview_df=None, report_text=""):
    rows = []

    def add(check, result, details=""):
        rows.append({"Check": check, "Result": result, "Details": details})

    if df is None or getattr(df, "empty", True):
        add("Import result", "FAILED / EMPTY", "No usable rows were loaded.")
        return pd.DataFrame(rows)

    add("Rows loaded", "OK", f"{len(df):,} normalized rows")
    add("Columns loaded", "OK", f"{len(df.columns):,} columns")

    required = ["medicine_name", "current_stock", "avg_monthly_consumption", "branch_name"]
    missing_required = [c for c in required if c not in df.columns]
    add(
        "Required columns",
        "OK" if not missing_required else "NEEDS REVIEW",
        ", ".join(missing_required) if missing_required else "medicine, stock, consumption and branch are present",
    )

    if "medicine_name" in df.columns:
        empty_names = int(df["medicine_name"].astype(str).str.strip().eq("").sum())
        add("Medicine names", "OK" if empty_names == 0 else "NEEDS REVIEW", f"empty medicine names: {empty_names:,}")
        try:
            add("Unique medicines", "INFO", f"{df['medicine_name'].astype(str).nunique():,}")
        except Exception:
            pass

    if "branch_name" in df.columns:
        branches = sorted(df["branch_name"].dropna().astype(str).unique().tolist())
        add(
            "Branches / stores",
            "OK" if branches else "NEEDS REVIEW",
            f"{len(branches):,} detected: " + ", ".join(branches[:12]) + (" ..." if len(branches) > 12 else ""),
        )

    if "current_stock" in df.columns:
        stock = pd.to_numeric(df["current_stock"], errors="coerce")
        add(
            "Stock values",
            "OK" if stock.notna().any() else "NEEDS REVIEW",
            f"missing/non-numeric stock rows: {int(stock.isna().sum()):,}; total stock: {float(stock.fillna(0).sum()):,.0f}",
        )

    if "avg_monthly_consumption" in df.columns:
        cons = pd.to_numeric(df["avg_monthly_consumption"], errors="coerce")
        zero_cons = int((cons.fillna(0) <= 0).sum())
        add(
            "Monthly consumption",
            "OK" if cons.notna().any() else "NEEDS REVIEW",
            f"zero/blank consumption rows: {zero_cons:,}; total monthly consumption: {float(cons.fillna(0).sum()):,.0f}",
        )

    if "expiry_date" in df.columns:
        exp = pd.to_datetime(df["expiry_date"], errors="coerce")
        missing_exp = int(exp.isna().sum())
        if missing_exp == len(df):
            add("Expiry / batch", "MISSING", "No usable expiry dates were found. This is expected for consumption-only workbooks. Import a batch-expiry file later.")
        else:
            add("Expiry / batch", "OK" if missing_exp < len(df) else "NEEDS REVIEW", f"missing expiry rows: {missing_exp:,}")
    else:
        add("Expiry / batch", "MISSING", "No expiry_date column was found.")

    if "item_code" in df.columns:
        duplicate_cols = ["item_code", "branch_name"] if "branch_name" in df.columns else ["item_code"]
        duplicate_codes = int(df.duplicated(subset=duplicate_cols).sum())
        add("Duplicate item-code/branch", "OK" if duplicate_codes == 0 else "NEEDS REVIEW", f"duplicates: {duplicate_codes:,}")

    if report_text:
        report_lower = str(report_text).lower()
        if "special government" in report_lower or "branches detected" in report_lower or "rows extracted" in report_lower:
            add("Importer mode", "SPECIAL GOVERNMENT EXCEL", "Arabic multi-row workbook / branch columns were detected and normalized.")
        elif "csv" in report_lower or "google sheet" in report_lower:
            add("Importer mode", "CSV / SHEET", "Standard CSV/Google Sheet parser was used.")
        else:
            add("Importer mode", "GENERIC", "Generic Excel mapping was used.")

    return pd.DataFrame(rows)


def build_import_quick_summary(df, sheet_used="", report_text=""):
    if df is None or getattr(df, "empty", True):
        return "Import summary: no data loaded yet."
    parts = [f"Sheet/source: {sheet_used}", f"Rows: {len(df):,}"]
    if "medicine_name" in df.columns:
        parts.append(f"Medicines: {df['medicine_name'].astype(str).nunique():,}")
    if "branch_name" in df.columns:
        parts.append(f"Branches/stores: {df['branch_name'].dropna().astype(str).nunique():,}")
    if "current_stock" in df.columns:
        stock_sum = pd.to_numeric(df["current_stock"], errors="coerce").fillna(0).sum()
        parts.append(f"Total stock: {stock_sum:,.0f}")
    if "avg_monthly_consumption" in df.columns:
        cons_sum = pd.to_numeric(df["avg_monthly_consumption"], errors="coerce").fillna(0).sum()
        parts.append(f"Monthly consumption: {cons_sum:,.0f}")
    return " | ".join(parts)

def create_demo_df():
    return normalize_dataframe(pd.DataFrame(DEMO_DATA))
def make_scroll_page(widget):
    area = QScrollArea()
    area.setWidgetResizable(True)
    area.setWidget(widget)
    area.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
    area.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
    return area
def make_card(title, value, subtitle=""):
    frame = QFrame()
    frame.setObjectName("Card")
    layout = QVBoxLayout(frame)
    title_label = QLabel(title)
    title_label.setStyleSheet("font-size:14px; font-weight:600;")
    value_label = QLabel(str(value))
    value_label.setStyleSheet("font-size:28px; font-weight:700;")
    sub_label = QLabel(subtitle)
    sub_label.setWordWrap(True)
    sub_label.setStyleSheet("color:#64748b;")
    layout.addWidget(title_label)
    layout.addWidget(value_label)
    layout.addWidget(sub_label)
    return frame
class MatplotlibChart(QWidget):
    def __init__(self, title=""):
        super().__init__()
        layout = QVBoxLayout(self)
        self.title = QLabel(title)
        self.title.setStyleSheet("font-weight:600; font-size:14px;")
        layout.addWidget(self.title)
        if HAS_MATPLOTLIB:
            self.figure = Figure(figsize=(5, 3), constrained_layout=True)
            self.canvas = FigureCanvas(self.figure)
            layout.addWidget(self.canvas)
        else:
            self.figure = None
            self.canvas = QLabel(t("ar" if QApplication.layoutDirection() == Qt.RightToLeft else "en", "matplotlib_missing"))
            layout.addWidget(self.canvas)
    def plot_bar(self, labels, values, title=""):
        if not HAS_MATPLOTLIB:
            return
        self.title.setText(title)
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        shaped_title = ui_text(title)
        shaped_labels = [ui_text(x) for x in labels]
        positions = list(range(len(shaped_labels)))
        ax.bar(positions, values)
        ax.set_xticks(positions)
        ax.set_xticklabels(shaped_labels, rotation=25, ha='right', fontsize=8, fontname=chart_font_family())
        ax.set_title(shaped_title, fontname=chart_font_family())
        ax.tick_params(axis='y', labelsize=8)
        for tick in ax.get_yticklabels():
            tick.set_fontname(chart_font_family())
        self.figure.subplots_adjust(left=0.10, right=0.98, top=0.90, bottom=0.32)
        self.canvas.draw_idle()
    def plot_pie(self, labels, values, title=""):
        if not HAS_MATPLOTLIB:
            return
        self.title.setText(title)
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        shaped_title = ui_text(title)
        shaped_labels = [ui_text(x) for x in labels]
        wedges, texts, autotexts = ax.pie(values, labels=shaped_labels, autopct='%1.0f%%')
        ax.set_title(shaped_title, fontname=chart_font_family())
        for txt in list(texts) + list(autotexts):
            txt.set_fontname(chart_font_family())
        self.figure.subplots_adjust(left=0.06, right=0.94, top=0.88, bottom=0.08)
        self.canvas.draw_idle()
class PharmaGuardMainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.lang = "en"
        self.dark_mode = False
        self.icon_path = Path(__file__).resolve().with_name("pharmaguard_icon.ico")
        if self.icon_path.exists():
            self.setWindowIcon(QIcon(str(self.icon_path)))
        self.setWindowTitle(APP_TITLE_EN)
        self.resize(1560, 980)
        self.df_raw = create_demo_df()
        self.df = compute_metrics(self.df_raw)
        self.current_sheet = "Demo"
        self.mapping_report_text = "Demo data loaded."
        self.preview_df = self.df.head(50).copy()
        self.import_diagnosis_df = build_import_diagnosis(self.df_raw, self.preview_df, self.mapping_report_text)
        self.import_summary_text = build_import_quick_summary(self.df_raw, self.current_sheet, self.mapping_report_text)
        self._build_ui()
        self.apply_theme()
        self.refresh_all_views()
    # ---------- UI ----------
    def _build_ui(self):
        self.toolbar = QToolBar("Main Toolbar")
        self.toolbar.setIconSize(QSize(18, 18))
        self.addToolBar(Qt.TopToolBarArea, self.toolbar)
        self.nav_dock = QDockWidget(t(self.lang, "dock_navigation"), self)
        self.nav_dock.setFeatures(QDockWidget.DockWidgetClosable | QDockWidget.DockWidgetMovable)
        self.nav_dock.setMinimumWidth(165)
        self.nav_dock.setMaximumWidth(205)
        self.nav_list = QListWidget()
        self.nav_list.setSpacing(3)
        self.nav_list.setAlternatingRowColors(False)
        self.nav_list.setUniformItemSizes(True)
        for key in NAV_ORDER:
            self.nav_list.addItem(QListWidgetItem(key))
        self.nav_list.currentRowChanged.connect(self.goto_screen)
        self.nav_dock.setWidget(self.nav_list)
        self.addDockWidget(Qt.LeftDockWidgetArea, self.nav_dock)
        self.tools_dock = QDockWidget(t(self.lang, "dock_tools"), self)
        self.tools_dock.setFeatures(QDockWidget.DockWidgetClosable | QDockWidget.DockWidgetMovable)
        self.tools_dock.setMinimumWidth(195)
        self.tools_dock.setMaximumWidth(240)
        tools_widget = QWidget()
        tools_layout = QVBoxLayout(tools_widget)
        tools_layout.setContentsMargins(6, 6, 6, 6)
        tools_layout.setSpacing(4)
        self.quick_screen = QComboBox()
        self.quick_screen.currentIndexChanged.connect(self.goto_screen)
        self.quick_label = QLabel()
        tools_layout.addWidget(self.quick_label)
        tools_layout.addWidget(self.quick_screen)
        self.search_edit = QLineEdit()
        self.search_edit.textChanged.connect(self.refresh_filters)
        self.search_label = QLabel()
        tools_layout.addWidget(self.search_label)
        tools_layout.addWidget(self.search_edit)
        self.alerts_only = QCheckBox()
        self.alerts_only.stateChanged.connect(self.refresh_filters)
        tools_layout.addWidget(self.alerts_only)
        self.branch_filter = QComboBox()
        self.branch_filter.currentIndexChanged.connect(self.refresh_filters)
        self.risk_filter = QComboBox()
        self.risk_filter.currentIndexChanged.connect(self.refresh_filters)
        self.priority_filter = QComboBox()
        self.priority_filter.currentIndexChanged.connect(self.refresh_filters)
        self.branch_label = QLabel()
        self.risk_label = QLabel()
        self.priority_label = QLabel()
        for label_widget, widget in [(self.branch_label, self.branch_filter), (self.risk_label, self.risk_filter), (self.priority_label, self.priority_filter)]:
            tools_layout.addWidget(label_widget)
            tools_layout.addWidget(widget)
        tools_layout.addStretch(1)
        self.tools_dock.setWidget(tools_widget)
        self.addDockWidget(Qt.RightDockWidgetArea, self.tools_dock)
        self.import_action = QAction("Import", self)
        self.import_action.triggered.connect(self.import_file_dialog)
        self.import_google_action = QAction("Google", self)
        self.import_google_action.triggered.connect(self.import_google_sheet_dialog)
        self.demo_action = QAction("Demo", self)
        self.demo_action.triggered.connect(self.load_demo)
        self.refresh_action = QAction("Refresh", self)
        self.refresh_action.triggered.connect(self.recompute)
        self.export_action = QAction("Export", self)
        self.export_action.triggered.connect(self.export_workbook)
        self.report_action = QAction("Final Report", self)
        self.report_action.triggered.connect(self.export_final_report)
        self.toggle_nav_action = QAction("Nav", self)
        self.toggle_nav_action.triggered.connect(lambda: self.nav_dock.setVisible(not self.nav_dock.isVisible()))
        self.toggle_tools_action = QAction("Tools", self)
        self.toggle_tools_action.triggered.connect(lambda: self.tools_dock.setVisible(not self.tools_dock.isVisible()))
        self.collapse_action = QAction("Collapse", self)
        self.collapse_action.triggered.connect(self.collapse_side_panels)
        self.expand_action = QAction("Expand", self)
        self.expand_action.triggered.connect(self.expand_side_panels)
        for act in [self.import_action, self.import_google_action, self.demo_action, self.refresh_action, self.export_action, self.report_action,
                    self.toggle_nav_action, self.toggle_tools_action, self.collapse_action, self.expand_action]:
            self.toolbar.addAction(act)
        top_widget = QWidget()
        top_layout = QHBoxLayout(top_widget)
        top_layout.setContentsMargins(0, 0, 0, 0)
        top_layout.setSpacing(6)
        self.title_label = QLabel()
        self.title_label.setStyleSheet("font-size:18px; font-weight:700;")
        self.subtitle_label = QLabel()
        self.subtitle_label.setWordWrap(True)
        self.subtitle_label.setStyleSheet("color:#64748b;")
        title_col = QVBoxLayout()
        title_col.setSpacing(2)
        title_col.setContentsMargins(0, 0, 0, 0)
        title_col.addWidget(self.title_label)
        title_col.addWidget(self.subtitle_label)
        top_layout.addLayout(title_col)
        top_layout.addStretch(1)
        self.lang_combo = QComboBox()
        self.lang_combo.addItems(["English", "العربية"])
        self.lang_combo.currentIndexChanged.connect(self.on_language_change)
        self.theme_combo = QComboBox()
        self.theme_combo.addItems(["Light", "Dark"])
        self.theme_combo.currentIndexChanged.connect(self.on_theme_change)
        top_layout.addWidget(self.lang_combo)
        top_layout.addWidget(self.theme_combo)
        self.stack = QStackedWidget()
        self.pages = {}
        for key in NAV_ORDER:
            page_widget = self.build_page(key)
            self.pages[key] = page_widget
            self.stack.addWidget(make_scroll_page(page_widget))
        central = QWidget()
        central_layout = QVBoxLayout(central)
        central_layout.setContentsMargins(2, 2, 2, 2)
        central_layout.setSpacing(4)
        central_layout.addWidget(top_widget)
        central_layout.addWidget(self.stack)
        self.setCentralWidget(central)
        self.status = QStatusBar()
        self.setStatusBar(self.status)
        self.nav_list.setCurrentRow(0)
    def build_page(self, key):
        container = QWidget()
        lay = QVBoxLayout(container)
        lay.setContentsMargins(2, 2, 2, 2)
        lay.setSpacing(6)
        if key == "dashboard":
            self.kpi_grid = QGridLayout()
            lay.addLayout(self.kpi_grid)
            charts_row = QHBoxLayout()
            self.chart_reorder = MatplotlibChart()
            self.chart_risk = MatplotlibChart()
            charts_row.addWidget(self.chart_reorder)
            charts_row.addWidget(self.chart_risk)
            lay.addLayout(charts_row)
            self.dashboard_table = self.make_table()
            lay.addWidget(self.dashboard_table)
        elif key == "import":
            self.import_group = QGroupBox()
            form = QVBoxLayout(self.import_group)
            btns = QHBoxLayout()
            self.import_btn = QPushButton()
            self.import_btn.clicked.connect(self.import_file_dialog)
            self.import_google_btn = QPushButton()
            self.import_google_btn.clicked.connect(self.import_google_sheet_dialog)
            self.live_sync_center_btn = QPushButton("🔄 Live Sync Center")
            self.live_sync_center_btn.clicked.connect(self.open_live_sync_center)
            self.demo_btn = QPushButton()
            self.demo_btn.clicked.connect(self.load_demo)
            form.addWidget(self.import_btn)
            form.addWidget(self.import_google_btn)
            form.addWidget(self.live_sync_center_btn)
            form.addWidget(self.demo_btn)
            form.addLayout(btns)
            self.sheet_used_label = QLabel()
            form.addWidget(self.sheet_used_label)
            self.scalability_note = QLabel("Google Sheets مناسب للعمليات اليومية الصغيرة والمتوسطة فقط. للأحجام الضخمة استخدم CSV/Parquet/DuckDB." if self.lang == "ar" else "Google Sheets is suitable for small/medium operational sync only. For very large volumes use CSV/Parquet/DuckDB.")
            self.scalability_note.setWordWrap(True)
            self.mapping_report = QPlainTextEdit()
            self.mapping_report.setReadOnly(True)
            self.mapping_report.setMinimumHeight(170)
            self.import_summary_label = QLabel()
            self.import_summary_label.setWordWrap(True)
            self.import_summary_label.setStyleSheet("font-weight:700; padding:6px; border:1px solid #cbd5e1; border-radius:8px;")
            self.import_diagnosis_title = QLabel("Import diagnosis / تشخيص الاستيراد")
            self.import_diagnosis_title.setStyleSheet("font-size:14px; font-weight:800; margin-top:8px;")
            self.import_diagnosis_table = self.make_table()
            self.import_diagnosis_table.setMinimumHeight(220)
            self.preview_title = QLabel("Data preview / معاينة البيانات")
            self.preview_title.setStyleSheet("font-size:14px; font-weight:800; margin-top:8px;")
            self.preview_table = self.make_table()
            self.preview_table.setMinimumHeight(300)
            form.addWidget(self.import_summary_label)
            form.addWidget(self.import_diagnosis_title)
            form.addWidget(self.import_diagnosis_table)
            form.addWidget(self.preview_title)
            form.addWidget(self.preview_table)
            form.addWidget(QLabel("Mapping / processing log"))
            form.addWidget(self.mapping_report)
            lay.addWidget(self.import_group)
        elif key == "inventory":
            self.inventory_top_group = QGroupBox()
            fl = QFormLayout(self.inventory_top_group)
            fl.setContentsMargins(8, 8, 8, 6)
            fl.setVerticalSpacing(6)
            self.medicine_combo = QComboBox()
            self.medicine_combo.setEditable(True)
            self.medicine_combo.setInsertPolicy(QComboBox.NoInsert)
            self.medicine_combo.currentIndexChanged.connect(self.refresh_inventory_details)
            self.medicine_search_timer = QTimer(self)
            self.medicine_search_timer.setSingleShot(True)
            self.medicine_search_timer.timeout.connect(self._apply_medicine_text_search)
            self.medicine_combo.currentTextChanged.connect(self.on_medicine_text_changed)
            self.medicine_completer_model = QStringListModel([])
            self.medicine_completer = QCompleter(self.medicine_completer_model, self.medicine_combo)
            self.medicine_completer.setCaseSensitivity(Qt.CaseInsensitive)
            self.medicine_completer.setFilterMode(Qt.MatchContains)
            self.medicine_completer.setCompletionMode(QCompleter.PopupCompletion)
            self.medicine_combo.setCompleter(self.medicine_completer)
            self.inventory_selector_label = QLabel()
            fl.addRow(self.inventory_selector_label, self.medicine_combo)
            lay.addWidget(self.inventory_top_group)
            self.inventory_summary = QTextEdit()
            self.inventory_summary.setReadOnly(True)
            self.inventory_summary.setMinimumHeight(140)
            lay.addWidget(self.inventory_summary)
            self.compare_group = QGroupBox()
            compare_layout = QVBoxLayout(self.compare_group)
            self.export_compare_btn = QPushButton()
            self.export_compare_btn.clicked.connect(self.export_branch_compare)
            self.branch_compare_table = self.make_table()
            self.branch_compare_table.setMinimumHeight(260)
            compare_layout.addWidget(self.export_compare_btn)
            compare_layout.addWidget(self.branch_compare_table)
            lay.addWidget(self.compare_group)
            self.scenario_group = QGroupBox()
            scenario_form = QFormLayout(self.scenario_group)
            self.scen_lead = QDoubleSpinBox(); self.scen_lead.setRange(0, 365); self.scen_lead.setDecimals(0)
            self.scen_pending = QDoubleSpinBox(); self.scen_pending.setRange(0, 1000000); self.scen_pending.setDecimals(0)
            self.scen_extra = QDoubleSpinBox(); self.scen_extra.setRange(0, 1000000); self.scen_extra.setDecimals(0)
            self.scenario_btn = QPushButton()
            self.scenario_btn.clicked.connect(self.run_scenario)
            self.scenario_result = QTextEdit(); self.scenario_result.setReadOnly(True)
            self.scen_lead_label = QLabel(); self.scen_pending_label = QLabel(); self.scen_extra_label = QLabel()
            scenario_form.addRow(self.scen_lead_label, self.scen_lead)
            scenario_form.addRow(self.scen_pending_label, self.scen_pending)
            scenario_form.addRow(self.scen_extra_label, self.scen_extra)
            scenario_form.addRow("", self.scenario_btn)
            scenario_form.addRow(self.scenario_result)
            lay.addWidget(self.scenario_group)
        elif key == "forecast":
            self.forecast_group = QGroupBox()
            fwrap = QVBoxLayout(self.forecast_group)
            fctrl = QHBoxLayout()
            self.forecast_horizon_combo = QComboBox()
            self.forecast_horizon_combo.addItems(["30", "60", "90"])
            self.forecast_horizon_combo.currentIndexChanged.connect(self.refresh_current_view)
            self.forecast_page_size_combo = QComboBox()
            self.forecast_page_size_combo.addItems(["50", "100", "200", "500"])
            self.forecast_page_size_combo.setCurrentText("100")
            self.forecast_page_size_combo.currentIndexChanged.connect(self._reset_forecast_page)
            self.forecast_prev_btn = QPushButton()
            self.forecast_next_btn = QPushButton()
            self.forecast_prev_btn.clicked.connect(self._prev_forecast_page)
            self.forecast_next_btn.clicked.connect(self._next_forecast_page)
            self.forecast_page_label = QLabel("1 / 1")
            self.forecast_horizon_label = QLabel()
            self.forecast_page_size_label = QLabel()
            fctrl.addWidget(self.forecast_horizon_label)
            fctrl.addWidget(self.forecast_horizon_combo)
            fctrl.addSpacing(8)
            fctrl.addWidget(self.forecast_page_size_label)
            fctrl.addWidget(self.forecast_page_size_combo)
            fctrl.addStretch(1)
            fctrl.addWidget(self.forecast_prev_btn)
            fctrl.addWidget(self.forecast_page_label)
            fctrl.addWidget(self.forecast_next_btn)
            self.forecast_kpi_grid = QGridLayout()
            charts_row = QHBoxLayout()
            self.chart_forecast_gap = MatplotlibChart()
            self.chart_forecast_mix = MatplotlibChart()
            charts_row.addWidget(self.chart_forecast_gap)
            charts_row.addWidget(self.chart_forecast_mix)
            self.forecast_table = self.make_table()
            fwrap.addLayout(fctrl)
            fwrap.addLayout(self.forecast_kpi_grid)
            fwrap.addLayout(charts_row)
            fwrap.addWidget(self.forecast_table)
            lay.addWidget(self.forecast_group)
        elif key == "stockbalance":
            self.stockbalance_group = QGroupBox()
            swrap = QVBoxLayout(self.stockbalance_group)
            sctrl = QHBoxLayout()
            self.balance_status_combo = QComboBox()
            self.balance_status_combo.addItems(["All", "Shortage", "Overstock", "Balanced"])
            self.balance_status_combo.currentIndexChanged.connect(self._reset_stockbalance_page)
            self.stockbalance_page_size_combo = QComboBox()
            self.stockbalance_page_size_combo.addItems(["50", "100", "200", "500"])
            self.stockbalance_page_size_combo.setCurrentText("100")
            self.stockbalance_page_size_combo.currentIndexChanged.connect(self._reset_stockbalance_page)
            self.stockbalance_prev_btn = QPushButton()
            self.stockbalance_next_btn = QPushButton()
            self.stockbalance_prev_btn.clicked.connect(self._prev_stockbalance_page)
            self.stockbalance_next_btn.clicked.connect(self._next_stockbalance_page)
            self.stockbalance_page_label = QLabel("1 / 1")
            self.balance_status_label = QLabel()
            self.stockbalance_page_size_label = QLabel()
            sctrl.addWidget(self.balance_status_label)
            sctrl.addWidget(self.balance_status_combo)
            sctrl.addSpacing(8)
            sctrl.addWidget(self.stockbalance_page_size_label)
            sctrl.addWidget(self.stockbalance_page_size_combo)
            sctrl.addStretch(1)
            sctrl.addWidget(self.stockbalance_prev_btn)
            sctrl.addWidget(self.stockbalance_page_label)
            sctrl.addWidget(self.stockbalance_next_btn)
            self.stockbalance_kpi_grid = QGridLayout()
            scharts = QHBoxLayout()
            self.chart_stock_mix = MatplotlibChart()
            self.chart_stock_gap = MatplotlibChart()
            scharts.addWidget(self.chart_stock_mix)
            scharts.addWidget(self.chart_stock_gap)
            self.stockbalance_table = self.make_table()
            swrap.addLayout(sctrl)
            swrap.addLayout(self.stockbalance_kpi_grid)
            swrap.addLayout(scharts)
            swrap.addWidget(self.stockbalance_table)
            lay.addWidget(self.stockbalance_group)
        elif key == "connector":
            self.connector_group = QGroupBox()
            cwrap = QVBoxLayout(self.connector_group)
            brow = QHBoxLayout()
            self.connector_import_excel_btn = QPushButton()
            self.connector_import_google_btn = QPushButton()
            self.connector_import_db_btn = QPushButton()
            self.connector_import_excel_btn.clicked.connect(self.import_file_dialog)
            self.connector_import_google_btn.clicked.connect(self.import_google_sheet_dialog)
            self.connector_import_db_btn.clicked.connect(self.import_database_dialog)
            brow.addWidget(self.connector_import_excel_btn)
            brow.addWidget(self.connector_import_google_btn)
            brow.addWidget(self.connector_import_db_btn)
            self.connector_info = QTextEdit()
            self.connector_info.setReadOnly(True)
            self.connector_preview_table = self.make_table()
            cwrap.addLayout(brow)
            cwrap.addWidget(self.connector_info)
            cwrap.addWidget(self.connector_preview_table)
            lay.addWidget(self.connector_group)
        elif key == "purchase":
            self.purchase_table = self.make_table()
            self.purchase_controls = QGroupBox()
            pform = QFormLayout(self.purchase_controls)
            self.purchase_horizon_combo = QComboBox()
            self.purchase_horizon_combo.addItems(["1", "2", "3", "6", "12"])
            self.purchase_horizon_combo.currentIndexChanged.connect(self.refresh_purchase_view_only)
            self.purchase_item_combo = QComboBox()
            self.purchase_item_combo.currentIndexChanged.connect(self.refresh_draft)
            self.purchase_horizon_label = QLabel()
            self.purchase_item_label = QLabel()
            pform.addRow(self.purchase_horizon_label, self.purchase_horizon_combo)
            pform.addRow(self.purchase_item_label, self.purchase_item_combo)
            self.draft_edit = QPlainTextEdit()
            self.save_draft_btn = QPushButton()
            self.save_draft_btn.clicked.connect(self.save_draft)
            lay.addWidget(self.purchase_controls)
            lay.addWidget(self.purchase_table)
            lay.addWidget(self.draft_edit)
            lay.addWidget(self.save_draft_btn)
        elif key == "redistribution":
            self.redistribution_table = self.make_table()
            lay.addWidget(self.redistribution_table)
        elif key == "fefo":
            self.fefo_table = self.make_table()
            lay.addWidget(self.fefo_table)
        elif key == "abcven":
            self.abcven_table = self.make_table()
            charts_row = QHBoxLayout()
            self.chart_abc = MatplotlibChart()
            self.chart_ven = MatplotlibChart()
            charts_row.addWidget(self.chart_abc)
            charts_row.addWidget(self.chart_ven)
            lay.addLayout(charts_row)
            lay.addWidget(self.abcven_table)
        elif key == "supplier":
            self.supplier_table = self.make_table()
            self.chart_supplier = MatplotlibChart()
            lay.addWidget(self.chart_supplier)
            lay.addWidget(self.supplier_table)
        elif key == "safety":
            self.safety_table = self.make_table()
            lay.addWidget(self.safety_table)
        elif key == "emergency":
            self.emergency_table = self.make_table()
            lay.addWidget(self.emergency_table)
        elif key == "quality":
            self.quality_text = QTextEdit()
            self.quality_text.setReadOnly(True)
            self.quality_table = self.make_table()
            lay.addWidget(self.quality_text)
            lay.addWidget(self.quality_table)
        elif key == "governance":
            self.gov_text = QTextEdit()
            self.gov_text.setReadOnly(True)
            self.prompt_combo = QComboBox()
            self.prompt_combo.currentIndexChanged.connect(self.refresh_prompt)
            self.prompt_text = QPlainTextEdit()
            self.prompt_text.setReadOnly(True)
            lay.addWidget(self.gov_text)
            lay.addWidget(self.prompt_combo)
            lay.addWidget(self.prompt_text)
        lay.addStretch(1)
        return container
    def make_table(self):
        table = QTableWidget()
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setAlternatingRowColors(True)
        table.setWordWrap(False)
        table.setTextElideMode(Qt.ElideNone)
        table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        table.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
        table.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOn)
        table.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        table.horizontalHeader().setStretchLastSection(False)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        table.horizontalHeader().setMinimumSectionSize(90)
        table.verticalHeader().setVisible(False)
        table.verticalHeader().setDefaultSectionSize(28)
        table.setStyleSheet("QTableWidget{gridline-color:#d9e2ef;font-size:11px;} QHeaderView::section{padding:6px 8px;font-weight:700;}")
        return table
    def apply_theme(self):
        self.setStyleSheet(DARK_QSS if self.dark_mode else LIGHT_QSS)
    # ---------- events ----------
    def on_language_change(self, index):
        self.lang = "en" if index == 0 else "ar"
        self.update_labels()
    def on_theme_change(self, index):
        self.dark_mode = index == 1
        self.apply_theme()
    def collapse_side_panels(self):
        self.nav_dock.hide()
        self.tools_dock.hide()
    def expand_side_panels(self):
        self.nav_dock.show()
        self.tools_dock.show()
    def goto_screen(self, index):
        if index < 0:
            return
        self.stack.setCurrentIndex(index)
        if self.nav_list.currentRow() != index:
            self.nav_list.blockSignals(True)
            self.nav_list.setCurrentRow(index)
            self.nav_list.blockSignals(False)
        if self.quick_screen.currentIndex() != index:
            self.quick_screen.blockSignals(True)
            self.quick_screen.setCurrentIndex(index)
            self.quick_screen.blockSignals(False)
        self.refresh_current_view()
    def update_labels(self):
        self.setWindowTitle(t(self.lang, "title"))
        self.setLayoutDirection(Qt.RightToLeft if self.lang == "ar" else Qt.LeftToRight)
        self.title_label.setText(t(self.lang, "welcome"))
        self.subtitle_label.setText(t(self.lang, "subtitle"))
        self.nav_dock.setWindowTitle(t(self.lang, "dock_navigation"))
        self.tools_dock.setWindowTitle(t(self.lang, "dock_tools"))
        self.quick_label.setText(t(self.lang, "quick_short"))
        self.search_label.setText(t(self.lang, "search_short"))
        self.branch_label.setText(t(self.lang, "branch"))
        self.risk_label.setText(t(self.lang, "risk"))
        self.priority_label.setText(t(self.lang, "priority"))
        self.import_action.setText(t(self.lang, "import_excel"))
        if hasattr(self, "import_google_action"): self.import_google_action.setText(t(self.lang, "import_google"))
        self.demo_action.setText(t(self.lang, "load_demo"))
        self.refresh_action.setText(t(self.lang, "refresh"))
        self.export_action.setText(t(self.lang, "export_workbook"))
        self.report_action.setText(t(self.lang, "export_report"))
        self.toggle_nav_action.setText(t(self.lang, "toggle_nav"))
        self.toggle_tools_action.setText(t(self.lang, "toggle_side"))
        self.collapse_action.setText(t(self.lang, "collapse_all"))
        self.expand_action.setText(t(self.lang, "expand_all"))
        self.alerts_only.setText(t(self.lang, "alerts_only"))
        self.search_edit.setPlaceholderText(t(self.lang, "search"))
        self.lang_combo.blockSignals(True)
        self.lang_combo.setItemText(0, "English")
        self.lang_combo.setItemText(1, "العربية")
        self.lang_combo.setCurrentIndex(0 if self.lang == "en" else 1)
        self.lang_combo.blockSignals(False)
        self.theme_combo.blockSignals(True)
        self.theme_combo.setItemText(0, t(self.lang, "light"))
        self.theme_combo.setItemText(1, t(self.lang, "dark"))
        self.theme_combo.blockSignals(False)
        self.quick_screen.blockSignals(True)
        self.quick_screen.clear()
        self.nav_list.blockSignals(True)
        self.nav_list.clear()
        for key in NAV_ORDER:
            label = t(self.lang, NAV_TEXT[key])
            self.quick_screen.addItem(label)
            self.nav_list.addItem(QListWidgetItem(label))
        self.nav_list.blockSignals(False)
        self.quick_screen.blockSignals(False)
        current = self.stack.currentIndex()
        self.nav_list.setCurrentRow(current)
        self.quick_screen.setCurrentIndex(current)
        all_label = t(self.lang, "all")
        self.branch_filter.blockSignals(True)
        branch = self.branch_filter.currentText()
        self.branch_filter.clear()
        self.branch_filter.addItems([all_label])
        self.branch_filter.addItems(sorted(self.df["branch_name"].dropna().astype(str).unique().tolist()))
        self.branch_filter.blockSignals(False)
        self.risk_filter.blockSignals(True)
        self.risk_filter.clear()
        self.risk_filter.addItems([all_label, translate_display_value("Low", self.lang), translate_display_value("Moderate", self.lang), translate_display_value("High", self.lang), translate_display_value("Critical", self.lang)])
        self.risk_filter.blockSignals(False)
        self.priority_filter.blockSignals(True)
        self.priority_filter.clear()
        self.priority_filter.addItems([all_label, translate_display_value("Low", self.lang), "Medium" if self.lang=="en" else "متوسط", translate_display_value("High", self.lang), translate_display_value("Critical", self.lang)])
        self.priority_filter.blockSignals(False)
        self.import_group.setTitle(t(self.lang, "import_group_title"))
        self.inventory_top_group.setTitle(t(self.lang, "medicine_lookup"))
        self.inventory_selector_label.setText(t(self.lang, "inventory_selector"))
        if self.medicine_combo.lineEdit() is not None:
            self.medicine_combo.lineEdit().setPlaceholderText(t(self.lang, "type_to_search"))
        self.compare_group.setTitle(t(self.lang, "branch_compare_title"))
        self.export_compare_btn.setText(t(self.lang, "export_compare"))
        self.scenario_group.setTitle(t(self.lang, "scenario_title"))
        self.scen_lead_label.setText(t(self.lang, "lead_time"))
        self.scen_pending_label.setText(t(self.lang, "pending_po"))
        self.scen_extra_label.setText(t(self.lang, "extra_stock"))
        self.import_btn.setText(t(self.lang, "import_excel"))
        if hasattr(self, "import_google_btn"): self.import_google_btn.setText(t(self.lang, "import_google"))
        if hasattr(self, "purchase_horizon_label"): self.purchase_horizon_label.setText(t(self.lang, "purchase_horizon"))
        if hasattr(self, "purchase_item_label"): self.purchase_item_label.setText(t(self.lang, "medicine_lookup"))
        self.demo_btn.setText(t(self.lang, "load_demo"))
        if hasattr(self, "live_sync_center_btn"):
            self.live_sync_center_btn.setText("🔄 Live Sync Center / مركز المزامنة")
        self.save_draft_btn.setText(t(self.lang, "save_txt"))
        self.scenario_btn.setText(t(self.lang, "run_scenario"))
        self.sheet_used_label.setText(f"{t(self.lang, 'sheet_used')}: {self.current_sheet}")
        self.mapping_report.setPlainText(self.mapping_report_text)
        self.prompt_combo.blockSignals(True)
        self.prompt_combo.clear()
        self.prompt_combo.addItems(list(PROMPTS.keys()))
        self.prompt_combo.blockSignals(False)
        self.refresh_prompt()
        if hasattr(self, 'connector_import_excel_btn'):
            self.connector_group.setTitle(t(self.lang, 'connector_center'))
            self.connector_import_excel_btn.setText(t(self.lang, 'import_excel'))
            self.connector_import_google_btn.setText(t(self.lang, 'import_google'))
            self.connector_import_db_btn.setText(t(self.lang, 'import_db'))
        if hasattr(self, 'forecast_group'):
            self.forecast_group.setTitle(t(self.lang, 'forecast_center'))
            self.forecast_horizon_label.setText(t(self.lang, 'forecast_horizon'))
            self.forecast_page_size_label.setText(t(self.lang, 'page_size'))
            self.forecast_prev_btn.setText(t(self.lang, 'prev_page'))
            self.forecast_next_btn.setText(t(self.lang, 'next_page'))
        if hasattr(self, 'stockbalance_group'):
            self.stockbalance_group.setTitle(t(self.lang, 'stock_balance_center'))
            self.balance_status_label.setText(t(self.lang, 'status_view'))
            self.stockbalance_page_size_label.setText(t(self.lang, 'page_size'))
            self.stockbalance_prev_btn.setText(t(self.lang, 'prev_page'))
            self.stockbalance_next_btn.setText(t(self.lang, 'next_page'))
            labels = [t(self.lang, 'all'), t(self.lang, 'balance_shortage'), t(self.lang, 'balance_overstock'), t(self.lang, 'balance_balanced')]
            current_idx = self.balance_status_combo.currentIndex()
            self.balance_status_combo.blockSignals(True)
            self.balance_status_combo.clear()
            self.balance_status_combo.addItems(labels)
            self.balance_status_combo.setCurrentIndex(max(0, min(current_idx, len(labels)-1)))
            self.balance_status_combo.blockSignals(False)
        self.refresh_inventory_details()
        self.status.showMessage(t(self.lang, "decision_only"), 5000)
    # ---------- data ----------
    def load_demo(self):
        self.df_raw = create_demo_df()
        self.current_sheet = "Demo"
        self.mapping_report_text = "Demo data loaded."
        self.preview_df = self.df_raw.head(50)
        self.import_diagnosis_df = build_import_diagnosis(self.df_raw, self.preview_df, self.mapping_report_text)
        self.import_summary_text = build_import_quick_summary(self.df_raw, self.current_sheet, self.mapping_report_text)
        self.recompute()
        QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "import_ok"))

    def open_live_sync_center(self):
        import subprocess
        from pathlib import Path

        project_root = Path(__file__).resolve().parent
        script_path = project_root / "run_pharmaguard.py"

        if not script_path.exists():
            QMessageBox.critical(
                self,
                t(self.lang, "title"),
                f"لم يتم العثور على ملف مركز المزامنة:\n{script_path}"
            )
            return

        subprocess.Popen(
            [sys.executable, str(script_path)],
            cwd=str(project_root)
        )
    def import_file_dialog(self):
        path, _ = QFileDialog.getOpenFileName(
            self, t(self.lang, "import_excel"), "", "Excel / CSV Files (*.xlsx *.xlsm *.xls *.csv);;Excel Files (*.xlsx *.xlsm *.xls);;CSV Files (*.csv);;All Files (*.*)"
        )
        if not path:
            return
        try:
            df, sheet_used, report, preview = import_any_file(path)
            self.df_raw = df
            self.current_sheet = sheet_used
            self.mapping_report_text = report
            self.preview_df = preview
            self.import_diagnosis_df = build_import_diagnosis(self.df_raw, self.preview_df, self.mapping_report_text)
            self.import_summary_text = build_import_quick_summary(self.df_raw, self.current_sheet, self.mapping_report_text)
            self.recompute()
            QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "raw_detected"))
        except Exception as exc:
            traceback.print_exc()
            QMessageBox.critical(self, t(self.lang, "title"), f"{t(self.lang, 'import_fail')}\n{exc}")
    def import_google_sheet_dialog(self):
        url, ok = QInputDialog.getText(self, t(self.lang, "import_google"), t(self.lang, "google_sheet_prompt"))
        if not ok or not clean_text(url):
            return
        try:
            df, sheet_used, report, preview = import_any_file(clean_text(url))
            self.df_raw = df
            self.current_sheet = sheet_used
            self.mapping_report_text = report
            self.preview_df = preview
            self.recompute()
            QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "google_sheet_ok"))
        except Exception as exc:
            traceback.print_exc()
            QMessageBox.critical(self, t(self.lang, "title"), f"{t(self.lang, 'import_fail')}\n{exc}")
    def recompute(self):
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            self.df = compute_metrics(self.df_raw)
            self.search_blob = build_search_blob_series(
                self.df,
                ["medicine_name", "generic_name", "active_ingredient", "dosage_form_group", "brand_names_display", "source_medicine_names", "branch_name", "supplier_name", "item_code"]
            )
            self.populate_filters()
            self.refresh_all_views()
        finally:
            QApplication.restoreOverrideCursor()
    def populate_filters(self):
        current_branch = self.branch_filter.currentText() if self.branch_filter.count() else "All"
        current_risk = self.risk_filter.currentText() if self.risk_filter.count() else "All"
        current_priority = self.priority_filter.currentText() if self.priority_filter.count() else "All"
        self.branch_filter.blockSignals(True)
        self.branch_filter.clear()
        self.branch_filter.addItems([t(self.lang, "all")] + sorted(self.df["branch_name"].dropna().astype(str).unique().tolist()))
        self.risk_filter.blockSignals(True)
        self.risk_filter.clear()
        self.risk_filter.addItems([t(self.lang, "all"), translate_display_value("Low", self.lang), translate_display_value("Moderate", self.lang), translate_display_value("High", self.lang), translate_display_value("Critical", self.lang)])
        self.priority_filter.blockSignals(True)
        self.priority_filter.clear()
        self.priority_filter.addItems([t(self.lang, "all"), translate_display_value("Low", self.lang), "Medium" if self.lang=="en" else "متوسط", translate_display_value("High", self.lang), translate_display_value("Critical", self.lang)])
        for combo, val in [(self.branch_filter, current_branch), (self.risk_filter, current_risk), (self.priority_filter, current_priority)]:
            idx = combo.findText(val)
            combo.setCurrentIndex(max(idx, 0))
        self.branch_filter.blockSignals(False)
        self.risk_filter.blockSignals(False)
        self.priority_filter.blockSignals(False)
    def filtered_df(self):
        df = self.df
        if df.empty:
            return df
        mask = pd.Series(True, index=df.index)
        q_bundle = normalize_query_bundle(self.search_edit.text())
        if q_bundle["text"] or q_bundle["latin"] or q_bundle["ar_latin"]:
            blob = getattr(self, "search_blob", None)
            if blob is None or len(blob) != len(df):
                blob = build_search_blob_series(
                    df,
                    ["medicine_name", "generic_name", "active_ingredient", "dosage_form_group", "brand_names_display", "source_medicine_names", "branch_name", "supplier_name", "item_code"]
                )
                self.search_blob = blob
            mask &= blob.map(lambda x: query_in_blob(q_bundle, x))
        if self.alerts_only.isChecked():
            mask &= ((df["shortage_risk"].isin(["High", "Critical"])) | (df["expiry_risk"].isin(["High", "Critical", "Expired"])) | (df["escalation_level"] == "Red"))
        if self.branch_filter.currentText() not in ("", "All", t(self.lang, "all")):
            mask &= (df["branch_name"] == self.branch_filter.currentText())
        if self.risk_filter.currentText() not in ("", "All", t(self.lang, "all")):
            selected = self.risk_filter.currentText()
            selected_en = next((k for k,v in RISK_TRANSLATIONS.items() if v.get("ar")==selected), selected)
            mask &= (df["shortage_risk"] == selected_en)
        if self.priority_filter.currentText() not in ("", "All", t(self.lang, "all")):
            selected = self.priority_filter.currentText()
            selected_en = next((k for k,v in RISK_TRANSLATIONS.items() if v.get("ar")==selected), "Medium" if selected=="متوسط" else selected)
            mask &= (df["clinical_priority"] == selected_en)
        return df.loc[mask].copy()
    # ---------- render ----------
    def refresh_filters(self):
        self.refresh_current_view()
    def _page_size_from_combo(self, combo, default=100):
        try:
            return int(combo.currentText())
        except Exception:
            return default
    def _paged_slice(self, df, page, page_size):
        if df is None or df.empty:
            return df, 0, 0
        total_pages = max(1, math.ceil(len(df) / max(page_size, 1)))
        page = max(0, min(page, total_pages - 1))
        start = page * page_size
        end = start + page_size
        return df.iloc[start:end].copy(), page, total_pages
    def _reset_forecast_page(self):
        self._forecast_page = 0
        self.refresh_current_view()
    def _prev_forecast_page(self):
        self._forecast_page = max(0, getattr(self, '_forecast_page', 0) - 1)
        self.refresh_current_view()
    def _next_forecast_page(self):
        self._forecast_page = getattr(self, '_forecast_page', 0) + 1
        self.refresh_current_view()
    def _reset_stockbalance_page(self):
        self._stockbalance_page = 0
        self.refresh_current_view()
    def _prev_stockbalance_page(self):
        self._stockbalance_page = max(0, getattr(self, '_stockbalance_page', 0) - 1)
        self.refresh_current_view()
    def _next_stockbalance_page(self):
        self._stockbalance_page = getattr(self, '_stockbalance_page', 0) + 1
        self.refresh_current_view()
    def import_database_dialog(self):
        if self.lang == 'ar':
            choices = [t(self.lang, 'db_sqlite'), t(self.lang, 'db_parquet'), t(self.lang, 'db_csv')]
        else:
            choices = [t(self.lang, 'db_sqlite'), t(self.lang, 'db_parquet'), t(self.lang, 'db_csv')]
        choice, ok = QInputDialog.getItem(self, t(self.lang, 'import_db'), t(self.lang, 'db_prompt'), choices, 0, False)
        if not ok or not choice:
            return
        try:
            if choice == t(self.lang, 'db_sqlite'):
                path, _ = QFileDialog.getOpenFileName(self, t(self.lang, 'import_db'), '', 'SQLite (*.db *.sqlite *.sqlite3)')
                if not path:
                    return
                conn = sqlite3.connect(path)
                try:
                    tables = pd.read_sql_query("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name", conn)["name"].tolist()
                    if not tables:
                        raise ValueError('No tables found in database.')
                    item, ok2 = QInputDialog.getItem(self, t(self.lang, 'import_db'), t(self.lang, 'db_table_prompt'), tables + ['SQL: SELECT ...'], 0, False)
                    if not ok2 or not item:
                        return
                    if item == 'SQL: SELECT ...':
                        sql, ok3 = QInputDialog.getMultiLineText(self, t(self.lang, 'import_db'), t(self.lang, 'db_table_prompt'), 'SELECT * FROM ' + tables[0] + ' LIMIT 500000')
                        if not ok3 or not clean_text(sql):
                            return
                        df = pd.read_sql_query(sql, conn)
                        source_name = f'SQLite query'
                    else:
                        df = pd.read_sql_query(f'SELECT * FROM "{item}"', conn)
                        source_name = f'SQLite:{item}'
                finally:
                    conn.close()
            elif choice == t(self.lang, 'db_parquet'):
                path, _ = QFileDialog.getOpenFileName(self, t(self.lang, 'import_db'), '', 'Parquet / Feather (*.parquet *.feather)')
                if not path:
                    return
                if path.lower().endswith('.feather'):
                    df = pd.read_feather(path)
                else:
                    df = pd.read_parquet(path)
                source_name = Path(path).name
            else:
                path, _ = QFileDialog.getOpenFileName(self, t(self.lang, 'import_db'), '', 'Delimited files (*.csv *.tsv)')
                if not path:
                    return
                sep = '\t' if path.lower().endswith('.tsv') else ','
                df = pd.read_csv(path, sep=sep)
                source_name = Path(path).name
            self.df_raw = df
            self.current_sheet = source_name
            self.mapping_report_text = f"{t(self.lang, 'connector_loaded')}\n{t(self.lang, 'connector_rows')}: {len(df):,}\n{t(self.lang, 'connector_mode')}: SQLite/Parquet/CSV"
            self.preview_df = df.head(200).copy()
            self.recompute()
            QMessageBox.information(self, t(self.lang, 'title'), t(self.lang, 'connector_loaded'))
        except Exception as exc:
            traceback.print_exc()
            QMessageBox.critical(self, t(self.lang, 'title'), f"{t(self.lang, 'import_fail')}\n{exc}")
    def refresh_forecast_center(self, df):
        working = self.build_network_prediction_df(df)
        self.clear_layout(self.forecast_kpi_grid)
        horizon = int(getattr(self, 'forecast_horizon_combo').currentText()) if hasattr(self, 'forecast_horizon_combo') else 30
        if working.empty:
            self.set_table_from_df(self.forecast_table, pd.DataFrame())
            self.forecast_page_label.setText('0 / 0')
            return
        if horizon == 30:
            demand_col, gap_col = 'predicted_next_month_demand', 'store_predicted_gap_next_month'
        elif horizon == 60:
            working['predicted_60d_demand'] = np.round(working['network_monthly_consumption'].astype(float).fillna(0) * 2, 1)
            working['predicted_60d_gap'] = np.round(np.maximum(working['predicted_60d_demand'] - working.get('central_store_available', 0).astype(float).fillna(0), 0), 1)
            demand_col, gap_col = 'predicted_60d_demand', 'predicted_60d_gap'
        else:
            demand_col, gap_col = 'predicted_3_month_demand', 'store_predicted_gap_3_months'
        working['forecast_balance'] = np.round(working.get('central_store_available', 0).astype(float).fillna(0) - working[demand_col].astype(float).fillna(0), 1)
        kpis = [
            (t(self.lang, 'forecast_horizon'), f'{horizon}d', self.current_sheet),
            (t(self.lang, 'forecast_gap'), int(pd.to_numeric(working[gap_col], errors='coerce').fillna(0).sum()), ''),
            (t(self.lang, 'balance_shortage'), int((pd.to_numeric(working[gap_col], errors='coerce').fillna(0) > 0).sum()), ''),
            (t(self.lang, 'balance_overstock'), int((working.get('overstock_risk', pd.Series(['Low']*len(working))).isin(['Moderate','High'])).sum()), ''),
        ]
        for i, (title, value, sub) in enumerate(kpis):
            self.forecast_kpi_grid.addWidget(make_card(title, value, sub), i // 2, i % 2)
        cols = [c for c in ['active_ingredient','generic_name','dosage_form_group','medicine_name','network_monthly_consumption', demand_col,'central_store_available','forecast_balance', gap_col,'highest_shortage_risk','store_support_status','suggested_network_action'] if c in working.columns]
        show = working[cols].sort_values([gap_col, demand_col], ascending=[False, False])
        page_size = self._page_size_from_combo(self.forecast_page_size_combo, 100)
        page_df, page, total_pages = self._paged_slice(show, getattr(self, '_forecast_page', 0), page_size)
        self._forecast_page = page
        self.forecast_page_label.setText(f'{page+1} / {total_pages}')
        self.set_table_from_df(self.forecast_table, page_df)
        if not show.empty and HAS_MATPLOTLIB:
            top = show.head(CHART_TOP_N)
            self.chart_forecast_gap.plot_bar(top['generic_name'].astype(str).tolist(), pd.to_numeric(top[gap_col], errors='coerce').fillna(0).tolist(), t(self.lang, 'top_gap_chart'))
            mix = show['store_support_status'].fillna('Covered').value_counts().head(5)
            self.chart_forecast_mix.plot_pie(mix.index.tolist(), mix.values.tolist(), t(self.lang, 'stock_mix_chart'))
    def refresh_stock_balance_center(self, df):
        working = df.copy()
        self.clear_layout(self.stockbalance_kpi_grid)
        if working.empty:
            self.set_table_from_df(self.stockbalance_table, pd.DataFrame())
            self.stockbalance_page_label.setText('0 / 0')
            return
        selected = self.balance_status_combo.currentText() if hasattr(self, 'balance_status_combo') else 'All'
        selected_en = next((k for k,v in RISK_TRANSLATIONS.items() if v.get('ar') == selected), selected)
        if selected_en == 'Shortage':
            working = working[working['stock_status'] == 'Shortage']
        elif selected_en == 'Overstock':
            working = working[working['stock_status'] == 'Overstock']
        elif selected_en == 'Balanced':
            working = working[working['stock_status'] == 'Balanced']
        kpis = [
            (t(self.lang, 'balance_shortage'), int((df['stock_status'] == 'Shortage').sum()), ''),
            (t(self.lang, 'balance_overstock'), int((df['stock_status'] == 'Overstock').sum()), ''),
            (t(self.lang, 'balance_balanced'), int((df['stock_status'] == 'Balanced').sum()), ''),
            (t(self.lang, 'forecast_gap'), int((pd.to_numeric(df.get('recommended_reorder_qty',0), errors='coerce').fillna(0)).sum()), ''),
        ]
        for i, (title, value, sub) in enumerate(kpis):
            self.stockbalance_kpi_grid.addWidget(make_card(title, value, sub), i // 2, i % 2)
        cols = [c for c in ['branch_name','generic_name','dosage_form_group','current_stock','dynamic_min_stock','days_of_stock_left','coverage_months','shortage_risk','overstock_risk','stock_status','predicted_stock_30d','predicted_stock_60d','predicted_stock_90d','recommended_action'] if c in working.columns]
        show = working[cols].sort_values(['stock_status','overstock_risk','shortage_risk','coverage_months'], ascending=[True, False, False, False])
        page_size = self._page_size_from_combo(self.stockbalance_page_size_combo, 100)
        page_df, page, total_pages = self._paged_slice(show, getattr(self, '_stockbalance_page', 0), page_size)
        self._stockbalance_page = page
        self.stockbalance_page_label.setText(f'{page+1} / {total_pages}')
        self.set_table_from_df(self.stockbalance_table, page_df)
        if not df.empty and HAS_MATPLOTLIB:
            mix = df['stock_status'].fillna('Balanced').value_counts()
            self.chart_stock_mix.plot_pie(mix.index.tolist(), mix.values.tolist(), t(self.lang, 'stock_mix_chart'))
            top = df.sort_values('days_of_stock_left', ascending=True).head(CHART_TOP_N)
            self.chart_stock_gap.plot_bar(top['generic_name'].astype(str).tolist(), pd.to_numeric(top['days_of_stock_left'], errors='coerce').fillna(0).tolist(), t(self.lang, 'chart_risk_title'))
    def refresh_connector_center(self):
        info = []
        is_ar = self.lang == 'ar'
        info.append(f"{t(self.lang, 'sheet_used')}: {self.current_sheet}")
        info.append(f"{t(self.lang, 'connector_rows')}: {len(self.df_raw) if hasattr(self, 'df_raw') else 0:,}")
        info.append(f"{t(self.lang, 'connector_mode')}: Google Sheet / CSV / Parquet / SQLite")
        info.append(t(self.lang, 'google_sync_note'))
        info.append('')
        if is_ar:
            info.append('لأحجام ضخمة جدًا: استخدم SQLite أو DuckDB أو Parquet مع تجميع الاستهلاك شهريًا قبل العرض.')
            info.append('الواجهة الحالية مناسبة للعرض التحليلي، بينما السحب المباشر لعشرات أو مئات الملايين من الحركات يحتاج قاعدة بيانات واستعلامات تجميع.')
        else:
            info.append('For very large volumes, use SQLite/DuckDB/Parquet and pre-aggregate monthly movement before loading into the UI.')
            info.append('This UI is optimized for analytical views. Very large raw transaction logs should be queried and aggregated before display.')
        self.connector_info.setPlainText('\n'.join(info))
        self.set_table_from_df(self.connector_preview_table, self.preview_df.head(100).copy() if self.preview_df is not None else pd.DataFrame())
    def refresh_all_views(self):
        self.update_labels()
        self.sheet_used_label.setText(f"{t(self.lang, 'sheet_used')}: {self.current_sheet}")
        self.mapping_report.setPlainText(self.mapping_report_text)
        if hasattr(self, "import_summary_label"):
            self.import_summary_label.setText(getattr(self, "import_summary_text", build_import_quick_summary(self.df_raw, self.current_sheet, self.mapping_report_text)))
        if hasattr(self, "import_diagnosis_table"):
            self.set_table_from_df(self.import_diagnosis_table, getattr(self, "import_diagnosis_df", build_import_diagnosis(self.df_raw, self.preview_df, self.mapping_report_text)))
        self.set_table_from_df(self.preview_table, self.preview_df)
        self.refresh_current_view()
    def refresh_current_view(self):
        df = self.filtered_df()
        current_key = NAV_ORDER[self.stack.currentIndex()]
        if current_key == "dashboard":
            self.refresh_dashboard(df)
        elif current_key == "import":
            self.sheet_used_label.setText(f"{t(self.lang, 'sheet_used')}: {self.current_sheet}")
            self.mapping_report.setPlainText(self.mapping_report_text)
            if hasattr(self, "import_summary_label"):
                self.import_summary_label.setText(getattr(self, "import_summary_text", build_import_quick_summary(self.df_raw, self.current_sheet, self.mapping_report_text)))
            if hasattr(self, "import_diagnosis_table"):
                self.set_table_from_df(self.import_diagnosis_table, getattr(self, "import_diagnosis_df", build_import_diagnosis(self.df_raw, self.preview_df, self.mapping_report_text)))
            self.set_table_from_df(self.preview_table, self.preview_df)
        elif current_key == "inventory":
            self.refresh_inventory_page(df)
        elif current_key == "forecast":
            self.refresh_forecast_center(df)
        elif current_key == "stockbalance":
            self.refresh_stock_balance_center(df)
        elif current_key == "connector":
            self.refresh_connector_center()
        elif current_key == "purchase":
            self.refresh_purchase(df)
        elif current_key == "redistribution":
            self.refresh_redistribution(df)
        elif current_key == "fefo":
            self.refresh_fefo(df)
        elif current_key == "abcven":
            self.refresh_abcven(df)
        elif current_key == "supplier":
            self.refresh_supplier(df)
        elif current_key == "safety":
            self.refresh_safety(df)
        elif current_key == "emergency":
            self.refresh_emergency(df)
        elif current_key == "quality":
            self.refresh_quality(df)
        elif current_key == "governance":
            self.refresh_governance()
    def refresh_dashboard(self, df):
        self.clear_layout(self.kpi_grid)
        kpis = [
            (t(self.lang, "kpi_total"), len(df), self.current_sheet),
            (t(self.lang, "kpi_critical"), int((df["shortage_risk"] == "Critical").sum()), ""),
            (t(self.lang, "kpi_expiry"), int(df["expiry_risk"].isin(["High", "Critical", "Expired"]).sum()), ""),
            (t(self.lang, "kpi_reorder"), int(df["recommended_reorder_qty"].sum()), ""),
            (t(self.lang, "kpi_red"), int((df["escalation_level"] == "Red").sum()), ""),
        ]
        for i, (title, value, sub) in enumerate(kpis):
            self.kpi_grid.addWidget(make_card(title, value, sub), i // 3, i % 3)
        show = df[[
            "generic_name", "brand_names_display", "dosage_form_group", "branch_name", "current_stock", "dynamic_min_stock",
            "days_of_stock_left", "shortage_risk", "overstock_risk", "stock_status", "predicted_stock_30d", "predicted_stock_60d", "predicted_stock_90d",
            "expiry_risk", "recommended_reorder_qty", "escalation_level"
        ]].sort_values(["escalation_level", "shortage_risk", "expiry_risk"])
        self.set_table_from_df(self.dashboard_table, show)
        if not df.empty and HAS_MATPLOTLIB:
            top = df.sort_values("recommended_reorder_qty", ascending=False).head(CHART_TOP_N)
            self.chart_reorder.plot_bar(top["medicine_name"].tolist(), top["recommended_reorder_qty"].tolist(), t(self.lang, "chart_reorder_title"))
            risk_counts = df["shortage_risk"].value_counts().reindex(["Low", "Moderate", "High", "Critical"]).fillna(0)
            self.chart_risk.plot_pie(risk_counts.index.tolist(), risk_counts.values.tolist(), t(self.lang, "chart_risk_title"))
    def on_medicine_text_changed(self, text):
        if not hasattr(self, "medicine_combo") or self.medicine_combo.count() == 0:
            return
        self._pending_medicine_text = text
        if hasattr(self, "medicine_search_timer"):
            self.medicine_search_timer.start(220)
        else:
            self._apply_medicine_text_search()
    def _apply_medicine_text_search(self):
        if not hasattr(self, "medicine_combo") or self.medicine_combo.count() == 0:
            return
        text = getattr(self, "_pending_medicine_text", self.medicine_combo.currentText())
        query = normalize_search_text(text)
        if not query:
            return
        best_idx = -1
        for idx in range(self.medicine_combo.count()):
            display = normalize_search_text(self.medicine_combo.itemText(idx))
            generic = normalize_search_text(self.medicine_combo.itemData(idx) or "")
            if query in display or query in generic:
                best_idx = idx
                break
        if best_idx >= 0:
            self.medicine_combo.blockSignals(True)
            self.medicine_combo.setCurrentIndex(best_idx)
            self.medicine_combo.blockSignals(False)
            if self.medicine_combo.lineEdit() is not None:
                self.medicine_combo.lineEdit().setText(text)
                self.medicine_combo.lineEdit().setCursorPosition(len(text))
            self.refresh_inventory_details()
    def get_selected_medicine_name(self):
        if not hasattr(self, "medicine_combo") or self.medicine_combo.count() == 0:
            return ""
        idx = self.medicine_combo.currentIndex()
        if idx >= 0:
            data = self.medicine_combo.itemData(idx)
            if data:
                return clean_text(data)
        raw = clean_text(self.medicine_combo.currentText())
        if not raw:
            return clean_text(self.medicine_combo.itemData(0) or self.medicine_combo.itemText(0))
        bundle = normalize_query_bundle(raw)
        cache = getattr(self, "_medicine_display_cache", [])
        best_name = ""
        best_score = 0.0
        for display, analysis_name, blob, generic_blob, aliases in cache:
            score = max(query_match_score(bundle, generic_blob), query_match_score(bundle, blob))
            if score > best_score:
                best_score = score
                best_name = analysis_name
        return clean_text(best_name or self.medicine_combo.itemData(0) or self.medicine_combo.itemText(0))
    def get_medicine_display_map(self, df):
        if df.empty:
            return []
        grouped = (
            df.groupby("medicine_name", dropna=False)
              .agg(
                  generic_name=("generic_name", lambda s: base._join_unique(s, sep=" | ")),
                  active_ingredient=("active_ingredient", lambda s: base._join_unique(s, sep=" | ")),
                  brand_names_display=("brand_names_display", lambda s: base._join_unique(s, sep=" | ")),
                  source_medicine_names=("source_medicine_names", lambda s: base._join_unique(s, sep=" | ")),
                  dosage_form_group=("dosage_form_group", lambda s: base._join_unique(s, sep=" | ")),
              )
              .reset_index()
        )
        items = []
        grouped = grouped.sort_values(["generic_name", "medicine_name"], kind="stable")
        for _, r in grouped.iterrows():
            analysis_name = base.clean_text(r["medicine_name"])
            scientific = base.clean_text(r.get("generic_name", ""))
            active = base.clean_text(r.get("active_ingredient", ""))
            brands = base.clean_text(r.get("brand_names_display", ""))
            source_names = base.clean_text(r.get("source_medicine_names", ""))
            dosage = base.clean_text(r.get("dosage_form_group", ""))
            display = choose_best_scientific_name(scientific, source_names, analysis_name, dosage)
            if not display:
                display = scientific or source_names or analysis_name
            aliases = " | ".join([p for p in [display, scientific, active, brands, source_names, analysis_name, dosage] if p])
            items.append((display, analysis_name, aliases))
        return items
    def refresh_inventory_page(self, df):
        items = self.get_medicine_display_map(df)
        current = self.get_selected_medicine_name() if hasattr(self, "medicine_combo") else ""
        self.medicine_combo.blockSignals(True)
        self.medicine_combo.clear()
        cache = []
        for display, analysis_name, aliases in items:
            self.medicine_combo.addItem(display, analysis_name)
            blob = build_search_blob_from_values(display, analysis_name, aliases)
            generic_blob = build_search_blob_from_values(display, aliases)
            cache.append((display, analysis_name, blob, generic_blob, aliases))
        self._medicine_display_cache = cache
        self.medicine_completer_model.setStringList([d for d, _, _, _, _ in cache][:250])
        idx = -1
        if current:
            current_bundle = normalize_query_bundle(current)
            best_score = 0.0
            for pos, (_, analysis_name, blob, generic_blob, aliases) in enumerate(cache):
                score = max(query_match_score(current_bundle, generic_blob), query_match_score(current_bundle, blob))
                if score > best_score:
                    best_score = score
                    idx = pos
        if idx < 0 and self.medicine_combo.count() > 0:
            idx = 0
        if idx >= 0:
            self.medicine_combo.setCurrentIndex(idx)
        self.medicine_combo.blockSignals(False)
        self.refresh_inventory_details()
    def export_branch_compare(self):
        df = self.filtered_df()
        if df.empty:
            return
        name = self.get_selected_medicine_name()
        sub = df[df["medicine_name"] == name].copy()
        if sub.empty:
            return
        path, _ = QFileDialog.getSaveFileName(self, t(self.lang, "export_compare"), f"branch_compare_{slug(name)}.xlsx", "Excel Files (*.xlsx);;CSV Files (*.csv)")
        if not path:
            return
        if path.lower().endswith('.csv'):
            sub.to_csv(path, index=False, encoding='utf-8-sig')
        else:
            if not path.lower().endswith('.xlsx'):
                path += '.xlsx'
            sub.to_excel(path, index=False)
        QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "save_ok"))
    def refresh_inventory_page(self, df):
        items = self.get_medicine_display_map(df)
        current = self.get_selected_medicine_name() if hasattr(self, "medicine_combo") else ""
        self.medicine_combo.blockSignals(True)
        self.medicine_combo.clear()
        displays = [d for d, _ in items]
        for display, generic in items:
            self.medicine_combo.addItem(display, generic)
        self.medicine_completer_model.setStringList(displays)
        idx = -1
        if current:
            for pos, (_, generic) in enumerate(items):
                if normalize_search_text(current) == normalize_search_text(generic) or normalize_search_text(current) in normalize_search_text(generic):
                    idx = pos
                    break
        if idx >= 0:
            self.medicine_combo.setCurrentIndex(idx)
        self.medicine_combo.blockSignals(False)
        self.refresh_inventory_details()
    def refresh_inventory_details(self):
        df = self.filtered_df()
        if df.empty or self.medicine_combo.count() == 0:
            self.inventory_summary.setPlainText(t(self.lang, "no_data"))
            self.set_table_from_df(self.branch_compare_table, pd.DataFrame())
            return
        name = self.get_selected_medicine_name()
        sub = df[df["medicine_name"] == name].copy()
        if sub.empty:
            self.inventory_summary.setPlainText(t(self.lang, "no_data"))
            self.set_table_from_df(self.branch_compare_table, pd.DataFrame())
            return
        row = sub.sort_values(["escalation_level", "shortage_risk"], ascending=[True, False]).iloc[0]
        txt = (
            f"{t(self.lang, 'medicine_lookup')}: {row['medicine_name']}\n"
            f"{t(self.lang, 'summary_item_code')}: {row['item_code']}\n"
            f"{t(self.lang, 'summary_branch')}: {translate_display_value(row['branch_name'], self.lang)}\n"
            f"{t(self.lang, 'summary_current_stock')}: {row['current_stock']}\n"
            f"{t(self.lang, 'summary_dynamic_min')}: {row['dynamic_min_stock']}\n"
            f"{t(self.lang, 'summary_avg_daily')}: {row['avg_daily_consumption']:.2f}\n"
            f"{t(self.lang, 'summary_avg_monthly')}: {row['avg_monthly_consumption']:.2f}\n"
            f"{t(self.lang, 'summary_lead_time')}: {row['lead_time_days']}\n"
            f"{t(self.lang, 'summary_pending_po')}: {row['pending_po_qty']}\n"
            f"{t(self.lang, 'summary_shortage_risk')}: {translate_display_value(row['shortage_risk'], self.lang)}\n"
            f"{t(self.lang, 'summary_expiry_risk')}: {translate_display_value(row['expiry_risk'], self.lang)}\n"
            f"{t(self.lang, 'summary_escalation')}: {translate_display_value(row['escalation_level'], self.lang)}\n"
            f"{t(self.lang, 'summary_supplier')}: {translate_display_value(row['supplier_name'], self.lang)}\n"
            f"{t(self.lang, 'summary_on_time_rate')}: {row['supplier_on_time_rate']}\n"
            f"{t(self.lang, 'summary_reason')}: {row['reason_explanation']}\n"
            f"{t(self.lang, 'summary_recommended_action')}: {row['recommended_action']}"
        )
        self.inventory_summary.setPlainText(txt)
        compare_cols = [
            "branch_name", "medicine_name", "generic_name", "dosage_form_group", "source_medicine_names", "item_code", "current_stock", "dynamic_min_stock",
            "avg_monthly_consumption", "pending_po_qty", "days_of_stock_left", "shortage_risk", "overstock_risk", "stock_status", "predicted_stock_30d", "predicted_stock_60d", "predicted_stock_90d", "escalation_level", "recommended_action"
        ]
        cmp = sub[[c for c in compare_cols if c in sub.columns]].copy().sort_values(["escalation_level", "days_of_stock_left"])
        self.set_table_from_df(self.branch_compare_table, cmp)
        self.scen_lead.setValue(float(row["lead_time_days"]))
        self.scen_pending.setValue(float(row["pending_po_qty"]))
        self.scen_extra.setValue(0)
    def run_scenario(self):
        df = self.filtered_df()
        if df.empty or self.medicine_combo.count() == 0:
            self.scenario_result.setPlainText(t(self.lang, "no_data"))
            return
        name = self.get_selected_medicine_name()
        sub = df[df["medicine_name"] == name].copy()
        if sub.empty:
            self.scenario_result.setPlainText(t(self.lang, "no_data"))
            return
        row = sub.iloc[0].copy()
        row["lead_time_days"] = self.scen_lead.value()
        row["pending_po_qty"] = self.scen_pending.value()
        row["current_stock"] = row["current_stock"] + self.scen_extra.value()
        calc = compute_metrics(pd.DataFrame([row]))[[
            "medicine_name", "brand_names_display", "current_stock", "lead_time_days", "pending_po_qty",
            "days_of_stock_left", "expected_stock_at_lead_time", "shortage_risk",
            "recommended_reorder_qty", "escalation_level", "recommended_action"
        ]].iloc[0]
        self.scenario_result.setPlainText(
            f"{t(self.lang, 'scenario_result_title')} {calc['medicine_name']}\n"
            f"Current stock: {calc['current_stock']}\n"
            f"Lead time: {calc['lead_time_days']}\n"
            f"Pending PO: {calc['pending_po_qty']}\n"
            f"Coverage days: {calc['days_of_stock_left']}\n"
            f"Expected stock at lead time: {calc['expected_stock_at_lead_time']}\n"
            f"Shortage risk: {calc['shortage_risk']}\n"
            f"Recommended reorder: {calc['recommended_reorder_qty']}\n"
            f"Escalation: {calc['escalation_level']}\n"
            f"Action: {calc['recommended_action']}"
        )
    def refresh_purchase_view_only(self):
        self.refresh_purchase(self.filtered_df())
    def refresh_purchase(self, df):
        horizon = 1
        if hasattr(self, "purchase_horizon_combo"):
            try:
                horizon = int(self.purchase_horizon_combo.currentText())
            except Exception:
                horizon = 1
        working = df.copy()
        if not working.empty:
            working["selected_horizon_months"] = horizon
            available = working["current_stock"].astype(float).fillna(0) + working["pending_po_qty"].astype(float).fillna(0)
            monthly = working["avg_monthly_consumption"].astype(float).fillna(working["avg_daily_consumption"].astype(float).fillna(0) * 30)
            dyn = working["dynamic_min_stock"].astype(float).fillna(0)
            working["recommended_reorder_qty_horizon"] = np.maximum(np.ceil((monthly * horizon) + dyn - available), 0).astype(int)
        purchase = working[working["recommended_reorder_qty_horizon"] > 0][[
            "medicine_name", "generic_name", "dosage_form_group", "brand_names_display", "branch_name", "shortage_risk", "overstock_risk", "stock_status", "escalation_level",
            "recommended_reorder_qty_horizon", "clinical_priority", "reason_explanation", "predicted_stock_30d", "predicted_stock_60d", "predicted_stock_90d"
        ]].sort_values(["escalation_level", "recommended_reorder_qty_horizon"], ascending=[True, False])
        self.set_table_from_df(self.purchase_table, purchase)
        current = self.purchase_item_combo.currentText()
        self.purchase_item_combo.blockSignals(True)
        self.purchase_item_combo.clear()
        items = (working[working["recommended_reorder_qty_horizon"] > 0]["generic_name"].fillna(working["medicine_name"]) + " | " + working[working["recommended_reorder_qty_horizon"] > 0]["branch_name"]).tolist()
        self.purchase_item_combo.addItems(items)
        idx = self.purchase_item_combo.findText(current)
        self.purchase_item_combo.setCurrentIndex(max(idx, 0))
        self.purchase_item_combo.blockSignals(False)
        self.refresh_draft()
    def refresh_draft(self):
        df = self.filtered_df().copy()
        if self.purchase_item_combo.count() == 0:
            self.draft_edit.setPlainText("")
            return
        try:
            horizon = int(self.purchase_horizon_combo.currentText()) if hasattr(self, "purchase_horizon_combo") else 1
        except Exception:
            horizon = 1
        item = self.purchase_item_combo.currentText()
        if " | " not in item:
            self.draft_edit.setPlainText("")
            return
        med, branch = item.rsplit(" | ", 1)
        sub = df[(df["branch_name"] == branch) & ((df["generic_name"] == med) | (df["medicine_name"] == med))]
        if sub.empty:
            self.draft_edit.setPlainText("")
            return
        row = sub.iloc[0].copy()
        row["selected_horizon_months"] = horizon
        monthly = float(row.get("avg_monthly_consumption", 0) or 0)
        available = float(row.get("current_stock", 0) or 0) + float(row.get("pending_po_qty", 0) or 0)
        dyn = float(row.get("dynamic_min_stock", 0) or 0)
        row["recommended_reorder_qty_horizon"] = max(int(math.ceil((monthly * horizon) + dyn - available)), 0)
        self.draft_edit.setPlainText(draft_purchase_request(row))
    def save_draft(self):
        text = self.draft_edit.toPlainText()
        if not text.strip():
            return
        path, _ = QFileDialog.getSaveFileName(self, t(self.lang, "save_txt"), "", "Text files (*.txt)")
        if path:
            Path(path).write_text(text, encoding="utf-8")
            QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "save_ok"))
    def refresh_redistribution(self, df):
        if df.empty:
            self.set_table_from_df(self.redistribution_table, pd.DataFrame())
            return
        rows = []
        for med, g in df.groupby("medicine_name"):
            needy = g[g["current_stock"] < g["dynamic_min_stock"]]
            surplus = g[g["current_stock"] > (g["dynamic_min_stock"] * 1.4)]
            if needy.empty or surplus.empty:
                continue
            for _, n in needy.iterrows():
                best = surplus.assign(surplus_qty=surplus["current_stock"] - surplus["dynamic_min_stock"]).sort_values("surplus_qty", ascending=False).iloc[0]
                transferable = max(int(best["current_stock"] - best["dynamic_min_stock"]), 0)
                need_qty = max(int(n["dynamic_min_stock"] - n["current_stock"]), 0)
                if transferable > 0 and need_qty > 0:
                    rows.append({
                        "medicine_name": med,
                        "from_branch": best["branch_name"],
                        "to_branch": n["branch_name"],
                        "suggested_transfer_qty": min(transferable, need_qty),
                        "source_surplus": transferable,
                        "target_gap": need_qty,
                        "note": "Human approval required before redistribution.",
                    })
        self.set_table_from_df(self.redistribution_table, pd.DataFrame(rows))
    def refresh_fefo(self, df):
        if df.empty:
            self.set_table_from_df(self.fefo_table, pd.DataFrame())
            return
        fefo = df[[
            "medicine_name", "generic_name", "dosage_form_group", "branch_name", "expiry_date", "days_to_expiry",
            "expiry_risk", "current_stock", "avg_daily_consumption",
            "expected_unused_before_expiry", "recommended_action"
        ]].sort_values(["days_to_expiry", "expected_unused_before_expiry"])
        self.set_table_from_df(self.fefo_table, fefo)
    def refresh_abcven(self, df):
        if df.empty:
            self.set_table_from_df(self.abcven_table, pd.DataFrame())
            return
        table = df[[
            "medicine_name", "generic_name", "dosage_form_group", "branch_name", "unit_cost", "avg_monthly_consumption",
            "abc_class", "ven_class", "clinical_priority"
        ]].sort_values(["abc_class", "ven_class"])
        self.set_table_from_df(self.abcven_table, table)
        if HAS_MATPLOTLIB:
            abc = df["abc_class"].value_counts().reindex(["A", "B", "C"]).fillna(0)
            ven = df["ven_class"].astype(str).str.upper().value_counts().reindex(["V", "E", "N"]).fillna(0)
            self.chart_abc.plot_pie(abc.index.tolist(), abc.values.tolist(), t(self.lang, "chart_abc_title"))
            self.chart_ven.plot_pie(ven.index.tolist(), ven.values.tolist(), t(self.lang, "chart_ven_title"))
    def refresh_supplier(self, df):
        if df.empty:
            self.set_table_from_df(self.supplier_table, pd.DataFrame())
            return
        sup = df.groupby("supplier_name", dropna=False).agg(
            items=("generic_name", "count"),
            avg_on_time_rate=("supplier_on_time_rate", "mean"),
            critical_items=("clinical_priority", lambda s: int((s == "Critical").sum())),
            critical_shortages=("shortage_risk", lambda s: int((s == "Critical").sum()))
        ).reset_index().sort_values(["avg_on_time_rate", "critical_shortages"], ascending=[False, False])
        self.set_table_from_df(self.supplier_table, sup)
        if HAS_MATPLOTLIB and not sup.empty:
            top = sup.head(CHART_TOP_N)
            self.chart_supplier.plot_bar(top["supplier_name"].tolist(), top["avg_on_time_rate"].round(2).tolist(), t(self.lang, "chart_supplier_title"))
    def refresh_safety(self, df):
        if df.empty:
            self.set_table_from_df(self.safety_table, pd.DataFrame())
            return
        table = df[[
            "medicine_name", "generic_name", "dosage_form_group", "branch_name", "min_stock_level", "dynamic_min_stock",
            "current_stock", "avg_daily_consumption", "lead_time_days", "recommended_reorder_qty"
        ]].sort_values(["recommended_reorder_qty"], ascending=False)
        self.set_table_from_df(self.safety_table, table)
    def refresh_emergency(self, df):
        red = df[df["escalation_level"] == "Red"][[
            "medicine_name", "generic_name", "dosage_form_group", "branch_name", "clinical_priority", "current_stock",
            "dynamic_min_stock", "days_of_stock_left", "shortage_risk", "recommended_action"
        ]].sort_values(["clinical_priority", "days_of_stock_left"], ascending=[False, True])
        self.set_table_from_df(self.emergency_table, red)
    def refresh_quality(self, df):
        raw = self.df_raw.copy()
        missing = {
            col: int(raw[col].isna().sum()) if col in raw.columns else len(raw)
            for col in ALL_COLUMNS
        }
        dup = int(raw.duplicated(subset=["generic_name", "item_code", "branch_name"]).sum() if "generic_name" in raw.columns else raw.duplicated(subset=["medicine_name", "item_code", "branch_name"]).sum()) if not raw.empty else 0
        if self.lang == "ar":
            text = (
                f"{t(self.lang, 'data_quality')}\n\n"
                f"عدد الصفوف: {len(raw)}\n"
                f"السجلات المكررة: {dup}\n"
                f"صفوف الاستهلاك اليومي الصفري: {int((self.df['avg_daily_consumption'] <= 0).sum())}\n"
                f"الأكواد المولدة تلقائيًا: {int(self.df['item_code'].astype(str).str.startswith('AUTO-').sum())}\n"
                f"الصفوف بدون صلاحية: {int(self.df['expiry_date'].isna().sum())}\n"
                f"الصفوف التي تم استنتاج مخزونها من المتوسط الشهري: {int((self.df_raw['current_stock'].fillna(0) == self.df_raw['avg_monthly_consumption'].fillna(-999)).sum())}\n"
            )
        else:
            text = (
                f"{t(self.lang, 'data_quality')}\n\n"
                f"Rows: {len(raw)}\n"
                f"Duplicate records: {dup}\n"
                f"Rows with zero avg_daily_consumption: {int((self.df['avg_daily_consumption'] <= 0).sum())}\n"
                f"Rows with inferred item codes: {int(self.df['item_code'].astype(str).str.startswith('AUTO-').sum())}\n"
                f"Rows with missing expiry: {int(self.df['expiry_date'].isna().sum())}\n"
                f"Rows with stock inferred from monthly averages: {int((self.df_raw['current_stock'].fillna(0) == self.df_raw['avg_monthly_consumption'].fillna(-999)).sum())}\n"
            )
        self.quality_text.setPlainText(text)
        self.set_table_from_df(self.quality_table, pd.DataFrame(list(missing.items()), columns=["column", "missing_count"]))
    def refresh_governance(self):
        gov_text = (
            "Governance guardrails\n"
            "- Decision support only; no automatic purchasing.\n"
            "- Human review by Pharmacy Director / authorized staff.\n"
            "- Daily data refresh and timestamp review.\n"
            "- Audit trail for alerts and drafts.\n"
            "- Clinical priority must be considered, not consumption alone.\n\n"
            "Failure modes\n"
            "- Unsupported recommendation -> metric: unsupported output rate.\n"
            "- Wrong shortage prioritization -> metric: false negative rate for critical items.\n"
            "- Outdated inventory -> metric: data freshness score.\n"
        )
        if self.lang == "ar":
            gov_text = (
                "ضوابط الحوكمة\n"
                "- النظام للدعم واتخاذ القرار فقط ولا ينفذ الشراء تلقائيًا.\n"
                "- المراجعة البشرية مطلوبة من مدير الصيدلية أو المفوض.\n"
                "- مراجعة يومية لتحديث البيانات والتاريخ.\n"
                "- الاحتفاظ بسجل للتنبيهات والمسودات.\n"
                "- يجب مراعاة الأولوية السريرية وليس الاستهلاك فقط.\n\n"
                "أنماط الفشل\n"
                "- توصية غير مدعومة -> المقياس: معدل المخرجات غير المدعومة.\n"
                "- ترتيب خاطئ للنقص -> المقياس: معدل السلبية الكاذبة للأصناف الحرجة.\n"
                "- بيانات قديمة -> المقياس: درجة حداثة البيانات.\n"
            )
        self.gov_text.setPlainText(gov_text)
        if self.prompt_combo.count() == 0:
            self.prompt_combo.addItems(list(PROMPTS.keys()))
        self.refresh_prompt()
    def refresh_prompt(self):
        if self.prompt_combo.count() == 0:
            self.prompt_text.setPlainText("")
            return
        self.prompt_text.setPlainText(PROMPTS[self.prompt_combo.currentText()])
    # ---------- helpers ----------
    def set_table_from_df(self, table, df):
        table.clear()
        if df is None or df.empty:
            table.setColumnCount(0)
            table.setRowCount(0)
            return
        view = df.copy()
        raw_columns = [str(c) for c in view.columns]
        if len(view) > MAX_TABLE_ROWS:
            view = view.head(MAX_TABLE_ROWS).copy()
            self.status.showMessage(t(self.lang, "showing_rows").format(n=MAX_TABLE_ROWS), 5000)
        for col in view.columns:
            if pd.api.types.is_datetime64_any_dtype(view[col]):
                view[col] = view[col].dt.strftime("%Y-%m-%d")
        view = view.fillna("")
        view = localize_dataframe_for_display(view, self.lang)
        display_columns = [str(c) for c in view.columns]
        table.setUpdatesEnabled(False)
        table.setSortingEnabled(False)
        table.setColumnCount(len(display_columns))
        table.setHorizontalHeaderLabels(display_columns)
        table.setRowCount(len(view))
        values = view.astype(str).to_numpy()
        for i in range(values.shape[0]):
            for j in range(values.shape[1]):
                item = QTableWidgetItem(values[i, j])
                table.setItem(i, j, item)
        table.setUpdatesEnabled(True)
        table.resizeColumnsToContents()
        header = table.horizontalHeader()
        try:
            header.setSectionResizeMode(QHeaderView.Interactive)
        except Exception:
            pass
        header.setStretchLastSection(False)
        width_defaults = {
            'branch': 150, 'الفرع': 150,
            'medicine': 170, 'scientific': 170, 'generic': 150, 'الاسم': 170,
            'trade': 180, 'تجاري': 180,
            'source': 220, 'أسماء': 220,
            'item code': 110, 'كود': 110,
            'stock': 105, 'الرصيد': 105,
            'dynamic': 110, 'min': 110, 'الأدنى': 110,
            'daily': 105, 'يومي': 105,
            'monthly': 110, 'شهري': 110,
            'pending': 100, 'معلقة': 100,
            'lead': 95, 'توريد': 95,
            'days': 100, 'أيام': 100,
            'risk': 105, 'خطورة': 105,
            'escalation': 105, 'تصعيد': 105,
            'supplier': 140, 'المورد': 140,
            'action': 240, 'الإجراء': 240,
        }
        for idx, (raw_label, display_label) in enumerate(zip(raw_columns, display_columns)):
            raw_l = raw_label.lower()
            disp_l = display_label.lower()
            width = max(table.columnWidth(idx) + 18, 95)
            for token, token_width in width_defaults.items():
                if token in raw_l or token in disp_l:
                    width = max(width, token_width)
            if 'recommended_action' in raw_l or 'recommended action' in disp_l.lower() or 'الإجراء' in disp_l:
                width = max(width, 260)
            if table is getattr(self, 'branch_compare_table', None):
                width = min(width, 240)
            else:
                width = min(width, 300)
            table.setColumnWidth(idx, width)
        if table.columnCount() > 0:
            last = table.columnCount() - 1
            if table is getattr(self, 'branch_compare_table', None):
                table.setColumnWidth(last, max(table.columnWidth(last), 220))
            elif table is getattr(self, 'inventory_detail_table', None):
                table.setColumnWidth(last, max(table.columnWidth(last), 230))
    def clear_layout(self, layout):
        while layout.count():
            item = layout.takeAt(0)
            widget = item.widget()
            if widget:
                widget.deleteLater()
    def is_store_branch_name(self, name):
        s = clean_text(name).lower()
        store_tokens = [
            "central stock", "central store", "main store", "warehouse", "store", "stock",
            "مخزن", "المخزن", "المخزن الرئيسي", "المخزن المركزي", "مخزن رئيسي", "مركزي"
        ]
        branch_tokens = ["satellite", "clinic", "pharmacy", "صيدلية", "عيادة", "فرع"]
        if any(tok in s for tok in store_tokens):
            return True
        if any(tok in s for tok in branch_tokens):
            return False
        return False
    def build_purchase_df(self, df):
        if df is None or df.empty:
            return pd.DataFrame()
        return df[df["recommended_reorder_qty"] > 0][[
            "medicine_name", "generic_name", "dosage_form_group", "brand_names_display", "branch_name",
            "shortage_risk", "escalation_level", "recommended_reorder_qty", "clinical_priority", "reason_explanation"
        ]].sort_values(["escalation_level", "recommended_reorder_qty"], ascending=[True, False])
    def build_redistribution_df(self, df):
        if df is None or df.empty:
            return pd.DataFrame()
        rows = []
        for med, g in df.groupby("medicine_name"):
            needy = g[g["current_stock"] < g["dynamic_min_stock"]]
            surplus = g[g["current_stock"] > (g["dynamic_min_stock"] * 1.4)]
            if needy.empty or surplus.empty:
                continue
            for _, n in needy.iterrows():
                best = surplus.assign(surplus_qty=surplus["current_stock"] - surplus["dynamic_min_stock"]).sort_values("surplus_qty", ascending=False).iloc[0]
                transferable = max(int(best["current_stock"] - best["dynamic_min_stock"]), 0)
                need_qty = max(int(n["dynamic_min_stock"] - n["current_stock"]), 0)
                if transferable > 0 and need_qty > 0:
                    rows.append({
                        "medicine_name": med,
                        "generic_name": n.get("generic_name", ""),
                        "dosage_form_group": n.get("dosage_form_group", "Miscellaneous"),
                        "from_branch": best["branch_name"],
                        "to_branch": n["branch_name"],
                        "suggested_transfer_qty": min(transferable, need_qty),
                        "source_surplus": transferable,
                        "target_gap": need_qty,
                        "note": "Human approval required before redistribution.",
                    })
        return pd.DataFrame(rows)
    def build_fefo_df(self, df):
        if df is None or df.empty:
            return pd.DataFrame()
        return df[[
            "medicine_name", "generic_name", "dosage_form_group", "branch_name", "expiry_date", "days_to_expiry",
            "expiry_risk", "current_stock", "avg_daily_consumption", "expected_unused_before_expiry", "recommended_action"
        ]].sort_values(["days_to_expiry", "expected_unused_before_expiry"])
    def build_abcven_df(self, df):
        if df is None or df.empty:
            return pd.DataFrame()
        return df[[
            "medicine_name", "generic_name", "dosage_form_group", "branch_name", "unit_cost", "avg_monthly_consumption",
            "abc_class", "ven_class", "clinical_priority"
        ]].sort_values(["abc_class", "ven_class"])
    def build_supplier_df(self, df):
        if df is None or df.empty:
            return pd.DataFrame()
        return df.groupby("supplier_name", dropna=False).agg(
            items=("medicine_name", "count"),
            avg_on_time_rate=("supplier_on_time_rate", "mean"),
            critical_items=("clinical_priority", lambda s: int((s == "Critical").sum())),
            critical_shortages=("shortage_risk", lambda s: int((s == "Critical").sum()))
        ).reset_index().sort_values(["avg_on_time_rate", "critical_shortages"], ascending=[False, False])
    def build_emergency_df(self, df):
        if df is None or df.empty:
            return pd.DataFrame()
        return df[df["escalation_level"] == "Red"][[
            "medicine_name", "generic_name", "dosage_form_group", "branch_name", "clinical_priority", "current_stock",
            "dynamic_min_stock", "days_of_stock_left", "shortage_risk", "recommended_action"
        ]].sort_values(["days_of_stock_left", "current_stock"])
    def split_store_branch_views(self, df):
        if df is None or df.empty:
            return df.copy(), df.copy(), df.copy()
        working = df.copy()
        working["_is_store_branch"] = working["branch_name"].map(self.is_store_branch_name)
        combined = working.drop(columns=["_is_store_branch"]).copy()
        store_df = working[working["_is_store_branch"]].drop(columns=["_is_store_branch"]).copy()
        branch_df = working[~working["_is_store_branch"]].drop(columns=["_is_store_branch"]).copy()
        return combined, store_df, branch_df
    def build_store_branch_support(self, df):
        if df is None or df.empty or "medicine_name" not in df.columns:
            return pd.DataFrame()
        combined, store_df, branch_df = self.split_store_branch_views(df)
        if store_df.empty or branch_df.empty:
            return pd.DataFrame()
        store_by_generic = (
            store_df.groupby("medicine_name", dropna=False)
            .agg(
                store_branches=("branch_name", lambda s: " | ".join(sorted({clean_text(v) for v in s if pd.notna(v) and str(v).strip()}))),
                central_store_stock=("current_stock", "sum"),
                central_store_pending_po=("pending_po_qty", "sum"),
                central_store_days_cover=("days_of_stock_left", "max"),
                store_brand_names=("brand_names_display", lambda s: _join_unique(s, sep=" | ")),
            )
            .reset_index()
        )
        needs = branch_df[
            (branch_df["shortage_risk"].isin(["High", "Critical"])) |
            (branch_df["escalation_level"].isin(["Amber", "Red"])) |
            (pd.to_numeric(branch_df["recommended_reorder_qty"], errors="coerce").fillna(0) > 0)
        ].copy()
        if needs.empty:
            return pd.DataFrame()
        rep = needs.merge(store_by_generic, on="medicine_name", how="left")
        rep["central_store_stock"] = pd.to_numeric(rep["central_store_stock"], errors="coerce").fillna(0)
        rep["central_store_pending_po"] = pd.to_numeric(rep["central_store_pending_po"], errors="coerce").fillna(0)
        rep["store_can_support"] = np.where(
            rep["central_store_stock"] > rep["recommended_reorder_qty"], "Yes", "No"
        )
        rep["suggested_network_action"] = np.where(
            rep["central_store_stock"] > rep["recommended_reorder_qty"],
            "Check transfer from central store before urgent purchase",
            "Central store not sufficient; keep urgent purchase / supplier escalation",
        )
        cols = [
            "medicine_name", "generic_name", "dosage_form_group", "brand_names_display", "branch_name", "current_stock", "dynamic_min_stock",
            "days_of_stock_left", "shortage_risk", "escalation_level", "recommended_reorder_qty",
            "store_branches", "central_store_stock", "central_store_pending_po", "central_store_days_cover",
            "store_can_support", "suggested_network_action"
        ]
        for c in cols:
            if c not in rep.columns:
                rep[c] = ""
        return rep[cols].sort_values(
            ["store_can_support", "recommended_reorder_qty"],
            ascending=[True, False]
        )
    def form_filtered_df(self, df, form_key):
        if df is None or df.empty:
            return pd.DataFrame()
        if not form_key or form_key == "all":
            return df.copy()
        if form_key == "tablets":
            target = "Tablets"
        elif form_key == "ampoules":
            target = "Ampoules"
        elif form_key == "solutions":
            target = "Solutions"
        elif form_key == "miscellaneous":
            target = "Miscellaneous"
        else:
            target = form_key
        if "dosage_form_group" not in df.columns:
            return pd.DataFrame()
        mask = df["dosage_form_group"].fillna("").astype(str).str.strip().str.casefold() == str(target).casefold()
        return df.loc[mask].copy()
    def dosage_form_display_name(self, form_key):
        names = {
            "tablets": ("الأقراص", "Tablets"),
            "ampoules": ("الأمبولات", "Ampoules"),
            "solutions": ("المحاليل", "Solutions"),
            "miscellaneous": ("المتنوعات", "Miscellaneous"),
        }
        ar, en = names.get(form_key, (str(form_key), str(form_key)))
        return ar if self.lang == "ar" else en
    def _report_scope_dialog(self):
        if self.lang == "ar":
            label = "اختر نطاق التقرير النهائي"
            items = [
                "العرض الحالي حسب الفلاتر",
                "المخزن فقط",
                "الصيدليات الفرعية فقط",
                "تقرير الشبكة بالكامل (المخزن + الفروع)",
                "تقرير مفصل: مخزن + فروع + ربط بينهما",
                "تقرير الأشكال الصيدلانية منفصلة",
                "تقرير الأقراص فقط",
                "تقرير الأمبولات فقط",
                "تقرير المحاليل فقط",
                "تقرير المتنوعات فقط",
            ]
        else:
            label = "Choose final report scope"
            items = [
                "Current filtered view",
                "Central store only",
                "Branch pharmacies only",
                "Whole network report (store + branches)",
                "Detailed network report: store + branches + linkage",
                "Split report by dosage form",
                "Tablets only report",
                "Ampoules only report",
                "Solutions only report",
                "Miscellaneous only report",
            ]
        choice, ok = QInputDialog.getItem(self, t(self.lang, "export_report"), label, items, 0, False)
        if not ok or not choice:
            return None
        mapping = {
            items[0]: "current",
            items[1]: "store",
            items[2]: "branches",
            items[3]: "network",
            items[4]: "detailed",
            items[5]: "forms_split",
            items[6]: "tablets",
            items[7]: "ampoules",
            items[8]: "solutions",
            items[9]: "miscellaneous",
        }
        return mapping.get(choice, "current")
    def _report_section_html(self, title, df):
        if df is None or df.empty:
            empty = "لا توجد بيانات لهذا الجزء." if self.lang == "ar" else "No data for this section."
            return f"<h2>{title}</h2><p>{empty}</p>"
        view = localize_dataframe_for_display(df, self.lang)
        return f"<h2>{title}</h2>{view.to_html(index=False)}"
    def _scope_title(self, scope):
        if self.lang == "ar":
            return {
                "current": "العرض الحالي حسب الفلاتر",
                "store": "المخزن فقط",
                "branches": "الصيدليات الفرعية فقط",
                "network": "تقرير الشبكة بالكامل",
                "detailed": "تقرير الشبكة المفصل",
                "forms_split": "تقرير الأشكال الصيدلانية منفصلة",
                "tablets": "تقرير الأقراص فقط",
                "ampoules": "تقرير الأمبولات فقط",
                "solutions": "تقرير المحاليل فقط",
                "miscellaneous": "تقرير المتنوعات فقط",
            }.get(scope, "العرض الحالي حسب الفلاتر")
        return {
            "current": "Current filtered view",
            "store": "Central store only",
            "branches": "Branch pharmacies only",
            "network": "Whole network report",
            "detailed": "Detailed network report",
            "forms_split": "Split report by dosage form",
            "tablets": "Tablets only report",
            "ampoules": "Ampoules only report",
            "solutions": "Solutions only report",
            "miscellaneous": "Miscellaneous only report",
        }.get(scope, "Current filtered view")
    def _report_output_format_dialog(self):
        if self.lang == "ar":
            label = "اختر صيغة التصدير للتقرير النهائي"
            items = ["HTML فقط", "Excel فقط", "HTML وExcel معًا"]
        else:
            label = "Choose final report export format"
            items = ["HTML only", "Excel only", "HTML and Excel together"]
        choice, ok = QInputDialog.getItem(self, t(self.lang, "export_report"), label, items, 0, False)
        if not ok or not choice:
            return None
        mapping = {items[0]: "html", items[1]: "excel", items[2]: "both"}
        return mapping.get(choice, "html")
    def _save_report_base_path(self, output_format):
        if output_format == "html":
            default_name = "pharmaguard_final_report.html"
            filt = "HTML Files (*.html)"
        elif output_format == "excel":
            default_name = "pharmaguard_final_report.xlsx"
            filt = "Excel Files (*.xlsx)"
        else:
            default_name = "pharmaguard_final_report"
            filt = "Report Base Name (*.*)"
        path, _ = QFileDialog.getSaveFileName(self, t(self.lang, "export_report"), default_name, filt)
        if not path:
            return None
        p = Path(path)
        if output_format == "html" and p.suffix.lower() != ".html":
            p = p.with_suffix(".html")
        elif output_format == "excel" and p.suffix.lower() != ".xlsx":
            p = p.with_suffix(".xlsx")
        elif output_format == "both" and p.suffix.lower() in {".html", ".xlsx"}:
            p = p.with_suffix("")
        return p
    def _style_excel_worksheet(self, worksheet, df):
        if not HAS_OPENPYXL or worksheet is None:
            return
        worksheet.freeze_panes = "A2"
        header_fill = PatternFill(fill_type="solid", fgColor="DCE6F1") if PatternFill else None
        header_font = Font(bold=True) if Font else None
        wrap = Alignment(wrap_text=True, vertical="top", horizontal="center") if Alignment else None
        body_wrap = Alignment(wrap_text=True, vertical="top") if Alignment else None
        for row in worksheet.iter_rows():
            for cell in row:
                if row[0].row == 1:
                    if header_fill:
                        cell.fill = header_fill
                    if header_font:
                        cell.font = header_font
                    if wrap:
                        cell.alignment = wrap
                else:
                    if body_wrap:
                        cell.alignment = body_wrap
        max_col = min(worksheet.max_column, len(df.columns) if not df.empty else worksheet.max_column)
        for idx in range(1, max_col + 1):
            col_letter = get_column_letter(idx) if get_column_letter else None
            if not col_letter:
                continue
            header_val = worksheet[f"{col_letter}1"].value
            max_len = len(str(header_val)) if header_val is not None else 0
            sample_limit = min(worksheet.max_row, 400)
            for r in range(2, sample_limit + 1):
                val = worksheet[f"{col_letter}{r}"].value
                if val is None:
                    continue
                max_len = max(max_len, len(str(val)))
            worksheet.column_dimensions[col_letter].width = max(12, min(max_len + 4, 60))
        for row_idx in range(2, min(worksheet.max_row, 250) + 1):
            worksheet.row_dimensions[row_idx].height = 22
        worksheet.auto_filter.ref = worksheet.dimensions
    def _write_excel_sheets(self, path, sheet_map):
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for sheet_name, df_out in sheet_map.items():
                local_df = df_out.copy() if isinstance(df_out, pd.DataFrame) else pd.DataFrame()
                local_df = localize_dataframe_for_display(local_df, self.lang) if not local_df.empty else local_df
                if local_df.empty:
                    local_df = pd.DataFrame({t(self.lang, "notes") if "notes" in TEXT[self.lang] else "Notes": ["-"]})
                local_df.to_excel(writer, index=False, sheet_name=sheet_name[:31])
                ws = writer.book[sheet_name[:31]]
                self._style_excel_worksheet(ws, local_df)
    def _html_report_style(self):
        return f"""
        <style>
        body {{font-family: Arial, Tahoma, sans-serif; margin: 24px; direction: {'rtl' if self.lang == 'ar' else 'ltr'};}}
        h1,h2 {{color:#0f172a;}}
        .card {{border:1px solid #cbd5e1; border-radius:12px; padding:12px 16px; margin:8px 0; background:#f8fafc;}}
        table {{border-collapse: collapse; width:100%; margin-top:10px; table-layout:auto;}}
        th, td {{border:1px solid #cbd5e1; padding:8px; text-align:center; vertical-align:top; white-space:normal; word-break:break-word; max-width:420px;}}
        th {{background:#e2e8f0;}}
        .note {{margin-top:14px; color:#334155;}}
        </style>
        """
    def _report_excel_sections(self, scope, combined, store_df, branch_df, linkage_df, report_df, top_red, top_reorder, supplier):
        if self.lang == "ar":
            title_inventory = "ملخص التحليل"
            title_emergency = "أعلى حالات التصعيد"
            title_purchase = "أعلى أولويات طلبات الشراء"
            title_supplier = "ملخص الموردين"
            title_store = "تقرير المخزن"
            title_branches = "تقرير الصيدليات الفرعية"
            title_linkage = "ربط احتياجات الفروع بتغطية المخزن"
            title_combined = "تقرير الشبكة بالكامل"
        else:
            title_inventory = "Analysis summary"
            title_emergency = "Top escalation cases"
            title_purchase = "Top purchase priorities"
            title_supplier = "Supplier summary"
            title_store = "Central store report"
            title_branches = "Branch pharmacies report"
            title_linkage = "Store support for branch needs"
            title_combined = "Whole network report"
        if scope == "current":
            return {title_inventory: report_df.head(40), title_emergency: top_red, title_purchase: top_reorder, title_supplier: supplier}
        if scope == "store":
            return {title_store: store_df.head(120), title_supplier: supplier}
        if scope == "branches":
            return {title_branches: branch_df.head(120), title_emergency: top_red, title_purchase: top_reorder}
        if scope == "network":
            return {title_combined: combined.head(120), title_store: store_df.head(120), title_branches: branch_df.head(120), title_linkage: linkage_df.head(120), title_supplier: supplier}
        if scope == "detailed":
            return {title_combined: combined.head(200), title_store: store_df.head(200), title_branches: branch_df.head(200), title_linkage: linkage_df.head(200), title_emergency: top_red, title_purchase: top_reorder, title_supplier: supplier}
        if scope == "forms_split":
            return {self.dosage_form_display_name("tablets"): self.form_filtered_df(combined, "tablets"), self.dosage_form_display_name("ampoules"): self.form_filtered_df(combined, "ampoules"), self.dosage_form_display_name("solutions"): self.form_filtered_df(combined, "solutions"), self.dosage_form_display_name("miscellaneous"): self.form_filtered_df(combined, "miscellaneous")}
        title = self.dosage_form_display_name(scope) if scope in {"tablets", "ampoules", "solutions", "miscellaneous"} else title_inventory
        return {title: report_df.head(160), title_emergency: top_red, title_purchase: top_reorder, title_supplier: supplier}
    def export_final_report(self):
        if self.df.empty:
            return
        base_df = self.filtered_df().copy()
        if base_df.empty:
            base_df = self.df.copy()
        scope = self._report_scope_dialog()
        if not scope:
            return
        output_format = self._report_output_format_dialog()
        if not output_format:
            return
        base_path = self._save_report_base_path(output_format)
        if not base_path:
            return
        combined, store_df, branch_df = self.split_store_branch_views(base_df)
        linkage_df = self.build_store_branch_support(base_df)
        if scope == "store":
            report_df = store_df
        elif scope == "branches":
            report_df = branch_df
        elif scope in {"tablets", "ampoules", "solutions", "miscellaneous"}:
            report_df = self.form_filtered_df(combined, scope)
        else:
            report_df = combined
        if report_df.empty and scope != "forms_split":
            report_df = combined
        top_red = report_df[report_df["escalation_level"] == "Red"].head(20) if not report_df.empty else pd.DataFrame()
        top_reorder = report_df.sort_values("recommended_reorder_qty", ascending=False).head(20) if not report_df.empty else pd.DataFrame()
        supplier = report_df.groupby("supplier_name").agg(items=("generic_name", "count"), avg_on_time_rate=("supplier_on_time_rate", "mean")).reset_index() if not report_df.empty else pd.DataFrame()
        active_branch = self.branch_filter.currentText()
        active_branch = "" if active_branch in ("", "All", t(self.lang, "all")) else active_branch
        active_search = self.search_edit.text().strip()
        if self.lang == "ar":
            scope_title = self._scope_title(scope)
            filters_html = f"""
            <div class='card'><b>نطاق التقرير:</b> {scope_title}<br>
            <b>فلتر الفرع الحالي:</b> {active_branch or 'الكل'}<br>
            <b>بحث المستخدم:</b> {active_search or 'بدون'}<br>
            <b>ملاحظة تشغيلية:</b> الصيدليات الفرعية تُخدم من المخزن؛ لذلك تم فصل التحليل إلى المخزن، الفروع، وربط الدعم من المخزن إلى الفروع عند الحاجة. كما يمكن فصل التقارير حسب الشكل الصيدلاني لأن كل لوح/مخزن قد يتعامل مع الأقراص أو الأمبولات أو المحاليل بشكل مستقل.</div>
            """
            note = "تم إنشاء هذا التقرير من البرنامج مع اعتماد التحليل على الاسم العلمي + الشكل الصيدلاني، وإظهار الأسماء التجارية للتوضيح فقط."
        else:
            scope_title = self._scope_title(scope)
            filters_html = f"""
            <div class='card'><b>Report scope:</b> {scope_title}<br>
            <b>Active branch filter:</b> {active_branch or 'All'}<br>
            <b>User search:</b> {active_search or 'None'}<br>
            <b>Operational note:</b> Branch pharmacies are supplied by the central store; therefore this report separates central store, branches, and store-to-branch support linkage. Dosage-form specific reports are also available because real-world store boards often separate tablets, ampoules, solutions, and miscellaneous items.</div>
            """
            note = "This report was generated from the application using scientific name + dosage-form analysis, with trade names displayed for clarification only."
        excel_sections = self._report_excel_sections(scope, combined, store_df, branch_df, linkage_df, report_df, top_red, top_reorder, supplier)
        total_items = len(report_df) if scope != "forms_split" else len(combined)
        critical_count = int((report_df['shortage_risk'] == 'Critical').sum()) if scope != "forms_split" and not report_df.empty else int((combined['shortage_risk'] == 'Critical').sum())
        expiry_count = int(report_df['expiry_risk'].isin(['High','Critical','Expired']).sum()) if scope != "forms_split" and not report_df.empty else int(combined['expiry_risk'].isin(['High','Critical','Expired']).sum())
        reorder_count = int(pd.to_numeric(report_df['recommended_reorder_qty'], errors='coerce').fillna(0).sum()) if scope != "forms_split" and not report_df.empty else int(pd.to_numeric(combined['recommended_reorder_qty'], errors='coerce').fillna(0).sum())
        red_count = int((report_df['escalation_level'] == 'Red').sum()) if scope != "forms_split" and not report_df.empty else int((combined['escalation_level'] == 'Red').sum())
        if output_format in {"html", "both"}:
            sections = []
            for section_title, section_df in excel_sections.items():
                sections.append(self._report_section_html(section_title, section_df))
            summary_html = f"""
            <html><head><meta charset='utf-8'>{self._html_report_style()}</head><body>
            <h1>{t(self.lang, 'title')}</h1>
            {filters_html}
            <div class='card'><b>{t(self.lang, 'kpi_total')}:</b> {total_items}<br>
            <b>{t(self.lang, 'kpi_critical')}:</b> {critical_count}<br>
            <b>{t(self.lang, 'kpi_expiry')}:</b> {expiry_count}<br>
            <b>{t(self.lang, 'kpi_reorder')}:</b> {reorder_count}<br>
            <b>{t(self.lang, 'kpi_red')}:</b> {red_count}</div>
            {''.join(sections)}
            <p class='note'>{note}</p>
            <p style='margin-top:20px;color:#475569'>{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
            </body></html>
            """
            html_path = base_path.with_suffix('.html') if output_format == 'both' else base_path
            Path(html_path).write_text(summary_html, encoding='utf-8')
        if output_format in {"excel", "both"}:
            summary_sheet_name = "ملخص" if self.lang == "ar" else "Summary"
            summary_df = pd.DataFrame([
                {
                    ("نطاق التقرير" if self.lang == "ar" else "Report scope"): self._scope_title(scope),
                    t(self.lang, 'kpi_total'): total_items,
                    t(self.lang, 'kpi_critical'): critical_count,
                    t(self.lang, 'kpi_expiry'): expiry_count,
                    t(self.lang, 'kpi_reorder'): reorder_count,
                    t(self.lang, 'kpi_red'): red_count,
                    ("فلتر الفرع" if self.lang == "ar" else "Branch filter"): active_branch or ("الكل" if self.lang == "ar" else "All"),
                    ("نص البحث" if self.lang == "ar" else "Search text"): active_search or ("بدون" if self.lang == "ar" else "None"),
                }
            ])
            excel_path = base_path.with_suffix('.xlsx') if output_format == 'both' else base_path
            excel_map = {summary_sheet_name: summary_df}
            excel_map.update(excel_sections)
            self._write_excel_sheets(excel_path, excel_map)
        QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "final_report_saved"))
    def _analysis_export_dialog(self):
        if self.lang == "ar":
            label = "اختر نوع التحليل أو مجموعة الجداول المطلوب تصديرها"
            items = [
                "الكل (ملف شامل)",
                "العرض الحالي حسب الفلاتر",
                "المخزن فقط",
                "الصيدليات الفرعية فقط",
                "مقارنة الفروع للصنف المحدد",
                "إعادة التوزيع",
                "FEFO والصلاحية",
                "ABC-VEN",
                "طلبات الشراء",
                "التصعيد العاجل",
                "الموردون",
                "ربط المخزن بالفروع",
                "الأشكال الصيدلانية منفصلة",
                "الأقراص فقط",
                "الأمبولات فقط",
                "المحاليل فقط",
                "المتنوعات فقط",
            ]
        else:
            label = "Choose which analysis export you want"
            items = [
                "All (full workbook)",
                "Current filtered view",
                "Central store only",
                "Branch pharmacies only",
                "Branch comparison for selected item",
                "Redistribution",
                "FEFO and expiry",
                "ABC-VEN",
                "Purchase requests",
                "Emergency escalation",
                "Supplier summary",
                "Store-branch linkage",
                "Split by dosage form",
                "Tablets only",
                "Ampoules only",
                "Solutions only",
                "Miscellaneous only",
            ]
        choice, ok = QInputDialog.getItem(self, t(self.lang, "export_workbook"), label, items, 0, False)
        if not ok or not choice:
            return None
        mapping = {
            items[0]: "all", items[1]: "current", items[2]: "store", items[3]: "branches", items[4]: "compare",
            items[5]: "redistribution", items[6]: "fefo", items[7]: "abcven", items[8]: "purchase",
            items[9]: "emergency", items[10]: "supplier", items[11]: "linkage", items[12]: "forms_split",
            items[13]: "tablets", items[14]: "ampoules", items[15]: "solutions", items[16]: "miscellaneous",
        }
        return mapping.get(choice, "all")
    def export_workbook(self):
        if self.df.empty:
            return
        export_mode = self._analysis_export_dialog()
        if not export_mode:
            return
        filtered = self.filtered_df()
        if filtered.empty:
            filtered = self.df.copy()
        combined, store_df, branch_df = self.split_store_branch_views(self.df)
        linkage_df = self.build_store_branch_support(self.df)
        default_name = f"pharmaguard_{export_mode}_analysis"
        path, _ = QFileDialog.getSaveFileName(self, t(self.lang, "export_workbook"), default_name + ".xlsx", "Excel Files (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        def current_compare_df():
            name = self.get_selected_medicine_name()
            if not name:
                return pd.DataFrame()
            return filtered[filtered["medicine_name"] == name].copy()
        forms_map = {
            "Tablets": self.form_filtered_df(combined, "tablets"),
            "Ampoules": self.form_filtered_df(combined, "ampoules"),
            "Solutions": self.form_filtered_df(combined, "solutions"),
            "Miscellaneous": self.form_filtered_df(combined, "miscellaneous"),
        }
        data_map = {
            "current": {"Filtered_View": filtered},
            "store": {"Central_Store": store_df},
            "branches": {"Branch_Pharmacies": branch_df},
            "compare": {"Branch_Comparison": current_compare_df()},
            "redistribution": {"Redistribution": self.build_redistribution_df(filtered)},
            "fefo": {"FEFO_Expiry": self.build_fefo_df(filtered)},
            "abcven": {"ABC_VEN": self.build_abcven_df(filtered)},
            "purchase": {"Purchase_Requests": self.build_purchase_df(filtered)},
            "emergency": {"Emergency_Escalation": self.build_emergency_df(filtered)},
            "supplier": {"Supplier_Summary": self.build_supplier_df(filtered)},
            "linkage": {"Store_Branch_Link": linkage_df},
            "forms_split": {
                "Tablets": forms_map["Tablets"],
                "Ampoules": forms_map["Ampoules"],
                "Solutions": forms_map["Solutions"],
                "Miscellaneous": forms_map["Miscellaneous"],
            },
            "tablets": {"Tablets": forms_map["Tablets"]},
            "ampoules": {"Ampoules": forms_map["Ampoules"]},
            "solutions": {"Solutions": forms_map["Solutions"]},
            "miscellaneous": {"Miscellaneous": forms_map["Miscellaneous"]},
            "all": {
                "Analyzed_Inventory": self.df,
                "Filtered_View": filtered,
                "Network_All": combined,
                "Central_Store": store_df,
                "Branch_Pharmacies": branch_df,
                "Store_Branch_Link": linkage_df,
                "Branch_Comparison": current_compare_df(),
                "Redistribution": self.build_redistribution_df(filtered),
                "FEFO_Expiry": self.build_fefo_df(filtered),
                "ABC_VEN": self.build_abcven_df(filtered),
                "Purchase_Requests": self.build_purchase_df(filtered),
                "Emergency_Escalation": self.build_emergency_df(filtered),
                "Supplier_Summary": self.build_supplier_df(filtered),
                "Tablets": forms_map["Tablets"],
                "Ampoules": forms_map["Ampoules"],
                "Solutions": forms_map["Solutions"],
                "Miscellaneous": forms_map["Miscellaneous"],
                "Import_Report": pd.DataFrame({"mapping_report": self.mapping_report_text.splitlines()}),
            },
        }
        selected_map = data_map.get(export_mode, data_map["all"])
        self._write_excel_sheets(path, selected_map)
        QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "save_ok"))
def make_icon_files(app_dir):
    png_path = app_dir / "pharmaguard_icon.png"
    ico_path = app_dir / "pharmaguard_icon.ico"
    if png_path.exists() and ico_path.exists():
        return
    try:
        from PIL import Image, ImageDraw
        img = Image.new("RGBA", (256, 256), (37, 99, 235, 255))
        d = ImageDraw.Draw(img)
        d.rounded_rectangle((24, 24, 232, 232), radius=40, fill=(15, 23, 42, 255))
        d.ellipse((56, 56, 200, 200), fill=(59, 130, 246, 255))
        d.rectangle((110, 72, 146, 184), fill=(255, 255, 255, 255))
        d.rectangle((74, 108, 182, 144), fill=(255, 255, 255, 255))
        img.save(png_path)
        img.save(ico_path)
    except Exception:
        pass
def load_icon():
    app_dir = Path(__file__).resolve().parent
    make_icon_files(app_dir)
    ico_path = app_dir / "pharmaguard_icon.ico"
    png_path = app_dir / "pharmaguard_icon.png"
    if ico_path.exists():
        return QIcon(str(ico_path))
    if png_path.exists():
        return QIcon(str(png_path))
    return QIcon()
def show_splash(app, icon_path):
    if not icon_path.exists():
        return None
    pix = QPixmap(str(icon_path))
    if pix.isNull():
        return None
    splash = QSplashScreen(pix.scaled(240, 240, Qt.KeepAspectRatio, Qt.SmoothTransformation))
    splash.showMessage("Loading PharmaGuard AI Pro V6.8..." if QApplication.layoutDirection() != Qt.RightToLeft else "جارٍ تحميل فارماجارد برو 6.8...", Qt.AlignBottom | Qt.AlignCenter, Qt.white)
    splash.show()
    app.processEvents()
    return splash
# Alias current module as base so legacy code keeps working
import sys as _sys
base = _sys.modules[__name__]
# === Inventory search layer ===
import sys
from pathlib import Path
from datetime import datetime
import numpy as np
import pandas as pd
from PySide6.QtCore import Qt, QTimer
from PySide6.QtGui import QIcon
from PySide6.QtWidgets import QApplication, QFileDialog, QInputDialog, QMessageBox
APP_VERSION = "V6.9 Consumption & Predictions Arabic"
APP_TITLE_EN = f"PharmaGuard AI Pro {APP_VERSION}"
APP_TITLE_AR = f"فارماجارد برو {APP_VERSION}"
import re
# ----------------- UI / Arabic / dark-mode patches -----------------
# Stronger dark-mode contrast so side panels and selected items are clearly visible.
base.DARK_QSS = base.DARK_QSS + """
QDockWidget { background: #0b1220; border-left: 1px solid #334155; border-right: 1px solid #334155; }
QDockWidget > QWidget { background: #0b1220; }
QListWidget::item { padding: 8px 10px; border-radius: 8px; margin: 2px; }
QListWidget::item:selected { background: #2563eb; color: white; }
QTableWidget::item:selected { background: #1d4ed8; color: white; }
QComboBox QAbstractItemView { background: #0f172a; color: #e5e7eb; selection-background-color: #2563eb; selection-color: white; border: 1px solid #334155; }
QLineEdit:focus, QComboBox:focus, QPlainTextEdit:focus, QTextEdit:focus, QTableWidget:focus, QListWidget:focus { border: 2px solid #60a5fa; }
QFrame#Card { background: #172033; border: 1px solid #4b5d7a; border-radius: 14px; }
QGroupBox { background: #111827; }
QHeaderView::section { background: #334155; color: #f8fafc; font-weight: 700; }
"""
base.LIGHT_QSS = base.LIGHT_QSS + """
QListWidget::item { padding: 8px 10px; border-radius: 8px; margin: 2px; }
QListWidget::item:selected { background: #2563eb; color: white; }
QTableWidget::item:selected { background: #dbeafe; color: #0f172a; }
QComboBox QAbstractItemView { selection-background-color: #2563eb; selection-color: white; }
QLineEdit:focus, QComboBox:focus, QPlainTextEdit:focus, QTextEdit:focus, QTableWidget:focus, QListWidget:focus { border: 2px solid #3b82f6; }
"""
_AR_REPLACEMENTS = {
    "أقراص": "Tablets",
    "اقراص": "Tablets",
    "أمبولات": "Ampoules",
    "امبولات": "Ampoules",
    "أمبول": "Ampoule",
    "امبول": "Ampoule",
    "محاليل": "Solutions",
    "محلول": "Solution",
    "متنوعات": "Miscellaneous",
    "متنوع": "Miscellaneous",
    "فيال": "Vial",
    "زجاجة": "Bottle",
}
def _chart_safe_label(value, max_len=28):
    s = base.clean_text(value)
    if not s:
        return ""
    # If Arabic shaping libraries are available, use them.
    if base.HAS_ARABIC_SHAPING and base._contains_arabic(s):
        s = base.ui_text(s)
    elif base._contains_arabic(s):
        # Fallback: replace common Arabic dosage-form words with English so charts stay readable.
        for ar_word, en_word in _AR_REPLACEMENTS.items():
            s = s.replace(ar_word, en_word)
        # Remove remaining Arabic if it still breaks chart labels.
        s = re.sub(r'[\u0600-\u06FF]+', '', s).strip()
        s = re.sub(r'\s+', ' ', s).strip()
    if len(s) > max_len:
        s = s[:max_len - 1] + "…"
    return s
def _normalize_inventory_query(value):
    s = base.clean_text(value)
    if not s:
        return ""
    s = s.lower()
    trans = str.maketrans({
        "أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا",
        "ى": "ي", "ئ": "ي", "ؤ": "و",
        "ة": "ه", "ـ": "",
    })
    s = s.translate(trans)
    s = re.sub(r'[ً-ٰٟۖ-ۭ]', '', s)
    s = re.sub(r'[^\w\s\-/\[\]()]+', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s
class BetterMatplotlibChart(base.QWidget):
    def __init__(self, title=""):
        super().__init__()
        self.dark_mode = False
        self.layout = base.QVBoxLayout(self)
        self.layout.setContentsMargins(6, 6, 6, 6)
        self.layout.setSpacing(6)
        self.title = base.QLabel(title)
        self.title.setWordWrap(True)
        self.title.setStyleSheet("font-weight:700; font-size:14px; padding: 2px 4px;")
        self.layout.addWidget(self.title)
        if base.HAS_MATPLOTLIB:
            self.figure = base.Figure(figsize=(5, 3), dpi=100)
            self.canvas = base.FigureCanvas(self.figure)
            self.layout.addWidget(self.canvas)
            self.set_dark_mode(False)
        else:
            self.figure = None
            self.canvas = base.QLabel(base.t("ar" if base.QApplication.layoutDirection() == base.Qt.RightToLeft else "en", "matplotlib_missing"))
            self.layout.addWidget(self.canvas)
    def set_dark_mode(self, dark_mode=False):
        self.dark_mode = bool(dark_mode)
        if not base.HAS_MATPLOTLIB:
            return
        fig_bg = "#111827" if self.dark_mode else "#ffffff"
        axes_bg = "#111827" if self.dark_mode else "#ffffff"
        text_color = "#e5e7eb" if self.dark_mode else "#111827"
        self.figure.set_facecolor(fig_bg)
        if hasattr(self, "title"):
            self.title.setStyleSheet(
                f"font-weight:700; font-size:14px; padding: 2px 4px; color:{text_color};"
            )
    def _style_axes(self, ax):
        if not base.HAS_MATPLOTLIB:
            return
        fig_bg = "#111827" if self.dark_mode else "#ffffff"
        axes_bg = "#111827" if self.dark_mode else "#ffffff"
        text_color = "#e5e7eb" if self.dark_mode else "#111827"
        grid_color = "#334155" if self.dark_mode else "#e5e7eb"
        self.figure.set_facecolor(fig_bg)
        ax.set_facecolor(axes_bg)
        ax.tick_params(axis='x', colors=text_color, labelsize=8)
        ax.tick_params(axis='y', colors=text_color, labelsize=8)
        for spine in ax.spines.values():
            spine.set_color(grid_color)
        ax.yaxis.label.set_color(text_color)
        ax.xaxis.label.set_color(text_color)
        ax.title.set_color(text_color)
    def plot_bar(self, labels, values, title=""):
        if not base.HAS_MATPLOTLIB:
            return
        self.title.setText(title)  # Qt label renders Arabic correctly.
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        safe_labels = [_chart_safe_label(x) for x in labels]
        positions = list(range(len(safe_labels)))
        bar_color = "#3b82f6" if self.dark_mode else "#2563eb"
        ax.bar(positions, values, color=bar_color)
        ax.set_xticks(positions)
        ax.set_xticklabels(safe_labels, rotation=25, ha='right', fontsize=8, fontname=base.chart_font_family())
        # Avoid Arabic title inside matplotlib itself; the QLabel above is enough.
        self._style_axes(ax)
        self.figure.subplots_adjust(left=0.10, right=0.98, top=0.94, bottom=0.30)
        self.canvas.draw_idle()
    def plot_pie(self, labels, values, title=""):
        if not base.HAS_MATPLOTLIB:
            return
        self.title.setText(title)
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        safe_labels = [_chart_safe_label(x, max_len=24) for x in labels]
        text_color = "#e5e7eb" if self.dark_mode else "#111827"
        wedges, texts, autotexts = ax.pie(values, labels=safe_labels, autopct='%1.0f%%')
        self._style_axes(ax)
        for txt in list(texts) + list(autotexts):
            txt.set_fontname(base.chart_font_family())
            txt.set_color(text_color)
        self.figure.subplots_adjust(left=0.06, right=0.94, top=0.94, bottom=0.08)
        self.canvas.draw_idle()
base.MatplotlibChart = BetterMatplotlibChart
_base_apply_theme = base.PharmaGuardMainWindow.apply_theme
def _patched_apply_theme(self):
    _base_apply_theme(self)
    for chart in self.findChildren(base.MatplotlibChart):
        try:
            chart.set_dark_mode(self.dark_mode)
        except Exception:
            pass
base.PharmaGuardMainWindow.apply_theme = _patched_apply_theme
# Make reorder chart labels cleaner: scientific name + dosage form, not raw mixed Arabic text.
_base_refresh_dashboard = base.PharmaGuardMainWindow.refresh_dashboard
def _patched_refresh_dashboard(self, df):
    _base_refresh_dashboard(self, df)
    if df is None or df.empty or not base.HAS_MATPLOTLIB:
        return
    try:
        top = df.sort_values("recommended_reorder_qty", ascending=False).head(base.CHART_TOP_N).copy()
        if "generic_name" in top.columns:
            label_series = top["generic_name"].fillna(top["medicine_name"]).astype(str)
        else:
            label_series = top["medicine_name"].astype(str)
        if "dosage_form_group" in top.columns:
            label_series = label_series + " - " + top["dosage_form_group"].fillna("").astype(str)
        self.chart_reorder.plot_bar(label_series.tolist(), top["recommended_reorder_qty"].tolist(), base.t(self.lang, "chart_reorder_title"))
        risk_counts = df["shortage_risk"].value_counts().reindex(["Low", "Moderate", "High", "Critical"]).fillna(0)
        risk_labels = [
            {"ar": "منخفض", "en": "Low"}[self.lang],
            {"ar": "متوسط", "en": "Moderate"}[self.lang],
            {"ar": "عالٍ", "en": "High"}[self.lang],
            {"ar": "حرج", "en": "Critical"}[self.lang],
        ] if self.lang == "ar" else ["Low", "Moderate", "High", "Critical"]
        self.chart_risk.plot_pie(risk_labels, risk_counts.values.tolist(), base.t(self.lang, "chart_risk_title"))
    except Exception:
        pass
base.PharmaGuardMainWindow.refresh_dashboard = _patched_refresh_dashboard
# ----------------- global patches on imported base module -----------------
if "active_ingredient" not in base.OPTIONAL_COLUMNS:
    base.OPTIONAL_COLUMNS = list(base.OPTIONAL_COLUMNS) + ["active_ingredient"]
base.ALL_COLUMNS = base.CORE_COLUMNS + base.OPTIONAL_COLUMNS
base.COLUMN_ALIASES["active_ingredient"] = [
    "active_ingredient", "active ingredient", "ingredient", "active", "المادة الفعالة",
    "الماده الفعاله", "الاسم العلمي", "scientific_name", "generic", "generic_name"
]
base.COLUMN_LABELS["active_ingredient"] = {"en": "Active ingredient", "ar": "المادة الفعالة"}
base.COLUMN_LABELS["predicted_next_month_demand"] = {"en": "Predicted next-month demand", "ar": "التوقع للشهر القادم"}
base.COLUMN_LABELS["predicted_3_month_demand"] = {"en": "Predicted 3-month demand", "ar": "التوقع لثلاثة أشهر"}
base.COLUMN_LABELS["predicted_end_month_stock"] = {"en": "Predicted end-month stock", "ar": "الرصيد المتوقع نهاية الشهر"}
base.COLUMN_LABELS["predicted_gap_next_month"] = {"en": "Predicted next-month gap", "ar": "فجوة الشهر القادم"}
base.COLUMN_LABELS["predicted_gap_3_months"] = {"en": "Predicted 3-month gap", "ar": "فجوة 3 أشهر"}
base.COLUMN_LABELS["coverage_months"] = {"en": "Coverage (months)", "ar": "التغطية بالشهور"}
base.COLUMN_LABELS["branch_count"] = {"en": "Branch count", "ar": "عدد الصيدليات"}
base.COLUMN_LABELS["highest_shortage_risk"] = {"en": "Highest shortage risk", "ar": "أعلى خطورة نقص"}
base.COLUMN_LABELS["highest_escalation_level"] = {"en": "Highest escalation", "ar": "أعلى تصعيد"}
base.COLUMN_LABELS["network_monthly_consumption"] = {"en": "Network monthly consumption", "ar": "الاستهلاك الشهري للشبكة"}
base.COLUMN_LABELS["branch_monthly_consumption"] = {"en": "Branch monthly consumption", "ar": "الاستهلاك الشهري للصيدلية"}
base.COLUMN_LABELS["central_store_stock"] = {"en": "Central store stock", "ar": "رصيد المخزن"}
base.COLUMN_LABELS["central_store_pending_po"] = {"en": "Central store pending PO", "ar": "المعلق للمخزن"}
base.COLUMN_LABELS["central_store_available"] = {"en": "Central store available", "ar": "المتاح بالمخزن"}
base.COLUMN_LABELS["central_store_coverage_months"] = {"en": "Central store coverage (months)", "ar": "تغطية المخزن بالشهور"}
base.COLUMN_LABELS["store_support_status"] = {"en": "Store support status", "ar": "قدرة المخزن على التغطية"}
base.COLUMN_LABELS["predicted_next_month_issue"] = {"en": "Predicted next-month issue", "ar": "الصرف المتوقع الشهر القادم"}
base.COLUMN_LABELS["predicted_3_month_issue"] = {"en": "Predicted 3-month issue", "ar": "الصرف المتوقع لثلاثة أشهر"}
base.COLUMN_LABELS["store_predicted_gap_next_month"] = {"en": "Store next-month gap", "ar": "فجوة المخزن الشهر القادم"}
base.COLUMN_LABELS["store_predicted_gap_3_months"] = {"en": "Store 3-month gap", "ar": "فجوة المخزن 3 أشهر"}
_orig_normalize = base.normalize_dataframe
_orig_aggregate = base.aggregate_by_generic_branch
def _clean_active(val):
    if pd.isna(val):
        return ""
    s = base.clean_text(val)
    return s
def normalize_dataframe_v69(df):
    out = _orig_normalize(df)
    if "active_ingredient" not in out.columns:
        out["active_ingredient"] = ""
    out["active_ingredient"] = out["active_ingredient"].map(_clean_active)
    fallback = out.get("generic_name", pd.Series([""] * len(out), index=out.index)).fillna("").astype(str)
    out["active_ingredient"] = out["active_ingredient"].where(out["active_ingredient"].astype(str).str.strip() != "", fallback)
    out["active_ingredient"] = out["active_ingredient"].where(out["active_ingredient"].astype(str).str.strip() != "", out["medicine_name"].astype(str))
    # ensure no empties
    out["active_ingredient"] = out["active_ingredient"].fillna(out["generic_name"]).fillna(out["medicine_name"]).astype(str)
    for col in ["avg_monthly_consumption", "avg_daily_consumption", "current_stock", "pending_po_qty"]:
        out[col] = pd.to_numeric(out[col], errors="coerce").fillna(0)
    return out[base.ALL_COLUMNS]
def aggregate_by_generic_branch_v69(df):
    if df is None or len(df) == 0:
        return pd.DataFrame(columns=base.ALL_COLUMNS)
    working = base.normalize_dataframe(df)
    grouped = []
    for (analysis_name, branch_name), g in working.groupby(["medicine_name", "branch_name"], dropna=False):
        grouped.append({
            "medicine_name": analysis_name,
            "analysis_name": analysis_name,
            "active_ingredient": base._pick_first_nonempty(g.get("active_ingredient", pd.Series(dtype=str)), default=base._pick_first_nonempty(g.get("generic_name", pd.Series(dtype=str)), default=analysis_name)),
            "generic_name": base._pick_first_nonempty(g.get("generic_name", pd.Series(dtype=str)), default=analysis_name),
            "dosage_form_group": base._pick_first_nonempty(g.get("dosage_form_group", pd.Series(dtype=str)), default="Miscellaneous"),
            "brand_name": base._pick_first_nonempty(g.get("brand_name", pd.Series(dtype=str))),
            "brand_names_display": base._join_unique(g.get("brand_name", pd.Series(dtype=str))),
            "source_medicine_names": base._join_unique(g.get("source_medicine_names", g["medicine_name"]), sep=" | "),
            "item_code": base._join_unique(g["item_code"], sep=" | "),
            "current_stock": pd.to_numeric(g["current_stock"], errors="coerce").fillna(0).sum(),
            "min_stock_level": pd.to_numeric(g["min_stock_level"], errors="coerce").fillna(0).sum(),
            "avg_daily_consumption": pd.to_numeric(g["avg_daily_consumption"], errors="coerce").fillna(0).sum(),
            "avg_monthly_consumption": pd.to_numeric(g["avg_monthly_consumption"], errors="coerce").fillna(0).sum(),
            "lead_time_days": pd.to_numeric(g["lead_time_days"], errors="coerce").fillna(14).max(),
            "pending_po_qty": pd.to_numeric(g["pending_po_qty"], errors="coerce").fillna(0).sum(),
            "expiry_date": pd.to_datetime(g["expiry_date"], errors="coerce").min(),
            "clinical_priority": base._pick_max_priority(g["clinical_priority"]),
            "storage_location": base._join_unique(g["storage_location"], sep=" | "),
            "branch_name": branch_name,
            "unit_cost": pd.to_numeric(g["unit_cost"], errors="coerce").fillna(0).mean(),
            "supplier_name": base._join_unique(g["supplier_name"], sep=" | "),
            "supplier_on_time_rate": pd.to_numeric(g["supplier_on_time_rate"], errors="coerce").fillna(0.75).mean(),
            "ven_class": base._pick_first_nonempty(g["ven_class"], default="E") or "E",
            "abc_class": base._pick_first_nonempty(g.get("abc_class", pd.Series(dtype=str)), default=""),
        })
    out = pd.DataFrame(grouped)
    for col in base.ALL_COLUMNS:
        if col not in out.columns:
            out[col] = None
    return out[base.ALL_COLUMNS]
base.normalize_dataframe = normalize_dataframe_v69
base.aggregate_by_generic_branch = aggregate_by_generic_branch_v69
def add_prediction_columns(df, stock_col="current_stock", pending_col="pending_po_qty", monthly_col="avg_monthly_consumption"):
    out = df.copy()
    stock = pd.to_numeric(out.get(stock_col, 0), errors="coerce").fillna(0)
    pending = pd.to_numeric(out.get(pending_col, 0), errors="coerce").fillna(0)
    monthly = pd.to_numeric(out.get(monthly_col, 0), errors="coerce").fillna(0)
    available = stock + pending
    out["predicted_next_month_demand"] = np.ceil(monthly).astype(float)
    out["predicted_3_month_demand"] = np.ceil(monthly * 3).astype(float)
    out["predicted_end_month_stock"] = np.round(available - out["predicted_next_month_demand"], 1)
    out["predicted_gap_next_month"] = np.round(np.maximum(out["predicted_next_month_demand"] - available, 0), 1)
    out["predicted_gap_3_months"] = np.round(np.maximum(out["predicted_3_month_demand"] - available, 0), 1)
    with np.errstate(divide='ignore', invalid='ignore'):
        out["coverage_months"] = np.where(monthly > 0, np.round(available / monthly, 2), np.inf)
    return out
class InventorySearchMainWindow(base.PharmaGuardMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_TITLE_AR if self.lang == "ar" else APP_TITLE_EN)
        self.statusBar().showMessage("جاهز" if self.lang == "ar" else "Ready")
        self._medicine_search_timer = QTimer(self)
        self._medicine_search_timer.setSingleShot(True)
        self._medicine_search_timer.timeout.connect(self.apply_medicine_text_search)
        self._rebind_inventory_search_behavior()
    def _rebind_inventory_search_behavior(self):
        if not hasattr(self, "medicine_combo"):
            return
        try:
            self.medicine_combo.currentTextChanged.disconnect()
        except Exception:
            pass
        try:
            self.medicine_combo.currentTextChanged.disconnect(self.on_medicine_text_changed)
        except Exception:
            pass
        line = self.medicine_combo.lineEdit()
        if line is not None:
            try:
                line.textEdited.disconnect()
            except Exception:
                pass
            try:
                line.returnPressed.disconnect()
            except Exception:
                pass
            line.textEdited.connect(self.on_medicine_text_edited)
            line.returnPressed.connect(self.apply_medicine_text_search)
        try:
            self.medicine_completer.activated.disconnect()
        except Exception:
            pass
        self.medicine_completer.activated.connect(self.apply_medicine_text_search)
    def on_medicine_text_edited(self, text):
        if not hasattr(self, "medicine_combo") or self.medicine_combo.count() == 0:
            return
        query_bundle = normalize_query_bundle(text)
        cache = getattr(self, '_medicine_display_cache', [])
        if not any(query_bundle.values()):
            if cache:
                full_list = [d for d, _, _, _, _ in cache][:250]
                self.medicine_completer_model.setStringList(full_list)
            return
        ranked = []
        for display, generic, blob, generic_blob, aliases in cache:
            score = max(query_match_score(query_bundle, blob), query_match_score(query_bundle, generic_blob))
            if score > 0:
                ranked.append((score, display))
        ranked.sort(key=lambda x: (-x[0], x[1]))
        results = [d for _, d in ranked[:150]]
        if not results and cache:
            results = [d for d, _, _, _, _ in cache[:120]]
        self.medicine_completer_model.setStringList(results)
        try:
            self.medicine_completer.complete()
        except Exception:
            pass
        self._medicine_search_timer.start(120)
    def on_medicine_text_changed(self, text):
        # compatibility slot: route to debounced editor logic without forcing heavy refresh on each keystroke
        self.on_medicine_text_edited(text)
    def apply_medicine_text_search(self, *args):
        if not hasattr(self, "medicine_combo") or self.medicine_combo.count() == 0:
            return
        raw_text = self.medicine_combo.currentText()
        bundle = normalize_query_bundle(raw_text)
        if not any(bundle.values()):
            if self.medicine_combo.currentIndex() >= 0:
                self.refresh_inventory_details()
            return
        cache = getattr(self, '_medicine_display_cache', [])
        chosen_idx = -1
        best_score = 0.0
        for idx, (display, generic, blob, generic_blob, aliases) in enumerate(cache):
            score = max(query_match_score(bundle, generic_blob), query_match_score(bundle, blob))
            if score > best_score:
                best_score = score
                chosen_idx = idx
        if chosen_idx >= 0:
            self.medicine_combo.blockSignals(True)
            self.medicine_combo.setCurrentIndex(chosen_idx)
            self.medicine_combo.blockSignals(False)
            line = self.medicine_combo.lineEdit()
            if line is not None:
                line.setText(raw_text)
                line.setCursorPosition(len(raw_text))
            self.refresh_inventory_details()
    def get_selected_medicine_name(self):
        if not hasattr(self, "medicine_combo") or self.medicine_combo.count() == 0:
            return ""
        idx = self.medicine_combo.currentIndex()
        if idx >= 0:
            data = self.medicine_combo.itemData(idx)
            if data:
                return base.clean_text(data)
        raw = base.clean_text(self.medicine_combo.currentText())
        bundle = normalize_query_bundle(raw)
        cache = getattr(self, '_medicine_display_cache', [])
        if not any(bundle.values()) and cache:
            return base.clean_text(cache[0][1])
        for _, generic, blob, generic_blob, aliases in cache:
            if query_in_blob(bundle, generic_blob) or query_in_blob(bundle, blob):
                return base.clean_text(generic)
        qtext = bundle.get('text', '')
        if qtext:
            for _, generic, blob, generic_blob, aliases in cache:
                if blob.startswith(qtext) or generic_blob.startswith(qtext):
                    return base.clean_text(generic)
        if cache:
            return base.clean_text(cache[0][1])
        return ""
    def get_medicine_display_map(self, df):
        if df.empty:
            return []
        grouped = (
            df.groupby("medicine_name", dropna=False)
              .agg(
                  generic_name=("generic_name", lambda s: base._join_unique(s, sep=" | ")),
                  active_ingredient=("active_ingredient", lambda s: base._join_unique(s, sep=" | ")),
                  brand_names_display=("brand_names_display", lambda s: base._join_unique(s, sep=" | ")),
                  source_medicine_names=("source_medicine_names", lambda s: base._join_unique(s, sep=" | ")),
                  dosage_form_group=("dosage_form_group", lambda s: base._join_unique(s, sep=" | ")),
              )
              .reset_index()
        )
        items = []
        grouped = grouped.sort_values(["generic_name", "medicine_name"], kind="stable")
        for _, r in grouped.iterrows():
            medicine_name = base.clean_text(r["medicine_name"])
            scientific = base.clean_text(r.get("generic_name", ""))
            active = base.clean_text(r.get("active_ingredient", ""))
            brands = base.clean_text(r.get("brand_names_display", ""))
            source_names = base.clean_text(r.get("source_medicine_names", ""))
            dosage = base.clean_text(r.get("dosage_form_group", ""))
            # Prefer imported scientific/generic name in the visible list.
            display = scientific or source_names or medicine_name
            aliases = " | ".join([p for p in [medicine_name, scientific, active, brands, source_names, dosage] if p])
            items.append((display, medicine_name, aliases))
        return items
    def refresh_inventory_page(self, df):
        items = self.get_medicine_display_map(df)
        current = self.get_selected_medicine_name() if hasattr(self, "medicine_combo") else ""
        current_typed = base.clean_text(self.medicine_combo.currentText()) if hasattr(self, "medicine_combo") else ""
        self.medicine_combo.blockSignals(True)
        self.medicine_combo.clear()
        cache = []
        for display, medicine_name, item_aliases in items:
            self.medicine_combo.addItem(display, medicine_name)
            sub = df[df["medicine_name"] == medicine_name].copy()
            aliases = " | ".join(filter(None, [
                item_aliases,
                base._join_unique(sub.get("brand_names_display", base.pd.Series(dtype=str)), sep=" | "),
                base._join_unique(sub.get("source_medicine_names", base.pd.Series(dtype=str)), sep=" | "),
                base._join_unique(sub.get("generic_name", base.pd.Series(dtype=str)), sep=" | "),
                base._join_unique(sub.get("active_ingredient", base.pd.Series(dtype=str)), sep=" | "),
            ]))
            blob = build_search_blob_from_values(display, medicine_name, aliases)
            generic_blob = build_search_blob_from_values(medicine_name, aliases)
            cache.append((display, medicine_name, blob, generic_blob, aliases))
        self._medicine_display_cache = cache
        self._medicine_master_cache = list(cache)
        self.medicine_completer_model.setStringList([d for d, _, _, _, _ in cache][:250])
        idx = -1
        preferred_query = current_typed or current
        if preferred_query:
            current_bundle = normalize_query_bundle(preferred_query)
            for pos, (_, generic, blob, generic_blob, aliases) in enumerate(cache):
                if query_in_blob(current_bundle, generic_blob) or query_in_blob(current_bundle, blob):
                    idx = pos
                    break
        if idx >= 0:
            self.medicine_combo.setCurrentIndex(idx)
        elif self.medicine_combo.count() > 0:
            self.medicine_combo.setCurrentIndex(0)
        self.medicine_combo.blockSignals(False)
        self._rebind_inventory_search_behavior()
        line = self.medicine_combo.lineEdit()
        if line is not None:
            line.setPlaceholderText(t(self.lang, "type_to_search"))
        self.refresh_inventory_details()
    # ---- prediction / consumption builders ----
    def build_generic_monthly_consumption_df(self, df):
        working = df.copy()
        if working.empty:
            return pd.DataFrame()
        grp = working.groupby(["active_ingredient", "generic_name", "dosage_form_group", "medicine_name"], dropna=False)
        out = grp.agg(
            brand_names_display=("brand_names_display", lambda s: base._join_unique(s, sep=" | ")),
            branch_count=("branch_name", lambda s: len({base.clean_text(v) for v in s if pd.notna(v) and str(v).strip()})),
            network_monthly_consumption=("avg_monthly_consumption", "sum"),
            avg_daily_consumption=("avg_daily_consumption", "sum"),
            current_stock=("current_stock", "sum"),
            pending_po_qty=("pending_po_qty", "sum"),
            highest_shortage_risk=("shortage_risk", lambda s: max([str(v) for v in s if pd.notna(v)] or ["Low"], key=lambda x: {"Low":1,"Moderate":2,"High":3,"Critical":4}.get(x,1))),
            highest_escalation_level=("escalation_level", lambda s: max([str(v) for v in s if pd.notna(v)] or ["Green"], key=lambda x: {"Green":1,"Amber":2,"Red":3}.get(x,1))),
            recommended_action=("recommended_action", lambda s: base._pick_first_nonempty(s, default="")),
        ).reset_index()
        out = out.rename(columns={"network_monthly_consumption": "avg_monthly_consumption"})
        out = add_prediction_columns(out, monthly_col="avg_monthly_consumption")
        out = out.rename(columns={"avg_monthly_consumption": "network_monthly_consumption"})
        return out.sort_values(["predicted_gap_next_month", "network_monthly_consumption"], ascending=[False, False])
    def build_branch_monthly_consumption_df(self, df):
        working = df.copy()
        if working.empty:
            return pd.DataFrame()
        _, _, branch_df = self.split_store_branch_views(working)
        if branch_df.empty:
            branch_df = working.copy()
        grp = branch_df.groupby(["branch_name", "active_ingredient", "generic_name", "dosage_form_group", "medicine_name"], dropna=False)
        out = grp.agg(
            brand_names_display=("brand_names_display", lambda s: base._join_unique(s, sep=" | ")),
            branch_monthly_consumption=("avg_monthly_consumption", "sum"),
            avg_daily_consumption=("avg_daily_consumption", "sum"),
            current_stock=("current_stock", "sum"),
            pending_po_qty=("pending_po_qty", "sum"),
            shortage_risk=("shortage_risk", lambda s: max([str(v) for v in s if pd.notna(v)] or ["Low"], key=lambda x: {"Low":1,"Moderate":2,"High":3,"Critical":4}.get(x,1))),
            escalation_level=("escalation_level", lambda s: max([str(v) for v in s if pd.notna(v)] or ["Green"], key=lambda x: {"Green":1,"Amber":2,"Red":3}.get(x,1))),
            recommended_action=("recommended_action", lambda s: base._pick_first_nonempty(s, default="")),
        ).reset_index()
        out = out.rename(columns={"branch_monthly_consumption": "avg_monthly_consumption"})
        out = add_prediction_columns(out, monthly_col="avg_monthly_consumption")
        out = out.rename(columns={"avg_monthly_consumption": "branch_monthly_consumption"})
        return out.sort_values(["predicted_gap_next_month", "branch_monthly_consumption"], ascending=[False, False])
    def build_central_store_issue_df(self, df):
        working = df.copy()
        if working.empty:
            return pd.DataFrame()
        _, store_df, branch_df = self.split_store_branch_views(working)
        if branch_df.empty:
            return pd.DataFrame()
        branch_need = branch_df.groupby(["active_ingredient", "generic_name", "dosage_form_group", "medicine_name"], dropna=False).agg(
            branch_count=("branch_name", lambda s: len({base.clean_text(v) for v in s if pd.notna(v) and str(v).strip()})),
            network_monthly_consumption=("avg_monthly_consumption", "sum"),
            branch_brand_names=("brand_names_display", lambda s: base._join_unique(s, sep=" | ")),
        ).reset_index()
        if store_df.empty:
            out = branch_need.copy()
            out["store_branches"] = ""
            out["central_store_stock"] = 0
            out["central_store_pending_po"] = 0
        else:
            store_sum = store_df.groupby(["active_ingredient", "generic_name", "dosage_form_group", "medicine_name"], dropna=False).agg(
                store_branches=("branch_name", lambda s: base._join_unique(s, sep=" | ")),
                central_store_stock=("current_stock", "sum"),
                central_store_pending_po=("pending_po_qty", "sum"),
            ).reset_index()
            out = branch_need.merge(store_sum, on=["active_ingredient", "generic_name", "dosage_form_group", "medicine_name"], how="left")
            out["store_branches"] = out["store_branches"].fillna("")
            out["central_store_stock"] = pd.to_numeric(out["central_store_stock"], errors="coerce").fillna(0)
            out["central_store_pending_po"] = pd.to_numeric(out["central_store_pending_po"], errors="coerce").fillna(0)
        out["central_store_available"] = out["central_store_stock"] + out["central_store_pending_po"]
        out["predicted_next_month_issue"] = np.ceil(pd.to_numeric(out["network_monthly_consumption"], errors="coerce").fillna(0))
        out["predicted_3_month_issue"] = np.ceil(out["predicted_next_month_issue"] * 3)
        with np.errstate(divide='ignore', invalid='ignore'):
            out["central_store_coverage_months"] = np.where(out["predicted_next_month_issue"] > 0, np.round(out["central_store_available"] / out["predicted_next_month_issue"], 2), np.inf)
        out["store_predicted_gap_next_month"] = np.round(np.maximum(out["predicted_next_month_issue"] - out["central_store_available"], 0), 1)
        out["store_predicted_gap_3_months"] = np.round(np.maximum(out["predicted_3_month_issue"] - out["central_store_available"], 0), 1)
        out["store_support_status"] = np.where(out["store_predicted_gap_next_month"] > 0, "Urgent gap", np.where(out["store_predicted_gap_3_months"] > 0, "Short-term gap", "Covered"))
        out["suggested_network_action"] = np.where(
            out["store_predicted_gap_next_month"] > 0,
            "Urgent purchase / supplier escalation needed",
            np.where(out["store_predicted_gap_3_months"] > 0, "Monitor next purchase cycle and branch transfers", "Central store can support current branch demand")
        )
        return out.sort_values(["store_predicted_gap_next_month", "predicted_next_month_issue"], ascending=[False, False])
    def build_network_prediction_df(self, df):
        generic_df = self.build_generic_monthly_consumption_df(df)
        store_df = self.build_central_store_issue_df(df)
        if generic_df.empty and store_df.empty:
            return pd.DataFrame()
        if store_df.empty:
            out = generic_df.copy()
            out["central_store_available"] = 0
            out["central_store_coverage_months"] = 0
            out["store_predicted_gap_next_month"] = out.get("predicted_gap_next_month", 0)
            out["store_predicted_gap_3_months"] = out.get("predicted_gap_3_months", 0)
            out["store_support_status"] = np.where(out["predicted_gap_next_month"] > 0, "Urgent gap", "Covered")
            return out
        merge_cols = ["active_ingredient", "generic_name", "dosage_form_group", "medicine_name"]
        out = generic_df.merge(
            store_df[[*merge_cols, "central_store_stock", "central_store_pending_po", "central_store_available", "central_store_coverage_months", "store_predicted_gap_next_month", "store_predicted_gap_3_months", "store_support_status", "suggested_network_action"]],
            on=merge_cols,
            how="left"
        )
        out["central_store_stock"] = pd.to_numeric(out.get("central_store_stock", 0), errors="coerce").fillna(0)
        out["central_store_pending_po"] = pd.to_numeric(out.get("central_store_pending_po", 0), errors="coerce").fillna(0)
        out["central_store_available"] = pd.to_numeric(out.get("central_store_available", 0), errors="coerce").fillna(0)
        out["central_store_coverage_months"] = out.get("central_store_coverage_months", 0)
        out["store_predicted_gap_next_month"] = pd.to_numeric(out.get("store_predicted_gap_next_month", 0), errors="coerce").fillna(0)
        out["store_predicted_gap_3_months"] = pd.to_numeric(out.get("store_predicted_gap_3_months", 0), errors="coerce").fillna(0)
        out["store_support_status"] = out.get("store_support_status", "")
        return out.sort_values(["store_predicted_gap_next_month", "predicted_gap_next_month", "network_monthly_consumption"], ascending=[False, False, False])
    def _analysis_export_dialog(self):
        if self.lang == "ar":
            label = "اختر نوع التحليل أو مجموعة الجداول المطلوب تصديرها"
            items = [
                "الكل (ملف شامل)",
                "العرض الحالي حسب الفلاتر",
                "المخزن فقط",
                "الصيدليات الفرعية فقط",
                "مقارنة الفروع للصنف المحدد",
                "إعادة التوزيع",
                "FEFO والصلاحية",
                "ABC-VEN",
                "طلبات الشراء",
                "التصعيد العاجل",
                "الموردون",
                "ربط المخزن بالفروع",
                "الأشكال الصيدلانية منفصلة",
                "الأقراص فقط",
                "الأمبولات فقط",
                "المحاليل فقط",
                "المتنوعات فقط",
                "الاستهلاك الشهري حسب المادة الفعالة + الاسم العلمي + الشكل",
                "الاستهلاك الشهري لكل صيدلية",
                "صرف/تغطية المخزن المركزي للصيدليات",
                "توقعات الاحتياج حسب المادة الفعالة والاسم العلمي والشكل",
                "توقعات الاحتياج لكل صيدلية",
                "توقعات تغطية المخزن المركزي",
                "تقرير شامل: الاستهلاك والتوقعات",
            ]
        else:
            label = "Choose which analysis export you want"
            items = [
                "All (full workbook)", "Current filtered view", "Central store only", "Branch pharmacies only",
                "Branch comparison for selected item", "Redistribution", "FEFO and expiry", "ABC-VEN",
                "Purchase requests", "Emergency escalation", "Supplier summary", "Store-branch linkage",
                "Split by dosage form", "Tablets only", "Ampoules only", "Solutions only", "Miscellaneous only",
                "Monthly consumption by active ingredient + scientific name + dosage form",
                "Monthly consumption by branch pharmacy",
                "Central store issue / coverage for branches",
                "Need forecast by active ingredient + scientific name + dosage form",
                "Need forecast for each branch pharmacy",
                "Central store coverage forecast",
                "Consumption and predictions package",
            ]
        choice, ok = QInputDialog.getItem(self, base.t(self.lang, "export_workbook"), label, items, 0, False)
        if not ok or not choice:
            return None
        mapping = {
            items[0]: "all", items[1]: "current", items[2]: "store", items[3]: "branches", items[4]: "compare",
            items[5]: "redistribution", items[6]: "fefo", items[7]: "abcven", items[8]: "purchase",
            items[9]: "emergency", items[10]: "supplier", items[11]: "linkage", items[12]: "forms_split",
            items[13]: "tablets", items[14]: "ampoules", items[15]: "solutions", items[16]: "miscellaneous",
            items[17]: "monthly_generic", items[18]: "monthly_branch", items[19]: "store_issue",
            items[20]: "forecast_generic", items[21]: "forecast_branch", items[22]: "forecast_store", items[23]: "forecast_bundle",
        }
        return mapping.get(choice, "all")
    def _report_scope_dialog(self):
        if self.lang == "ar":
            label = "اختر نطاق التقرير النهائي"
            items = [
                "العرض الحالي حسب الفلاتر", "المخزن فقط", "الصيدليات الفرعية فقط",
                "تقرير الشبكة بالكامل (المخزن + الفروع)", "تقرير مفصل: مخزن + فروع + ربط بينهما",
                "تقرير الأشكال الصيدلانية منفصلة", "تقرير الأقراص فقط", "تقرير الأمبولات فقط",
                "تقرير المحاليل فقط", "تقرير المتنوعات فقط", "تقرير الاستهلاك والتوقعات"
            ]
        else:
            label = "Choose final report scope"
            items = [
                "Current filtered view", "Central store only", "Branch pharmacies only", "Whole network report (store + branches)",
                "Detailed network report: store + branches + linkage", "Split report by dosage form", "Tablets only report",
                "Ampoules only report", "Solutions only report", "Miscellaneous only report", "Consumption and predictions report"
            ]
        choice, ok = QInputDialog.getItem(self, base.t(self.lang, "export_report"), label, items, 0, False)
        if not ok or not choice:
            return None
        mapping = {
            items[0]: "current", items[1]: "store", items[2]: "branches", items[3]: "network", items[4]: "detailed",
            items[5]: "forms_split", items[6]: "tablets", items[7]: "ampoules", items[8]: "solutions", items[9]: "miscellaneous",
            items[10]: "consumption_predictions",
        }
        return mapping.get(choice, "current")
    def _scope_title(self, scope):
        base_map = super()._scope_title(scope)
        if scope != "consumption_predictions":
            return base_map
        return "تقرير الاستهلاك والتوقعات" if self.lang == "ar" else "Consumption and predictions report"
    def _build_prediction_export_map(self, source_df):
        generic = self.build_generic_monthly_consumption_df(source_df)
        branch = self.build_branch_monthly_consumption_df(source_df)
        store_issue = self.build_central_store_issue_df(source_df)
        network = self.build_network_prediction_df(source_df)
        return {
            "Monthly_By_Generic_Form": generic,
            "Monthly_By_Pharmacy": branch,
            "Central_Store_Issue_Coverage": store_issue,
            "Forecast_Generic_Form": generic[[c for c in [
                "active_ingredient", "generic_name", "dosage_form_group", "medicine_name", "brand_names_display", "network_monthly_consumption",
                "predicted_next_month_demand", "predicted_3_month_demand", "predicted_end_month_stock", "predicted_gap_next_month", "predicted_gap_3_months", "coverage_months",
                "highest_shortage_risk", "highest_escalation_level", "recommended_action"
            ] if c in generic.columns]].copy() if not generic.empty else pd.DataFrame(),
            "Forecast_By_Pharmacy": branch[[c for c in [
                "branch_name", "active_ingredient", "generic_name", "dosage_form_group", "medicine_name", "brand_names_display", "branch_monthly_consumption",
                "predicted_next_month_demand", "predicted_3_month_demand", "predicted_end_month_stock", "predicted_gap_next_month", "predicted_gap_3_months", "coverage_months",
                "shortage_risk", "escalation_level", "recommended_action"
            ] if c in branch.columns]].copy() if not branch.empty else pd.DataFrame(),
            "Forecast_Central_Store": store_issue[[c for c in [
                "active_ingredient", "generic_name", "dosage_form_group", "medicine_name", "branch_count", "network_monthly_consumption",
                "central_store_stock", "central_store_pending_po", "central_store_available", "predicted_next_month_issue", "predicted_3_month_issue",
                "central_store_coverage_months", "store_predicted_gap_next_month", "store_predicted_gap_3_months", "store_support_status", "suggested_network_action"
            ] if c in store_issue.columns]].copy() if not store_issue.empty else pd.DataFrame(),
            "Forecast_Network_Summary": network,
        }
    def export_workbook(self):
        if self.df.empty:
            return
        export_mode = self._analysis_export_dialog()
        if not export_mode:
            return
        filtered = self.filtered_df()
        if filtered.empty:
            filtered = self.df.copy()
        combined, store_df, branch_df = self.split_store_branch_views(self.df)
        linkage_df = self.build_store_branch_support(self.df)
        default_name = f"pharmaguard_{export_mode}_analysis"
        path, _ = QFileDialog.getSaveFileName(self, base.t(self.lang, "export_workbook"), default_name + ".xlsx", "Excel Files (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        def current_compare_df():
            name = self.get_selected_medicine_name()
            if not name:
                return pd.DataFrame()
            return filtered[filtered["medicine_name"] == name].copy()
        forms_map = {
            "Tablets": self.form_filtered_df(combined, "tablets"),
            "Ampoules": self.form_filtered_df(combined, "ampoules"),
            "Solutions": self.form_filtered_df(combined, "solutions"),
            "Miscellaneous": self.form_filtered_df(combined, "miscellaneous"),
        }
        pred_map = self._build_prediction_export_map(filtered)
        data_map = {
            "current": {"Filtered_View": filtered},
            "store": {"Central_Store": store_df},
            "branches": {"Branch_Pharmacies": branch_df},
            "compare": {"Branch_Comparison": current_compare_df()},
            "redistribution": {"Redistribution": self.build_redistribution_df(filtered)},
            "fefo": {"FEFO_Expiry": self.build_fefo_df(filtered)},
            "abcven": {"ABC_VEN": self.build_abcven_df(filtered)},
            "purchase": {"Purchase_Requests": self.build_purchase_df(filtered)},
            "emergency": {"Emergency_Escalation": self.build_emergency_df(filtered)},
            "supplier": {"Supplier_Summary": self.build_supplier_df(filtered)},
            "linkage": {"Store_Branch_Link": linkage_df},
            "forms_split": {"Tablets": forms_map["Tablets"], "Ampoules": forms_map["Ampoules"], "Solutions": forms_map["Solutions"], "Miscellaneous": forms_map["Miscellaneous"]},
            "tablets": {"Tablets": forms_map["Tablets"]},
            "ampoules": {"Ampoules": forms_map["Ampoules"]},
            "solutions": {"Solutions": forms_map["Solutions"]},
            "miscellaneous": {"Miscellaneous": forms_map["Miscellaneous"]},
            "monthly_generic": {"Monthly_By_Generic_Form": pred_map["Monthly_By_Generic_Form"]},
            "monthly_branch": {"Monthly_By_Pharmacy": pred_map["Monthly_By_Pharmacy"]},
            "store_issue": {"Central_Store_Issue_Coverage": pred_map["Central_Store_Issue_Coverage"]},
            "forecast_generic": {"Forecast_Generic_Form": pred_map["Forecast_Generic_Form"]},
            "forecast_branch": {"Forecast_By_Pharmacy": pred_map["Forecast_By_Pharmacy"]},
            "forecast_store": {"Forecast_Central_Store": pred_map["Forecast_Central_Store"]},
            "forecast_bundle": pred_map,
            "all": {
                "Analyzed_Inventory": self.df, "Filtered_View": filtered, "Network_All": combined, "Central_Store": store_df,
                "Branch_Pharmacies": branch_df, "Store_Branch_Link": linkage_df, "Branch_Comparison": current_compare_df(),
                "Redistribution": self.build_redistribution_df(filtered), "FEFO_Expiry": self.build_fefo_df(filtered),
                "ABC_VEN": self.build_abcven_df(filtered), "Purchase_Requests": self.build_purchase_df(filtered),
                "Emergency_Escalation": self.build_emergency_df(filtered), "Supplier_Summary": self.build_supplier_df(filtered),
                "Tablets": forms_map["Tablets"], "Ampoules": forms_map["Ampoules"], "Solutions": forms_map["Solutions"],
                "Miscellaneous": forms_map["Miscellaneous"], "Import_Report": pd.DataFrame({"mapping_report": self.mapping_report_text.splitlines()}),
                **pred_map,
            },
        }
        selected_map = data_map.get(export_mode, data_map["all"])
        self._write_excel_sheets(path, selected_map)
        QMessageBox.information(self, base.t(self.lang, "title"), base.t(self.lang, "save_ok"))
    def export_final_report(self):
        if self.df.empty:
            return
        base_df = self.filtered_df().copy()
        if base_df.empty:
            base_df = self.df.copy()
        scope = self._report_scope_dialog()
        if not scope:
            return
        output_format = self._report_output_format_dialog()
        if not output_format:
            return
        base_path = self._save_report_base_path(output_format)
        if not base_path:
            return
        combined, store_df, branch_df = self.split_store_branch_views(base_df)
        linkage_df = self.build_store_branch_support(base_df)
        pred_sections = self._build_prediction_export_map(base_df)
        if scope == "store":
            report_df = store_df
        elif scope == "branches":
            report_df = branch_df
        elif scope in {"tablets", "ampoules", "solutions", "miscellaneous"}:
            report_df = self.form_filtered_df(combined, scope)
        elif scope == "consumption_predictions":
            report_df = self.build_network_prediction_df(base_df)
        else:
            report_df = combined
        if report_df.empty and scope != "forms_split":
            report_df = combined
        top_red = report_df[report_df["escalation_level"] == "Red"].head(20) if not report_df.empty and "escalation_level" in report_df.columns else pd.DataFrame()
        top_reorder = report_df.sort_values("recommended_reorder_qty", ascending=False).head(20) if not report_df.empty and "recommended_reorder_qty" in report_df.columns else pd.DataFrame()
        supplier = report_df.groupby("supplier_name").agg(items=("generic_name", "count"), avg_on_time_rate=("supplier_on_time_rate", "mean")).reset_index() if not report_df.empty and "supplier_name" in report_df.columns else pd.DataFrame()
        active_branch = self.branch_filter.currentText()
        active_branch = "" if active_branch in ("", "All", base.t(self.lang, "all")) else active_branch
        active_search = self.search_edit.text().strip()
        note = ("هذا التقرير للتقييم والدعم فقط ولا يمثل قرار شراء أو صرف نهائي." if self.lang == "ar" else "This report is for decision support only and does not authorize purchasing or dispensing.")
        filters_html = f"<div class='card'><b>{'الفرع' if self.lang == 'ar' else 'Branch filter'}:</b> {active_branch or ('الكل' if self.lang=='ar' else 'All')}<br><b>{'نص البحث' if self.lang == 'ar' else 'Search text'}:</b> {active_search or ('بدون' if self.lang == 'ar' else 'None')}</div>"
        if scope == "current":
            excel_sections = {self._scope_title(scope): report_df}
        elif scope == "detailed":
            excel_sections = {
                ("الشبكة" if self.lang == "ar" else "Network_All"): combined,
                ("المخزن" if self.lang == "ar" else "Central_Store"): store_df,
                ("الفروع" if self.lang == "ar" else "Branch_Pharmacies"): branch_df,
                ("ربط المخزن بالفروع" if self.lang == "ar" else "Store_Branch_Link"): linkage_df,
                ("أعلى حالات التصعيد" if self.lang == "ar" else "Top_Red_Escalations"): top_red,
                ("أعلى احتياجات الشراء" if self.lang == "ar" else "Top_Reorder"): top_reorder,
                ("ملخص الموردين" if self.lang == "ar" else "Supplier_Summary"): supplier,
                **pred_sections,
            }
        elif scope == "forms_split":
            excel_sections = {
                ("أقراص" if self.lang == "ar" else "Tablets"): self.form_filtered_df(combined, "tablets"),
                ("أمبولات" if self.lang == "ar" else "Ampoules"): self.form_filtered_df(combined, "ampoules"),
                ("محاليل" if self.lang == "ar" else "Solutions"): self.form_filtered_df(combined, "solutions"),
                ("متنوعات" if self.lang == "ar" else "Miscellaneous"): self.form_filtered_df(combined, "miscellaneous"),
            }
        elif scope == "consumption_predictions":
            excel_sections = pred_sections
        else:
            excel_sections = {
                self._scope_title(scope): report_df,
                ("أعلى حالات التصعيد" if self.lang == "ar" else "Top_Red_Escalations"): top_red,
                ("أعلى احتياجات الشراء" if self.lang == "ar" else "Top_Reorder"): top_reorder,
                ("ملخص الموردين" if self.lang == "ar" else "Supplier_Summary"): supplier,
            }
        total_items = len(report_df) if report_df is not None else 0
        critical_count = int((report_df["shortage_risk"] == "Critical").sum()) if report_df is not None and "shortage_risk" in report_df.columns else 0
        expiry_count = int(report_df["expiry_risk"].isin(["Critical", "High", "Expired"]).sum()) if report_df is not None and "expiry_risk" in report_df.columns else 0
        reorder_count = int(pd.to_numeric(report_df["recommended_reorder_qty"], errors="coerce").fillna(0).sum()) if report_df is not None and "recommended_reorder_qty" in report_df.columns else 0
        red_count = int((report_df["escalation_level"] == "Red").sum()) if report_df is not None and "escalation_level" in report_df.columns else 0
        if output_format in {"html", "both"}:
            sections = []
            for section_title, section_df in excel_sections.items():
                sections.append(self._report_section_html(section_title, section_df))
            summary_html = f"""
            <html><head><meta charset='utf-8'>{self._html_report_style()}</head><body>
            <h1>{base.t(self.lang, 'title')}</h1>
            <h2>{self._scope_title(scope)}</h2>
            {filters_html}
            <div class='card'><b>{base.t(self.lang, 'kpi_total')}:</b> {total_items}<br>
            <b>{base.t(self.lang, 'kpi_critical')}:</b> {critical_count}<br>
            <b>{base.t(self.lang, 'kpi_expiry')}:</b> {expiry_count}<br>
            <b>{base.t(self.lang, 'kpi_reorder')}:</b> {reorder_count}<br>
            <b>{base.t(self.lang, 'kpi_red')}:</b> {red_count}</div>
            {''.join(sections)}
            <p class='note'>{note}</p>
            <p style='margin-top:20px;color:#475569'>{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
            </body></html>
            """
            html_path = base_path.with_suffix('.html') if output_format == 'both' else base_path
            Path(html_path).write_text(summary_html, encoding='utf-8')
        if output_format in {"excel", "both"}:
            summary_sheet_name = "ملخص" if self.lang == "ar" else "Summary"
            summary_df = pd.DataFrame([{
                ("نطاق التقرير" if self.lang == "ar" else "Report scope"): self._scope_title(scope),
                base.t(self.lang, 'kpi_total'): total_items,
                base.t(self.lang, 'kpi_critical'): critical_count,
                base.t(self.lang, 'kpi_expiry'): expiry_count,
                base.t(self.lang, 'kpi_reorder'): reorder_count,
                base.t(self.lang, 'kpi_red'): red_count,
                ("فلتر الفرع" if self.lang == "ar" else "Branch filter"): active_branch or ("الكل" if self.lang == "ar" else "All"),
                ("نص البحث" if self.lang == "ar" else "Search text"): active_search or ("بدون" if self.lang == "ar" else "None"),
            }])
            excel_path = base_path.with_suffix('.xlsx') if output_format == 'both' else base_path
            excel_map = {summary_sheet_name: summary_df}
            excel_map.update(excel_sections)
            self._write_excel_sheets(excel_path, excel_map)
        QMessageBox.information(self, base.t(self.lang, 'title'), base.t(self.lang, 'final_report_saved'))
# === Inventory tabs layer ===
import sys
from pathlib import Path
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QLabel, QFrame,
    QPushButton, QTabWidget, QTextEdit, QPlainTextEdit, QTableWidget,
    QHeaderView, QSizePolicy
)
APP_VERSION = "V6.9.3 Inventory Workspace"
APP_TITLE_EN = f"PharmaGuard AI Pro {APP_VERSION}"
APP_TITLE_AR = "فارماجارد برو V6.9.3 - مساحة عمل تحليل الصنف"
def _risk_rank(value: str) -> int:
    mapping = {"Low": 1, "Moderate": 2, "High": 3, "Critical": 4}
    return mapping.get(str(value), 0)
def _escalation_rank(value: str) -> int:
    mapping = {"Green": 1, "Amber": 2, "Red": 3}
    return mapping.get(str(value), 0)
class SummaryCard(QFrame):
    def __init__(self, title="", value="-", parent=None):
        super().__init__(parent)
        self.setObjectName("Card")
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(10, 8, 10, 8)
        lay.setSpacing(3)
        self.title_label = QLabel(title)
        self.value_label = QLabel(str(value))
        self.note_label = QLabel("")
        self.title_label.setStyleSheet("font-size:12px; font-weight:700;")
        self.value_label.setStyleSheet("font-size:24px; font-weight:800;")
        self.note_label.setStyleSheet("font-size:10px; color:#94a3b8;")
        self.note_label.setWordWrap(True)
        lay.addWidget(self.title_label)
        lay.addWidget(self.value_label)
        lay.addWidget(self.note_label)
    def set_content(self, title, value, note=""):
        self.title_label.setText(str(title))
        self.value_label.setText(str(value))
        self.note_label.setText(str(note))
class PharmaGuardMainWindow(InventorySearchMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_TITLE_AR if self.lang == "ar" else APP_TITLE_EN)
        self._upgrade_inventory_workspace()
        self._apply_inventory_texts()
        self._install_inventory_search_panel()
        self.refresh_inventory_page(self.filtered_df())
        self.refresh_inventory_details()
    def _upgrade_inventory_workspace(self):
        page = self.pages.get("inventory")
        if page is None:
            return
        lay = page.layout()
        if lay is None:
            return
        # Remove current widgets from the page layout so we can place them in a clearer workspace.
        for widget in [self.inventory_summary, self.compare_group, self.scenario_group]:
            try:
                lay.removeWidget(widget)
            except Exception:
                pass
            widget.setParent(None)
        # Quick navigation row
        self.inventory_nav_row = QHBoxLayout()
        self.inventory_nav_row.setContentsMargins(0, 0, 0, 0)
        self.inventory_nav_row.setSpacing(6)
        self.inventory_nav_buttons = []
        nav_specs = [
            ("overview", "ملخص الصنف", "Overview"),
            ("compare", "مقارنة الفروع", "Branch comparison"),
            ("details", "البيانات التفصيلية", "Detailed rows"),
            ("scenario", "سيناريو ماذا لو", "What-if scenario"),
        ]
        self.inventory_tabs = QTabWidget()
        self.inventory_tabs.setObjectName("InventoryTabs")
        # Overview tab
        self.inventory_overview_tab = QWidget()
        ov = QVBoxLayout(self.inventory_overview_tab)
        ov.setContentsMargins(6, 6, 6, 6)
        ov.setSpacing(6)
        self.inventory_cards_layout = QGridLayout()
        self.inventory_cards_layout.setSpacing(6)
        self.card_total_stock = SummaryCard()
        self.card_branches = SummaryCard()
        self.card_highest_risk = SummaryCard()
        self.card_monthly_use = SummaryCard()
        cards = [self.card_total_stock, self.card_branches, self.card_highest_risk, self.card_monthly_use]
        for i, card in enumerate(cards):
            self.inventory_cards_layout.addWidget(card, i // 2, i % 2)
        self.inventory_cards_layout.setHorizontalSpacing(6)
        self.inventory_cards_layout.setVerticalSpacing(6)
        ov.addLayout(self.inventory_cards_layout)
        self.inventory_snapshot = QTextEdit()
        self.inventory_snapshot.setReadOnly(True)
        self.inventory_snapshot.setMinimumHeight(80)
        self.inventory_snapshot.setMaximumHeight(112)
        self.inventory_summary.setMinimumHeight(68)
        self.inventory_summary.setMaximumHeight(96)
        ov.addWidget(self.inventory_snapshot)
        ov.addWidget(self.inventory_summary)
        # Compare tab
        self.compare_tab = QWidget()
        cv = QVBoxLayout(self.compare_tab)
        cv.setContentsMargins(4, 4, 4, 4)
        cv.setSpacing(4)
        self.compare_intro = QLabel()
        self.compare_intro.setWordWrap(True)
        self.compare_intro.setStyleSheet("font-size:12px; font-weight:600;")
        cv.addWidget(self.compare_intro)
        cv.addWidget(self.export_compare_btn)
        cv.addWidget(self.branch_compare_table)
        # Details tab
        self.details_tab = QWidget()
        dv = QVBoxLayout(self.details_tab)
        dv.setContentsMargins(4, 4, 4, 4)
        dv.setSpacing(4)
        self.inventory_detail_intro = QLabel()
        self.inventory_detail_intro.setWordWrap(True)
        self.inventory_detail_intro.setStyleSheet("font-size:12px; font-weight:600;")
        dv.addWidget(self.inventory_detail_intro)
        self.inventory_detail_table = self.make_table()
        dv.addWidget(self.inventory_detail_table)
        # Scenario tab
        self.scenario_tab = QWidget()
        sv = QVBoxLayout(self.scenario_tab)
        sv.setContentsMargins(6, 6, 6, 6)
        sv.setSpacing(6)
        self.scenario_intro = QLabel()
        self.scenario_intro.setWordWrap(True)
        self.scenario_intro.setStyleSheet("font-size:13px; font-weight:600;")
        sv.addWidget(self.scenario_intro)
        sv.addWidget(self.scenario_group)
        self.inventory_tabs.addTab(self.inventory_overview_tab, "")
        self.inventory_tabs.addTab(self.compare_tab, "")
        self.inventory_tabs.addTab(self.details_tab, "")
        self.inventory_tabs.addTab(self.scenario_tab, "")
        for idx, (_, ar, en) in enumerate(nav_specs):
            btn = QPushButton(ar if self.lang == "ar" else en)
            btn.clicked.connect(lambda _, i=idx: self.inventory_tabs.setCurrentIndex(i))
            self.inventory_nav_buttons.append(btn)
            self.inventory_nav_row.addWidget(btn)
        self.inventory_nav_row.addStretch(1)
        lay.addLayout(self.inventory_nav_row)
        lay.addWidget(self.inventory_tabs, 1)
        self._beautify_inventory_tables()
    def _beautify_inventory_tables(self):
        tables = [self.branch_compare_table, getattr(self, 'inventory_detail_table', None)]
        for table in tables:
            if table is None:
                continue
            table.setWordWrap(False)
            table.setTextElideMode(Qt.ElideNone)
            table.setHorizontalScrollMode(QTableWidget.ScrollPerPixel)
            table.setVerticalScrollMode(QTableWidget.ScrollPerPixel)
            table.setAlternatingRowColors(True)
            table.verticalHeader().setDefaultSectionSize(28)
            header = table.horizontalHeader()
            header.setStretchLastSection(False)
            try:
                header.setSectionResizeMode(QHeaderView.Interactive)
                header.setMinimumSectionSize(90)
            except Exception:
                pass
    def _apply_inventory_texts(self):
        is_ar = self.lang == "ar"
        self.inventory_tabs.setTabText(0, "ملخص الصنف" if is_ar else "Overview")
        self.inventory_tabs.setTabText(1, "مقارنة الفروع" if is_ar else "Branch comparison")
        self.inventory_tabs.setTabText(2, "البيانات التفصيلية" if is_ar else "Detailed rows")
        self.inventory_tabs.setTabText(3, "سيناريو ماذا لو" if is_ar else "What-if scenario")
        self.compare_intro.setText(
            "مقارنة الصنف عبر الفروع." if is_ar
            else "Compare the item across branches."
        )
        self.inventory_detail_intro.setText(
            "الصفوف التفصيلية بعد الفلاتر." if is_ar
            else "Detailed rows after filters."
        )
        self.scenario_intro.setText(
            "يمكنك هنا تجربة تأثير تغيير مدة التوريد أو الكمية المعلقة أو الزيادة في الرصيد قبل اتخاذ القرار." if is_ar
            else "Use this tab to test the effect of lead time, pending PO quantity, or extra stock before taking action."
        )
        self.card_total_stock.set_content("إجمالي الرصيد" if is_ar else "Total stock", "-")
        self.card_branches.set_content("عدد المواقع" if is_ar else "Locations", "-")
        self.card_highest_risk.set_content("أعلى خطورة" if is_ar else "Highest risk", "-")
        self.card_monthly_use.set_content("الاستهلاك الشهري" if is_ar else "Monthly use", "-")
    def _install_inventory_search_panel(self):
        if hasattr(self, "inventory_search_edit"):
            return
        form = self.inventory_top_group.layout()
        self.inventory_search_label = QLabel("ابحث بالاسم العلمي / العام / التجاري" if self.lang == "ar" else "Scientific / generic / trade")
        self.inventory_search_edit = QLineEdit()
        self.inventory_search_edit.setPlaceholderText("ابحث بالاسم العلمي أو العام أو التجاري" if self.lang == "ar" else "Search by scientific, generic, or trade name")
        self.inventory_search_edit.textChanged.connect(self.filter_inventory_medicine_list)
        try:
            form.insertRow(0, self.inventory_search_label, self.inventory_search_edit)
        except Exception:
            form.addRow(self.inventory_search_label, self.inventory_search_edit)
        self.medicine_combo.setMinimumHeight(30)
        self.medicine_combo.setMaxVisibleItems(30)
        line = self.medicine_combo.lineEdit()
        if line is not None:
            line.setPlaceholderText(t(self.lang, "type_to_search"))
    def filter_inventory_medicine_list(self, text=""):
        if not hasattr(self, "medicine_combo"):
            return
        cache = list(getattr(self, '_medicine_master_cache', getattr(self, '_medicine_display_cache', [])))
        selected = self.get_selected_medicine_name() if self.medicine_combo.count() else ""
        bundle = normalize_query_bundle(text)
        if any(bundle.values()):
            scored = []
            for rec in cache:
                score = max(query_match_score(bundle, rec[2]), query_match_score(bundle, rec[3]))
                if score > 0:
                    scored.append((score, rec))
            scored.sort(key=lambda x: -x[0])
            filtered = [rec for _, rec in scored]
        else:
            filtered = cache
        if not filtered:
            filtered = cache
        self.medicine_combo.blockSignals(True)
        self.medicine_combo.clear()
        for display, generic, blob, generic_blob, aliases in filtered:
            self.medicine_combo.addItem(display, generic)
        self._medicine_display_cache = filtered
        self.medicine_completer_model.setStringList([d for d, _, _, _, _ in filtered][:250])
        idx = -1
        if selected:
            for i, (_, generic, *_rest) in enumerate(filtered):
                if base.normalize_search_text(generic) == base.normalize_search_text(selected):
                    idx = i
                    break
        if idx < 0 and self.medicine_combo.count() > 0:
            idx = 0
        if idx >= 0:
            self.medicine_combo.setCurrentIndex(idx)
        self.medicine_combo.blockSignals(False)
        self.refresh_inventory_details()
    def refresh_inventory_page(self, df):
        super().refresh_inventory_page(df)
        # Keep a master cache and allow explicit search box filtering by scientific/trade names.
        self._medicine_master_cache = list(getattr(self, '_medicine_display_cache', []))
        if hasattr(self, 'inventory_search_edit'):
            self.filter_inventory_medicine_list(self.inventory_search_edit.text())
    def refresh_inventory_details(self):
        super().refresh_inventory_details()
        if not hasattr(self, "inventory_snapshot"):
            return
        df = self.filtered_df()
        if df.empty or self.medicine_combo.count() == 0:
            self.inventory_snapshot.setPlainText("لا توجد بيانات حالياً." if self.lang == "ar" else "No data currently available.")
            if hasattr(self, 'inventory_detail_table'):
                self.set_table_from_df(self.inventory_detail_table, base.pd.DataFrame())
            return
        name = self.get_selected_medicine_name()
        sub = df[df["medicine_name"] == name].copy()
        if sub.empty:
            self.inventory_snapshot.setPlainText("لا توجد بيانات للصنف المحدد." if self.lang == "ar" else "No rows found for the selected medicine.")
            if hasattr(self, 'inventory_detail_table'):
                self.set_table_from_df(self.inventory_detail_table, base.pd.DataFrame())
            return
        total_stock = float(base.pd.to_numeric(sub["current_stock"], errors="coerce").fillna(0).sum())
        total_pending = float(base.pd.to_numeric(sub["pending_po_qty"], errors="coerce").fillna(0).sum())
        monthly_use = float(base.pd.to_numeric(sub["avg_monthly_consumption"], errors="coerce").fillna(0).sum())
        locations = len({base.clean_text(v) for v in sub["branch_name"].tolist() if base.clean_text(v)})
        best_row = sub.sort_values(["escalation_level", "shortage_risk"], ascending=[True, False]).iloc[0]
        highest_risk = max([str(v) for v in sub["shortage_risk"].fillna("Low")], key=_risk_rank)
        highest_esc = max([str(v) for v in sub["escalation_level"].fillna("Green")], key=_escalation_rank)
        brands = base._join_unique(sub.get("brand_names_display", base.pd.Series(dtype=str)), sep=" | ")
        forms = base._join_unique(sub.get("dosage_form_group", base.pd.Series(dtype=str)), sep=" | ")
        source_names = base._join_unique(sub.get("source_medicine_names", base.pd.Series(dtype=str)), sep=" | ")
        is_ar = self.lang == "ar"
        self.card_total_stock.set_content(
            "إجمالي الرصيد" if is_ar else "Total stock",
            f"{total_stock:,.0f}",
            (f"معلق: {total_pending:,.0f}" if is_ar else f"Pending PO: {total_pending:,.0f}")
        )
        self.card_branches.set_content(
            "عدد المواقع" if is_ar else "Locations",
            str(locations),
            ("مخزن/صيدليات يظهر فيها الصنف" if is_ar else "Store/branches carrying the item")
        )
        self.card_highest_risk.set_content(
            "أعلى خطورة" if is_ar else "Highest risk",
            base.translate_display_value(highest_risk, self.lang),
            (f"التصعيد: {base.translate_display_value(highest_esc, self.lang)}" if is_ar else f"Escalation: {base.translate_display_value(highest_esc, self.lang)}")
        )
        self.card_monthly_use.set_content(
            "الاستهلاك الشهري" if is_ar else "Monthly use",
            f"{monthly_use:,.1f}",
            (f"شكل صيدلاني: {forms}" if is_ar else f"Dosage form: {forms}")
        )
        snapshot = []
        if is_ar:
            snapshot.append(f"الاسم العلمي/اسم التحليل: {name}")
            snapshot.append(f"الأسماء التجارية: {brands or 'غير متاح'}")
            snapshot.append(f"الأشكال الصيدلانية: {forms or 'غير متاح'}")
            snapshot.append(f"عدد المواقع التي يظهر فيها الصنف: {locations}")
            snapshot.append(f"إجمالي الرصيد عبر الشبكة: {total_stock:,.0f}")
            snapshot.append(f"إجمالي الاستهلاك الشهري التقديري: {monthly_use:,.1f}")
            snapshot.append(f"أعلى خطورة نقص: {base.translate_display_value(highest_risk, self.lang)}")
            snapshot.append(f"أعلى خطورة تكدس: {base.translate_display_value(max([str(v) for v in sub.get('overstock_risk', base.pd.Series(['Low'])).fillna('Low')], key=_risk_rank), self.lang)}")
            snapshot.append(f"الرصيد المتوقع بعد 30 يوم: {sub.get('predicted_stock_30d', base.pd.Series([0])).astype(float).sum():,.1f}")
            snapshot.append(f"أعلى مستوى تصعيد: {base.translate_display_value(highest_esc, self.lang)}")
            snapshot.append(f"أفضل إجراء مقترح الآن: {best_row.get('recommended_action', '')}")
            if source_names:
                snapshot.append(f"أسماء المصدر الخام: {source_names}")
        else:
            snapshot.append(f"Scientific / analytical item: {name}")
            snapshot.append(f"Trade names: {brands or 'N/A'}")
            snapshot.append(f"Dosage forms: {forms or 'N/A'}")
            snapshot.append(f"Locations carrying the item: {locations}")
            snapshot.append(f"Total network stock: {total_stock:,.0f}")
            snapshot.append(f"Estimated monthly consumption: {monthly_use:,.1f}")
            snapshot.append(f"Highest shortage risk: {base.translate_display_value(highest_risk, self.lang)}")
            snapshot.append(f"Highest overstock risk: {base.translate_display_value(max([str(v) for v in sub.get('overstock_risk', base.pd.Series(['Low'])).fillna('Low')], key=_risk_rank), self.lang)}")
            snapshot.append(f"Predicted stock after 30d: {sub.get('predicted_stock_30d', base.pd.Series([0])).astype(float).sum():,.1f}")
            snapshot.append(f"Highest escalation level: {base.translate_display_value(highest_esc, self.lang)}")
            snapshot.append(f"Recommended action now: {best_row.get('recommended_action', '')}")
            if source_names:
                snapshot.append(f"Raw source names: {source_names}")
        self.inventory_snapshot.setPlainText("\n".join(snapshot))

        # UI FIX: keep Branch comparison readable by showing only the most important columns
        # in this sub-screen. The full raw/details columns are still available in the
        # "Detailed rows" tab, so no functionality is removed.
        branch_cols = [
            "branch_name", "current_stock", "dynamic_min_stock", "avg_monthly_consumption",
            "pending_po_qty", "days_of_stock_left", "shortage_risk", "overstock_risk",
            "stock_status", "predicted_stock_30d", "escalation_level", "recommended_action",
        ]
        branch_view = sub[[c for c in branch_cols if c in sub.columns]].copy()
        if {"escalation_level", "days_of_stock_left"}.issubset(branch_view.columns):
            branch_view = branch_view.sort_values(["escalation_level", "days_of_stock_left"], ascending=[True, True])
        self.set_table_from_df(self.branch_compare_table, branch_view)
        try:
            self.branch_compare_table.setWordWrap(True)
            self.branch_compare_table.verticalHeader().setDefaultSectionSize(42)
            self.branch_compare_table.setMinimumHeight(360)
            readable_widths = {
                0: 170,  # Branch
                1: 110,  # Current stock
                2: 130,  # Dynamic min
                3: 140,  # Monthly use
                4: 120,  # Pending PO
                5: 120,  # Days left
                6: 125,  # Shortage risk
                7: 125,  # Overstock risk
                8: 125,  # Stock status
                9: 135,  # Predicted stock 30d
                10: 115, # Escalation
                11: 340, # Recommended action
            }
            for col_idx, width in readable_widths.items():
                if col_idx < self.branch_compare_table.columnCount():
                    self.branch_compare_table.setColumnWidth(col_idx, width)
            self.branch_compare_table.resizeRowsToContents()
        except Exception:
            pass

        detail_cols = [
            "branch_name", "medicine_name", "generic_name", "dosage_form_group", "brand_names_display", "source_medicine_names",
            "item_code", "current_stock", "dynamic_min_stock", "avg_daily_consumption", "avg_monthly_consumption",
            "pending_po_qty", "lead_time_days", "days_of_stock_left", "expected_stock_at_lead_time",
            "shortage_risk", "overstock_risk", "stock_status", "predicted_stock_30d", "predicted_stock_60d", "predicted_stock_90d", "expiry_risk", "escalation_level", "supplier_name", "recommended_action"
        ]
        detail = sub[[c for c in detail_cols if c in sub.columns]].copy()
        detail = detail.sort_values(["escalation_level", "days_of_stock_left"], ascending=[True, True])
        if hasattr(self, 'inventory_detail_table'):
            self.set_table_from_df(self.inventory_detail_table, detail)

# ========================= V10 ADD-ON: Stock Movement Ledger + SQLite foundation =========================
# This add-on keeps all existing features and adds a real movement log screen.
# It stores stock movements in a local SQLite database next to this Python file.

try:
    if "movement" not in NAV_ORDER:
        insert_pos = NAV_ORDER.index("purchase") if "purchase" in NAV_ORDER else len(NAV_ORDER)
        NAV_ORDER.insert(insert_pos, "movement")
    NAV_TEXT["movement"] = "nav_movement"
    TEXT["en"].update({
        "nav_movement": "📒 Stock Movement Ledger",
        "movement_title": "Stock Movement Ledger",
        "movement_note": "Safe mode: movements are stored in SQLite and exported for audit. Current snapshot stock is not changed automatically yet.",
        "movement_item": "Medicine",
        "movement_branch": "Branch / store",
        "movement_type": "Movement type",
        "movement_qty": "Quantity",
        "movement_batch": "Batch number",
        "movement_reason": "Reason / reference",
        "movement_user": "User",
        "movement_add": "Add movement",
        "movement_export": "Export movement ledger",
        "movement_summary": "Movement summary",
        "movement_db": "Local ledger DB",
    })
    TEXT["ar"].update({
        "nav_movement": "📒 سجل حركة الدواء",
        "movement_title": "سجل حركة الدواء",
        "movement_note": "وضع آمن: يتم حفظ الحركات في قاعدة SQLite وتصديرها للمراجعة. الرصيد الحالي في التحليل لا يتغير تلقائيًا في هذه النسخة.",
        "movement_item": "الصنف / الدواء",
        "movement_branch": "الفرع / المخزن",
        "movement_type": "نوع الحركة",
        "movement_qty": "الكمية",
        "movement_batch": "رقم الباتش",
        "movement_reason": "السبب / المرجع",
        "movement_user": "المستخدم",
        "movement_add": "إضافة حركة",
        "movement_export": "تصدير سجل الحركة",
        "movement_summary": "ملخص الحركة",
        "movement_db": "قاعدة سجل الحركة المحلية",
    })
except Exception:
    pass

MOVEMENT_COLUMNS = [
    "movement_id", "timestamp", "movement_type", "medicine_name", "generic_name",
    "item_code", "branch_name", "quantity", "batch_number", "reason", "user_name"
]

def _movement_db_path():
    try:
        return str(Path(__file__).resolve().with_name("pharmaguard_movement_ledger.db"))
    except Exception:
        return "pharmaguard_movement_ledger.db"

def _ensure_movement_db(db_path=None):
    path = db_path or _movement_db_path()
    conn = sqlite3.connect(path)
    try:
        cur = conn.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS stock_movements (
                movement_id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                movement_type TEXT NOT NULL,
                medicine_name TEXT NOT NULL,
                generic_name TEXT,
                item_code TEXT,
                branch_name TEXT,
                quantity REAL NOT NULL,
                batch_number TEXT,
                reason TEXT,
                user_name TEXT
            )
        """)
        conn.commit()
    finally:
        conn.close()
    return path

def _load_movement_ledger(db_path=None):
    path = _ensure_movement_db(db_path)
    conn = sqlite3.connect(path)
    try:
        df = pd.read_sql_query("SELECT * FROM stock_movements ORDER BY movement_id DESC", conn)
    except Exception:
        df = pd.DataFrame(columns=MOVEMENT_COLUMNS)
    finally:
        conn.close()
    for col in MOVEMENT_COLUMNS:
        if col not in df.columns:
            df[col] = None
    return df[MOVEMENT_COLUMNS]

def _insert_movement(row, db_path=None):
    path = _ensure_movement_db(db_path)
    conn = sqlite3.connect(path)
    try:
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO stock_movements
            (timestamp, movement_type, medicine_name, generic_name, item_code, branch_name, quantity, batch_number, reason, user_name)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            row.get("timestamp"), row.get("movement_type"), row.get("medicine_name"), row.get("generic_name"),
            row.get("item_code"), row.get("branch_name"), row.get("quantity"), row.get("batch_number"),
            row.get("reason"), row.get("user_name"),
        ))
        conn.commit()
    finally:
        conn.close()

class PharmaGuardMainWindow(PharmaGuardMainWindow):
    def __init__(self):
        self.movement_db_path = _ensure_movement_db()
        self.movement_df = _load_movement_ledger(self.movement_db_path)
        super().__init__()

    def build_page(self, key):
        if key != "movement":
            return super().build_page(key)
        container = QWidget()
        lay = QVBoxLayout(container)
        lay.setContentsMargins(8, 8, 8, 8)
        lay.setSpacing(8)

        self.movement_group = QGroupBox()
        wrap = QVBoxLayout(self.movement_group)

        self.movement_note_label = QLabel()
        self.movement_note_label.setWordWrap(True)
        self.movement_note_label.setStyleSheet("font-weight:700; padding:6px; border:1px solid #cbd5e1; border-radius:8px;")
        wrap.addWidget(self.movement_note_label)

        form_row1 = QHBoxLayout()
        form_row2 = QHBoxLayout()

        self.movement_item_label = QLabel()
        self.movement_item_combo = QComboBox()
        self.movement_item_combo.setEditable(True)
        self.movement_item_combo.setMinimumWidth(260)

        self.movement_branch_label = QLabel()
        self.movement_branch_combo = QComboBox()
        self.movement_branch_combo.setEditable(True)
        self.movement_branch_combo.setMinimumWidth(180)

        self.movement_type_label = QLabel()
        self.movement_type_combo = QComboBox()
        self.movement_type_combo.addItems(["Receive", "Issue", "Transfer", "Return", "Waste", "Adjustment"])

        self.movement_qty_label = QLabel()
        self.movement_qty_spin = QDoubleSpinBox()
        self.movement_qty_spin.setRange(0, 100000000)
        self.movement_qty_spin.setDecimals(2)
        self.movement_qty_spin.setValue(1)

        form_row1.addWidget(self.movement_item_label)
        form_row1.addWidget(self.movement_item_combo)
        form_row1.addWidget(self.movement_branch_label)
        form_row1.addWidget(self.movement_branch_combo)
        form_row1.addWidget(self.movement_type_label)
        form_row1.addWidget(self.movement_type_combo)
        form_row1.addWidget(self.movement_qty_label)
        form_row1.addWidget(self.movement_qty_spin)

        self.movement_batch_label = QLabel()
        self.movement_batch_edit = QLineEdit()
        self.movement_reason_label = QLabel()
        self.movement_reason_edit = QLineEdit()
        self.movement_user_label = QLabel()
        self.movement_user_edit = QLineEdit()
        self.movement_user_edit.setText("pharmacist")

        self.movement_add_btn = QPushButton()
        self.movement_add_btn.clicked.connect(self.add_stock_movement)
        self.movement_export_btn = QPushButton()
        self.movement_export_btn.clicked.connect(self.export_stock_movement_ledger)

        form_row2.addWidget(self.movement_batch_label)
        form_row2.addWidget(self.movement_batch_edit)
        form_row2.addWidget(self.movement_reason_label)
        form_row2.addWidget(self.movement_reason_edit)
        form_row2.addWidget(self.movement_user_label)
        form_row2.addWidget(self.movement_user_edit)
        form_row2.addWidget(self.movement_add_btn)
        form_row2.addWidget(self.movement_export_btn)

        self.movement_summary_label = QLabel()
        self.movement_summary_label.setWordWrap(True)
        self.movement_summary_label.setStyleSheet("font-weight:700; padding:6px;")
        self.movement_table = self.make_table()
        self.movement_table.setMinimumHeight(420)
        self.movement_table.setWordWrap(True)
        self.movement_table.verticalHeader().setDefaultSectionSize(40)

        wrap.addLayout(form_row1)
        wrap.addLayout(form_row2)
        wrap.addWidget(self.movement_summary_label)
        wrap.addWidget(self.movement_table)
        lay.addWidget(self.movement_group)
        lay.addStretch(1)
        return container

    def update_labels(self):
        super().update_labels()
        if not hasattr(self, "movement_group"):
            return
        self.movement_group.setTitle(t(self.lang, "movement_title"))
        self.movement_note_label.setText(t(self.lang, "movement_note"))
        self.movement_item_label.setText(t(self.lang, "movement_item"))
        self.movement_branch_label.setText(t(self.lang, "movement_branch"))
        self.movement_type_label.setText(t(self.lang, "movement_type"))
        self.movement_qty_label.setText(t(self.lang, "movement_qty"))
        self.movement_batch_label.setText(t(self.lang, "movement_batch"))
        self.movement_reason_label.setText(t(self.lang, "movement_reason"))
        self.movement_user_label.setText(t(self.lang, "movement_user"))
        self.movement_add_btn.setText(t(self.lang, "movement_add"))
        self.movement_export_btn.setText(t(self.lang, "movement_export"))
        self._refresh_movement_combos()
        self.refresh_movement_page()

    def refresh_all_views(self):
        super().refresh_all_views()
        self._refresh_movement_combos()
        if hasattr(self, "movement_table"):
            self.refresh_movement_page()

    def refresh_current_view(self):
        current_key = NAV_ORDER[self.stack.currentIndex()]
        if current_key == "movement":
            self.refresh_movement_page()
            return
        super().refresh_current_view()

    def _refresh_movement_combos(self):
        if not hasattr(self, "movement_item_combo"):
            return
        current_item = self.movement_item_combo.currentText()
        current_branch = self.movement_branch_combo.currentText() if hasattr(self, "movement_branch_combo") else ""
        items = []
        branches = []
        try:
            if hasattr(self, "df") and self.df is not None and not self.df.empty:
                name_col = "generic_name" if "generic_name" in self.df.columns else "medicine_name"
                items = sorted({clean_text(v) for v in self.df[name_col].dropna().astype(str).tolist() if clean_text(v)})
                if "branch_name" in self.df.columns:
                    branches = sorted({clean_text(v) for v in self.df["branch_name"].dropna().astype(str).tolist() if clean_text(v)})
        except Exception:
            pass
        self.movement_item_combo.blockSignals(True)
        self.movement_item_combo.clear()
        self.movement_item_combo.addItems(items)
        if current_item:
            idx = self.movement_item_combo.findText(current_item)
            if idx >= 0:
                self.movement_item_combo.setCurrentIndex(idx)
            else:
                self.movement_item_combo.setEditText(current_item)
        self.movement_item_combo.blockSignals(False)
        self.movement_branch_combo.blockSignals(True)
        self.movement_branch_combo.clear()
        self.movement_branch_combo.addItems(branches)
        if current_branch:
            idx = self.movement_branch_combo.findText(current_branch)
            if idx >= 0:
                self.movement_branch_combo.setCurrentIndex(idx)
            else:
                self.movement_branch_combo.setEditText(current_branch)
        self.movement_branch_combo.blockSignals(False)

    def _movement_selected_item_context(self, item_text, branch_text):
        ctx = {"medicine_name": item_text, "generic_name": item_text, "item_code": "", "branch_name": branch_text}
        try:
            if not hasattr(self, "df") or self.df is None or self.df.empty:
                return ctx
            q = normalize_search_text(item_text)
            working = self.df.copy()
            if branch_text and "branch_name" in working.columns:
                exact_branch = working["branch_name"].astype(str).map(clean_text).str.lower() == clean_text(branch_text).lower()
                if exact_branch.any():
                    working = working[exact_branch]
            for col in ["generic_name", "medicine_name", "analysis_name", "source_medicine_names"]:
                if col in working.columns:
                    mask = working[col].astype(str).map(normalize_search_text).str.contains(q, na=False, regex=False)
                    if mask.any():
                        row = working[mask].iloc[0]
                        ctx["medicine_name"] = clean_text(row.get("medicine_name", item_text))
                        ctx["generic_name"] = clean_text(row.get("generic_name", item_text))
                        ctx["item_code"] = clean_text(row.get("item_code", ""))
                        ctx["branch_name"] = clean_text(row.get("branch_name", branch_text)) or branch_text
                        break
        except Exception:
            pass
        return ctx

    def add_stock_movement(self):
        item_text = clean_text(self.movement_item_combo.currentText() if hasattr(self, "movement_item_combo") else "")
        branch_text = clean_text(self.movement_branch_combo.currentText() if hasattr(self, "movement_branch_combo") else "")
        qty = float(self.movement_qty_spin.value() if hasattr(self, "movement_qty_spin") else 0)
        if not item_text:
            QMessageBox.warning(self, t(self.lang, "title"), "Choose a medicine first. / اختر الصنف أولًا")
            return
        if qty <= 0:
            QMessageBox.warning(self, t(self.lang, "title"), "Quantity must be greater than zero. / الكمية لازم تكون أكبر من صفر")
            return
        ctx = self._movement_selected_item_context(item_text, branch_text)
        row = {
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "movement_type": clean_text(self.movement_type_combo.currentText()),
            "medicine_name": ctx.get("medicine_name") or item_text,
            "generic_name": ctx.get("generic_name") or item_text,
            "item_code": ctx.get("item_code", ""),
            "branch_name": ctx.get("branch_name") or branch_text,
            "quantity": qty,
            "batch_number": clean_text(self.movement_batch_edit.text()),
            "reason": clean_text(self.movement_reason_edit.text()),
            "user_name": clean_text(self.movement_user_edit.text()) or "pharmacist",
        }
        try:
            _insert_movement(row, self.movement_db_path)
            self.movement_df = _load_movement_ledger(self.movement_db_path)
            self.refresh_movement_page()
            self.movement_reason_edit.clear()
            self.movement_batch_edit.clear()
            QMessageBox.information(self, t(self.lang, "title"), "Movement saved. / تم حفظ الحركة")
        except Exception as exc:
            QMessageBox.critical(self, t(self.lang, "title"), f"Could not save movement.\n{exc}")

    def refresh_movement_page(self):
        if not hasattr(self, "movement_table"):
            return
        try:
            self.movement_df = _load_movement_ledger(self.movement_db_path)
        except Exception:
            if not hasattr(self, "movement_df") or self.movement_df is None:
                self.movement_df = pd.DataFrame(columns=MOVEMENT_COLUMNS)
        df = self.movement_df.copy()
        show_cols = [c for c in ["movement_id", "timestamp", "movement_type", "generic_name", "item_code", "branch_name", "quantity", "batch_number", "reason", "user_name"] if c in df.columns]
        self.set_table_from_df(self.movement_table, df[show_cols].head(500) if show_cols else pd.DataFrame())
        try:
            self.movement_table.setColumnWidth(0, 90)
            self.movement_table.setColumnWidth(1, 155)
            self.movement_table.setColumnWidth(2, 115)
            self.movement_table.setColumnWidth(3, 230)
            self.movement_table.setColumnWidth(5, 160)
            self.movement_table.setColumnWidth(8, 260)
            self.movement_table.resizeRowsToContents()
        except Exception:
            pass
        total_rows = len(df)
        total_qty = float(pd.to_numeric(df.get("quantity", pd.Series(dtype=float)), errors="coerce").fillna(0).sum()) if total_rows else 0
        db_txt = f"{t(self.lang, 'movement_db')}: {self.movement_db_path}"
        if self.lang == "ar":
            summary = f"{t(self.lang, 'movement_summary')}: عدد الحركات = {total_rows:,} | إجمالي الكميات المسجلة = {total_qty:,.2f}\n{db_txt}"
        else:
            summary = f"{t(self.lang, 'movement_summary')}: movements = {total_rows:,} | total recorded qty = {total_qty:,.2f}\n{db_txt}"
        self.movement_summary_label.setText(summary)

    def export_stock_movement_ledger(self):
        try:
            self.movement_df = _load_movement_ledger(self.movement_db_path)
            if self.movement_df.empty:
                QMessageBox.information(self, t(self.lang, "title"), "No movements to export. / لا توجد حركات للتصدير")
                return
            path, _ = QFileDialog.getSaveFileName(self, t(self.lang, "movement_export"), "stock_movement_ledger.xlsx", "Excel Workbook (*.xlsx);;CSV (*.csv)")
            if not path:
                return
            if str(path).lower().endswith(".csv"):
                self.movement_df.to_csv(path, index=False, encoding="utf-8-sig")
            else:
                if not str(path).lower().endswith(".xlsx"):
                    path = path + ".xlsx"
                self.movement_df.to_excel(path, index=False)
            QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "save_ok"))
        except Exception as exc:
            QMessageBox.critical(self, t(self.lang, "title"), f"Export failed.\n{exc}")
# ======================= END V10 ADD-ON =======================



# ======================= V11 ADD-ON: Enhanced Forecasting + Operational Database =======================
# هدف النسخة: تقوية Forecasting وربط النظام بقاعدة SQLite تشغيلية بدون حذف أي ميزة قديمة.
# - Forecast Center أصبح فيه أكثر من طريقة توقع.
# - قاعدة بيانات تشغيلية تحفظ آخر Snapshot ونتائج Forecast للمراجعة والرجوع لها.

try:
    TEXT["en"].update({
        "forecast_model": "Forecast method",
        "forecast_model_rule": "Rule-based",
        "forecast_model_movement": "Movement weighted",
        "forecast_model_safety": "Safety-adjusted",
        "forecast_save_db": "Save forecast to DB",
        "forecast_confidence": "Forecast confidence",
        "db_tools_title": "Operational SQLite Database",
        "db_tools_note": "The app keeps a local SQLite database for inventory snapshots and forecast results. This does not remove Excel import; it adds safer traceability.",
        "db_save_snapshot": "Save current snapshot",
        "db_load_snapshot": "Load latest snapshot",
        "db_save_forecast": "Save current forecast",
        "db_status": "Database status",
        "db_path": "Database path",
        "db_saved": "Saved to operational database.",
        "db_loaded": "Latest snapshot loaded from operational database.",
        "db_no_snapshot": "No saved inventory snapshot found.",
    })
    TEXT["ar"].update({
        "forecast_model": "طريقة التوقع",
        "forecast_model_rule": "قواعد بسيطة",
        "forecast_model_movement": "مع ترجيح سجل الحركة",
        "forecast_model_safety": "مع عامل أمان",
        "forecast_save_db": "حفظ التوقع في قاعدة البيانات",
        "forecast_confidence": "ثقة التوقع",
        "db_tools_title": "قاعدة بيانات تشغيلية SQLite",
        "db_tools_note": "البرنامج يحفظ نسخة تشغيلية محلية من المخزون ونتائج التوقع داخل SQLite. هذا لا يلغي استيراد Excel، لكنه يضيف تتبع ورجوع آمن للبيانات.",
        "db_save_snapshot": "حفظ Snapshot الحالي",
        "db_load_snapshot": "تحميل آخر Snapshot",
        "db_save_forecast": "حفظ التوقع الحالي",
        "db_status": "حالة قاعدة البيانات",
        "db_path": "مسار قاعدة البيانات",
        "db_saved": "تم الحفظ في قاعدة البيانات التشغيلية.",
        "db_loaded": "تم تحميل آخر Snapshot من قاعدة البيانات التشغيلية.",
        "db_no_snapshot": "لا يوجد Snapshot محفوظ حتى الآن.",
    })
except Exception:
    pass

for _col, _labels in {
    "forecast_model": {"en": "Forecast method", "ar": "طريقة التوقع"},
    "imported_monthly_consumption": {"en": "Imported monthly use", "ar": "الاستهلاك الشهري المستورد"},
    "ledger_30d_issue_qty": {"en": "Ledger 30d issue", "ar": "صرف آخر 30 يوم"},
    "ledger_projected_monthly_consumption": {"en": "Ledger projected monthly", "ar": "الاستهلاك المتوقع من سجل الحركة"},
    "forecast_monthly_demand": {"en": "Forecast monthly demand", "ar": "الاستهلاك الشهري المتوقع"},
    "forecast_horizon_demand": {"en": "Forecast horizon demand", "ar": "طلب فترة التوقع"},
    "forecast_available_stock": {"en": "Available stock", "ar": "الرصيد المتاح"},
    "forecast_gap": {"en": "Forecast gap", "ar": "فجوة التوقع"},
    "forecast_confidence": {"en": "Confidence", "ar": "الثقة"},
    "forecast_reason": {"en": "Forecast reason", "ar": "سبب التوقع"},
}.items():
    try:
        COLUMN_LABELS[_col] = _labels
    except Exception:
        pass

OPERATIONAL_DB_COLUMNS = [
    "snapshot_id", "row_no", "medicine_name", "generic_name", "item_code", "branch_name",
    "current_stock", "min_stock_level", "dynamic_min_stock", "avg_daily_consumption",
    "avg_monthly_consumption", "lead_time_days", "pending_po_qty", "expiry_date",
    "clinical_priority", "supplier_name", "supplier_on_time_rate", "shortage_risk",
    "overstock_risk", "stock_status", "recommended_reorder_qty", "recommended_action",
    "extra_json"
]


def _operational_db_path():
    try:
        return str(Path(__file__).resolve().with_name("pharmaguard_operational.db"))
    except Exception:
        return "pharmaguard_operational.db"


def _ensure_operational_db(db_path=None):
    path = db_path or _operational_db_path()
    conn = sqlite3.connect(path)
    try:
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("""
            CREATE TABLE IF NOT EXISTS inventory_snapshots (
                snapshot_id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                source_name TEXT,
                row_count INTEGER DEFAULT 0,
                note TEXT
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS inventory_snapshot_rows (
                snapshot_id INTEGER NOT NULL,
                row_no INTEGER NOT NULL,
                medicine_name TEXT,
                generic_name TEXT,
                item_code TEXT,
                branch_name TEXT,
                current_stock REAL,
                min_stock_level REAL,
                dynamic_min_stock REAL,
                avg_daily_consumption REAL,
                avg_monthly_consumption REAL,
                lead_time_days REAL,
                pending_po_qty REAL,
                expiry_date TEXT,
                clinical_priority TEXT,
                supplier_name TEXT,
                supplier_on_time_rate REAL,
                shortage_risk TEXT,
                overstock_risk TEXT,
                stock_status TEXT,
                recommended_reorder_qty REAL,
                recommended_action TEXT,
                extra_json TEXT,
                PRIMARY KEY (snapshot_id, row_no),
                FOREIGN KEY (snapshot_id) REFERENCES inventory_snapshots(snapshot_id) ON DELETE CASCADE
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS forecast_results (
                forecast_id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                source_name TEXT,
                horizon_days INTEGER,
                forecast_model TEXT,
                rows_count INTEGER DEFAULT 0,
                total_forecast_gap REAL DEFAULT 0,
                payload_json TEXT
            )
        """)
        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_inventory_snapshot_rows_item
            ON inventory_snapshot_rows(generic_name, branch_name)
        """)
        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_forecast_results_created
            ON forecast_results(created_at)
        """)
        conn.commit()
    finally:
        conn.close()
    return path


def _json_safe_records(df):
    safe = df.copy()
    for col in safe.columns:
        if pd.api.types.is_datetime64_any_dtype(safe[col]):
            safe[col] = safe[col].astype(str)
    safe = safe.replace({np.nan: None})
    return safe.to_dict(orient="records")


def _save_inventory_snapshot_to_db(df, source_name="", note="", db_path=None):
    path = _ensure_operational_db(db_path)
    if df is None:
        df = pd.DataFrame()
    working = df.copy()
    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = sqlite3.connect(path)
    try:
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO inventory_snapshots(created_at, source_name, row_count, note) VALUES (?, ?, ?, ?)",
            (created_at, clean_text(source_name), int(len(working)), clean_text(note))
        )
        snapshot_id = cur.lastrowid
        base_cols = [c for c in OPERATIONAL_DB_COLUMNS if c not in {"snapshot_id", "row_no", "extra_json"}]
        rows = []
        for row_no, (_, row) in enumerate(working.iterrows(), start=1):
            values = []
            for col in base_cols:
                value = row.get(col, None)
                if pd.isna(value):
                    value = None
                if isinstance(value, (pd.Timestamp, datetime, date)):
                    value = str(value)
                values.append(value)
            extra_cols = [c for c in working.columns if c not in base_cols]
            extra = {c: row.get(c, None) for c in extra_cols[:80]}
            extra_df = pd.DataFrame([extra])
            extra_json = extra_df.to_json(orient="records", force_ascii=False, date_format="iso") if extra else "[]"
            rows.append([snapshot_id, row_no, *values, extra_json])
        placeholders = ",".join(["?"] * len(OPERATIONAL_DB_COLUMNS))
        cur.executemany(
            f"INSERT INTO inventory_snapshot_rows({','.join(OPERATIONAL_DB_COLUMNS)}) VALUES ({placeholders})",
            rows
        )
        conn.commit()
        return snapshot_id, path
    finally:
        conn.close()


def _load_latest_inventory_snapshot_from_db(db_path=None):
    path = _ensure_operational_db(db_path)
    conn = sqlite3.connect(path)
    try:
        meta = pd.read_sql_query(
            "SELECT * FROM inventory_snapshots ORDER BY snapshot_id DESC LIMIT 1",
            conn
        )
        if meta.empty:
            return pd.DataFrame(), None
        snapshot_id = int(meta.iloc[0]["snapshot_id"])
        df = pd.read_sql_query(
            "SELECT * FROM inventory_snapshot_rows WHERE snapshot_id = ? ORDER BY row_no",
            conn,
            params=(snapshot_id,)
        )
        drop_cols = [c for c in ["snapshot_id", "row_no", "extra_json"] if c in df.columns]
        df = df.drop(columns=drop_cols)
        return df, meta.iloc[0].to_dict()
    finally:
        conn.close()


def _save_forecast_to_db(forecast_df, source_name="", horizon_days=30, forecast_model="", db_path=None):
    path = _ensure_operational_db(db_path)
    if forecast_df is None:
        forecast_df = pd.DataFrame()
    payload = pd.DataFrame(_json_safe_records(forecast_df)).to_json(orient="records", force_ascii=False)
    total_gap = 0.0
    if not forecast_df.empty and "forecast_gap" in forecast_df.columns:
        total_gap = float(pd.to_numeric(forecast_df["forecast_gap"], errors="coerce").fillna(0).sum())
    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = sqlite3.connect(path)
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO forecast_results(created_at, source_name, horizon_days, forecast_model, rows_count, total_forecast_gap, payload_json)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (created_at, clean_text(source_name), int(horizon_days), clean_text(forecast_model), int(len(forecast_df)), total_gap, payload)
        )
        conn.commit()
        return cur.lastrowid, path
    finally:
        conn.close()


def _movement_issue_projection(movement_df, horizon_reference_days=30):
    if movement_df is None or movement_df.empty:
        return pd.DataFrame(columns=["generic_key", "ledger_30d_issue_qty", "ledger_projected_monthly_consumption"])
    mov = movement_df.copy()
    if "timestamp" not in mov.columns or "generic_name" not in mov.columns:
        return pd.DataFrame(columns=["generic_key", "ledger_30d_issue_qty", "ledger_projected_monthly_consumption"])
    mov["_ts"] = pd.to_datetime(mov["timestamp"], errors="coerce")
    mov = mov[mov["_ts"].notna()].copy()
    if mov.empty:
        return pd.DataFrame(columns=["generic_key", "ledger_30d_issue_qty", "ledger_projected_monthly_consumption"])
    cutoff = pd.Timestamp(datetime.now()) - pd.Timedelta(days=30)
    mov = mov[mov["_ts"] >= cutoff].copy()
    if mov.empty:
        return pd.DataFrame(columns=["generic_key", "ledger_30d_issue_qty", "ledger_projected_monthly_consumption"])
    mov["_type"] = mov.get("movement_type", "").astype(str).str.lower()
    mov["_qty"] = pd.to_numeric(mov.get("quantity", 0), errors="coerce").fillna(0)
    # الصرف والهالك يعتبروا استهلاك. المرتجع يقلل الاستهلاك. الوارد والتحويل لا يعتبروا demand مباشر.
    mov["_signed_issue"] = np.select(
        [mov["_type"].isin(["issue", "waste"]), mov["_type"].isin(["return"])],
        [mov["_qty"], -mov["_qty"]],
        default=0
    )
    mov["generic_key"] = mov["generic_name"].astype(str).map(normalize_search_text)
    out = mov.groupby("generic_key", dropna=False).agg(
        ledger_30d_issue_qty=("_signed_issue", "sum"),
        first_movement=("_ts", "min"),
        last_movement=("_ts", "max"),
    ).reset_index()
    out["ledger_30d_issue_qty"] = pd.to_numeric(out["ledger_30d_issue_qty"], errors="coerce").fillna(0).clip(lower=0)
    span_days = (out["last_movement"] - out["first_movement"]).dt.days + 1
    span_days = span_days.clip(lower=1, upper=30)
    out["ledger_projected_monthly_consumption"] = np.round((out["ledger_30d_issue_qty"] / span_days) * 30, 2)
    return out[["generic_key", "ledger_30d_issue_qty", "ledger_projected_monthly_consumption"]]


class PharmaGuardMainWindow(PharmaGuardMainWindow):
    def __init__(self):
        self.operational_db_path = _ensure_operational_db()
        self._last_forecast_df = pd.DataFrame()
        super().__init__()

    def build_page(self, key):
        page = super().build_page(key)
        try:
            if key == "forecast" and hasattr(self, "forecast_group"):
                layout = self.forecast_group.layout()
                control_row = QHBoxLayout()
                self.forecast_model_label = QLabel()
                self.forecast_model_combo = QComboBox()
                self.forecast_model_combo.addItems(["Rule-based", "Movement weighted", "Safety-adjusted"])
                self.forecast_model_combo.currentIndexChanged.connect(self.refresh_current_view)
                self.forecast_save_db_btn = QPushButton()
                self.forecast_save_db_btn.clicked.connect(self.save_current_forecast_to_db)
                control_row.addWidget(self.forecast_model_label)
                control_row.addWidget(self.forecast_model_combo)
                control_row.addStretch(1)
                control_row.addWidget(self.forecast_save_db_btn)
                layout.insertLayout(1, control_row)
            if key == "connector" and hasattr(self, "connector_group"):
                layout = self.connector_group.layout()
                self.db_tools_group = QGroupBox()
                db_lay = QVBoxLayout(self.db_tools_group)
                self.db_tools_note_label = QLabel()
                self.db_tools_note_label.setWordWrap(True)
                self.db_status_label = QLabel()
                self.db_status_label.setWordWrap(True)
                db_buttons = QHBoxLayout()
                self.db_save_snapshot_btn = QPushButton()
                self.db_load_snapshot_btn = QPushButton()
                self.db_save_forecast_btn = QPushButton()
                self.db_save_snapshot_btn.clicked.connect(self.save_current_snapshot_to_db)
                self.db_load_snapshot_btn.clicked.connect(self.load_latest_snapshot_from_db)
                self.db_save_forecast_btn.clicked.connect(self.save_current_forecast_to_db)
                db_buttons.addWidget(self.db_save_snapshot_btn)
                db_buttons.addWidget(self.db_load_snapshot_btn)
                db_buttons.addWidget(self.db_save_forecast_btn)
                db_lay.addWidget(self.db_tools_note_label)
                db_lay.addLayout(db_buttons)
                db_lay.addWidget(self.db_status_label)
                layout.addWidget(self.db_tools_group)
        except Exception:
            traceback.print_exc()
        return page

    def update_labels(self):
        super().update_labels()
        try:
            if hasattr(self, "forecast_model_label"):
                self.forecast_model_label.setText(t(self.lang, "forecast_model"))
            if hasattr(self, "forecast_save_db_btn"):
                self.forecast_save_db_btn.setText(t(self.lang, "forecast_save_db"))
            if hasattr(self, "db_tools_group"):
                self.db_tools_group.setTitle(t(self.lang, "db_tools_title"))
                self.db_tools_note_label.setText(t(self.lang, "db_tools_note"))
                self.db_save_snapshot_btn.setText(t(self.lang, "db_save_snapshot"))
                self.db_load_snapshot_btn.setText(t(self.lang, "db_load_snapshot"))
                self.db_save_forecast_btn.setText(t(self.lang, "db_save_forecast"))
                self._update_db_status_label()
        except Exception:
            pass

    def recompute(self):
        super().recompute()
        # حفظ تلقائي خفيف للـ snapshot بعد الاستيراد أو التحديث، بدون تعطيل البرنامج لو فشل.
        try:
            if hasattr(self, "df") and self.df is not None and not self.df.empty and len(self.df) <= 150000:
                self._last_snapshot_id, _ = _save_inventory_snapshot_to_db(
                    self.df,
                    source_name=getattr(self, "current_sheet", ""),
                    note="auto snapshot after recompute",
                    db_path=self.operational_db_path,
                )
                self._update_db_status_label()
        except Exception:
            traceback.print_exc()

    def _update_db_status_label(self):
        if not hasattr(self, "db_status_label"):
            return
        try:
            path = self.operational_db_path
            conn = sqlite3.connect(path)
            try:
                snap_count = pd.read_sql_query("SELECT COUNT(*) AS n FROM inventory_snapshots", conn).iloc[0]["n"]
                fc_count = pd.read_sql_query("SELECT COUNT(*) AS n FROM forecast_results", conn).iloc[0]["n"]
            finally:
                conn.close()
            txt = f"{t(self.lang, 'db_status')}: snapshots={snap_count:,} | forecasts={fc_count:,}\n{t(self.lang, 'db_path')}: {path}"
            self.db_status_label.setText(txt)
        except Exception as exc:
            self.db_status_label.setText(f"{t(self.lang, 'db_status')}: {exc}")

    def save_current_snapshot_to_db(self):
        try:
            if not hasattr(self, "df") or self.df is None or self.df.empty:
                QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "no_data"))
                return
            snapshot_id, path = _save_inventory_snapshot_to_db(
                self.df,
                source_name=getattr(self, "current_sheet", ""),
                note="manual save from Connector Center",
                db_path=self.operational_db_path,
            )
            self._last_snapshot_id = snapshot_id
            self._update_db_status_label()
            QMessageBox.information(self, t(self.lang, "title"), f"{t(self.lang, 'db_saved')}\nSnapshot ID: {snapshot_id}\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, t(self.lang, "title"), f"Database save failed.\n{exc}")

    def load_latest_snapshot_from_db(self):
        try:
            df, meta = _load_latest_inventory_snapshot_from_db(self.operational_db_path)
            if df.empty:
                QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "db_no_snapshot"))
                return
            self.df_raw = df
            self.current_sheet = f"Operational DB snapshot #{meta.get('snapshot_id')}"
            self.mapping_report_text = f"Loaded from operational SQLite database. Snapshot: {meta.get('snapshot_id')} | Date: {meta.get('created_at')}"
            self.preview_df = df.head(200).copy()
            self.import_diagnosis_df = build_import_diagnosis(self.df_raw, self.preview_df, self.mapping_report_text) if 'build_import_diagnosis' in globals() else pd.DataFrame()
            self.import_summary_text = build_import_quick_summary(self.df_raw, self.current_sheet, self.mapping_report_text) if 'build_import_quick_summary' in globals() else self.mapping_report_text
            self.recompute()
            QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "db_loaded"))
        except Exception as exc:
            QMessageBox.critical(self, t(self.lang, "title"), f"Database load failed.\n{exc}")

    def _current_forecast_model(self):
        if hasattr(self, "forecast_model_combo"):
            return clean_text(self.forecast_model_combo.currentText()) or "Rule-based"
        return "Rule-based"

    def build_enhanced_forecast_df(self, df, horizon_days=30):
        working = self.build_network_prediction_df(df)
        if working is None or working.empty:
            return pd.DataFrame()
        working = working.copy()
        model = self._current_forecast_model()
        imported_monthly = pd.to_numeric(working.get("network_monthly_consumption", 0), errors="coerce").fillna(0)
        working["imported_monthly_consumption"] = np.round(imported_monthly, 2)
        working["generic_key"] = working.get("generic_name", working.get("medicine_name", "")).astype(str).map(normalize_search_text)
        try:
            movement_df = _load_movement_ledger(getattr(self, "movement_db_path", None))
        except Exception:
            movement_df = pd.DataFrame()
        ledger_proj = _movement_issue_projection(movement_df)
        if not ledger_proj.empty:
            working = working.merge(ledger_proj, on="generic_key", how="left")
        else:
            working["ledger_30d_issue_qty"] = 0
            working["ledger_projected_monthly_consumption"] = 0
        working["ledger_30d_issue_qty"] = pd.to_numeric(working.get("ledger_30d_issue_qty", 0), errors="coerce").fillna(0)
        working["ledger_projected_monthly_consumption"] = pd.to_numeric(working.get("ledger_projected_monthly_consumption", 0), errors="coerce").fillna(0)
        has_ledger = working["ledger_projected_monthly_consumption"] > 0
        if model == "Movement weighted":
            forecast_monthly = np.where(has_ledger, (working["ledger_projected_monthly_consumption"] * 0.60) + (imported_monthly * 0.40), imported_monthly)
            reason = np.where(has_ledger, "60% movement ledger + 40% imported monthly consumption", "No movement ledger yet; using imported monthly consumption")
        elif model == "Safety-adjusted":
            weighted = np.where(has_ledger, (working["ledger_projected_monthly_consumption"] * 0.60) + (imported_monthly * 0.40), imported_monthly)
            risk_mult = np.select(
                [working.get("highest_shortage_risk", "Low").astype(str).eq("Critical"), working.get("highest_shortage_risk", "Low").astype(str).eq("High"), working.get("store_support_status", "").astype(str).str.contains("gap", case=False, na=False)],
                [1.20, 1.12, 1.10],
                default=1.00
            )
            forecast_monthly = weighted * risk_mult
            reason = np.where(has_ledger, "Movement weighted demand plus safety factor for shortage/store gaps", "Imported monthly consumption plus safety factor for shortage/store gaps")
        else:
            forecast_monthly = imported_monthly
            reason = np.array(["Imported monthly consumption rule-based baseline"] * len(working))
        working["forecast_model"] = model
        working["forecast_monthly_demand"] = np.round(forecast_monthly, 2)
        working["forecast_horizon_demand"] = np.round(working["forecast_monthly_demand"] * (float(horizon_days) / 30.0), 2)
        available = pd.to_numeric(working.get("central_store_available", working.get("current_stock", 0)), errors="coerce").fillna(0)
        working["forecast_available_stock"] = np.round(available, 2)
        working["forecast_gap"] = np.round(np.maximum(working["forecast_horizon_demand"] - available, 0), 2)
        working["forecast_balance"] = np.round(available - working["forecast_horizon_demand"], 2)
        working["forecast_confidence"] = np.select(
            [has_ledger & (imported_monthly > 0), has_ledger | (imported_monthly > 0)],
            ["High", "Medium"],
            default="Low"
        )
        working["forecast_reason"] = reason
        working["suggested_network_action"] = np.where(
            working["forecast_gap"] > 0,
            "Prepare purchase/redistribution plan based on forecast gap",
            working.get("suggested_network_action", "Covered - routine monitoring")
        )
        return working.sort_values(["forecast_gap", "forecast_horizon_demand"], ascending=[False, False])

    def refresh_forecast_center(self, df):
        try:
            self.clear_layout(self.forecast_kpi_grid)
            horizon = int(getattr(self, "forecast_horizon_combo").currentText()) if hasattr(self, "forecast_horizon_combo") else 30
            working = self.build_enhanced_forecast_df(df, horizon_days=horizon)
            self._last_forecast_df = working.copy()
            if working.empty:
                self.set_table_from_df(self.forecast_table, pd.DataFrame())
                self.forecast_page_label.setText("0 / 0")
                return
            kpis = [
                (t(self.lang, "forecast_horizon"), f"{horizon}d", self._current_forecast_model()),
                (t(self.lang, "forecast_gap"), int(pd.to_numeric(working["forecast_gap"], errors="coerce").fillna(0).sum()), "total gap"),
                (t(self.lang, "balance_shortage"), int((pd.to_numeric(working["forecast_gap"], errors="coerce").fillna(0) > 0).sum()), "items"),
                (t(self.lang, "forecast_confidence"), str(working["forecast_confidence"].mode().iloc[0]) if not working["forecast_confidence"].empty else "Low", "overall"),
            ]
            for i, (title, value, sub) in enumerate(kpis):
                self.forecast_kpi_grid.addWidget(make_card(title, value, sub), i // 2, i % 2)
            cols = [c for c in [
                "generic_name", "dosage_form_group", "medicine_name", "imported_monthly_consumption",
                "ledger_30d_issue_qty", "ledger_projected_monthly_consumption", "forecast_monthly_demand",
                "forecast_horizon_demand", "forecast_available_stock", "forecast_balance", "forecast_gap",
                "forecast_confidence", "highest_shortage_risk", "store_support_status", "suggested_network_action", "forecast_reason"
            ] if c in working.columns]
            show = working[cols].copy()
            page_size = self._page_size_from_combo(self.forecast_page_size_combo, 100)
            page_df, page, total_pages = self._paged_slice(show, getattr(self, "_forecast_page", 0), page_size)
            self._forecast_page = page
            self.forecast_page_label.setText(f"{page+1} / {total_pages}")
            self.set_table_from_df(self.forecast_table, page_df)
            try:
                self.forecast_table.setWordWrap(True)
                self.forecast_table.verticalHeader().setDefaultSectionSize(42)
                for col_idx, width in {0:220, 1:130, 3:150, 4:140, 6:155, 9:120, 10:120, 11:110, 14:260, 15:360}.items():
                    if col_idx < self.forecast_table.columnCount():
                        self.forecast_table.setColumnWidth(col_idx, width)
            except Exception:
                pass
            if HAS_MATPLOTLIB:
                top = working.head(CHART_TOP_N)
                self.chart_forecast_gap.plot_bar(top["generic_name"].astype(str).tolist(), pd.to_numeric(top["forecast_gap"], errors="coerce").fillna(0).tolist(), t(self.lang, "top_gap_chart"))
                mix = working["forecast_confidence"].fillna("Low").value_counts().head(5)
                self.chart_forecast_mix.plot_pie(mix.index.tolist(), mix.values.tolist(), t(self.lang, "stock_mix_chart"))
        except Exception:
            traceback.print_exc()
            super().refresh_forecast_center(df)

    def save_current_forecast_to_db(self):
        try:
            if not hasattr(self, "_last_forecast_df") or self._last_forecast_df is None or self._last_forecast_df.empty:
                if hasattr(self, "df") and self.df is not None:
                    horizon = int(self.forecast_horizon_combo.currentText()) if hasattr(self, "forecast_horizon_combo") else 30
                    self._last_forecast_df = self.build_enhanced_forecast_df(self.filtered_df(), horizon_days=horizon)
            if self._last_forecast_df.empty:
                QMessageBox.information(self, t(self.lang, "title"), t(self.lang, "no_data"))
                return
            horizon = int(self.forecast_horizon_combo.currentText()) if hasattr(self, "forecast_horizon_combo") else 30
            forecast_id, path = _save_forecast_to_db(
                self._last_forecast_df,
                source_name=getattr(self, "current_sheet", ""),
                horizon_days=horizon,
                forecast_model=self._current_forecast_model(),
                db_path=self.operational_db_path,
            )
            self._update_db_status_label()
            QMessageBox.information(self, t(self.lang, "title"), f"{t(self.lang, 'db_saved')}\nForecast ID: {forecast_id}\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, t(self.lang, "title"), f"Forecast database save failed.\n{exc}")

# ======================= END V11 ADD-ON =======================

def main():
    app = QApplication(sys.argv)
    try:
        app.setWindowIcon(base.load_icon())
    except Exception:
        pass
    win = PharmaGuardMainWindow()
    win.showMaximized()
    sys.exit(app.exec())
if __name__ == "__main__":
    main()